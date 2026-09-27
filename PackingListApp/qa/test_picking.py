import sys
import os
import tempfile
import unittest
from collections import Counter
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
ROOT = Path(os.environ.get('PACKING_TEST_DATA', APP.parent))
sys.path.insert(0, str(APP / 'src'))
from docx import Document
from extractors import Batch, Item, Order, PackingError, parse_pdf
from picking import picking_rows
from documents import build_document
from engine import read_completed_doc


class PickingTests(unittest.TestCase):
    def test_summaries_use_original_product_and_variation_names(self):
        for group in ('SPX', 'FLASH', 'LAZADA_TIKTOK'):
            batches = [parse_pdf(p) for p in (ROOT / '20260926').rglob('*.pdf')]
            batches = [b for b in batches if b.group == group]
            approved = Document(ROOT / '20260926/000 Output' / f'{group}_Packing_Checklist_2026-09-26_1.docx')
            expected = Counter()
            for table in approved.tables[:-1]:
                for row in table.rows[2:]:
                    expected[(row.cells[1].text, row.cells[2].text)] += int(row.cells[3].text)
            self.assertEqual({(p, v): q for p, v, q in picking_rows(batches)}, dict(expected), group)

    def test_only_selected_orders_are_counted(self):
        batch = parse_pdf(ROOT / '20260926/SPX/20260926_SPX_PL_1.pdf')
        first = Batch(batch.path, batch.platform, batch.sha256, batch.orders[:1])
        self.assertEqual(sum(q for _, _, q in picking_rows([first])), first.orders[0].pieces)
        self.assertLess(sum(q for _, _, q in picking_rows([first])), batch.pieces)

    def test_unknown_items_preserve_size_grade_and_variation(self):
        items = [Item('New soap medical grade 1 L', 'Lavender', 2),
                 Item('New soap cosmetic grade 1 L', 'Lavender', 3),
                 Item('New soap medical grade 1 L', 'Rose', 4),
                 Item('New soap medical grade 1 L', 'Rose, 500mL', 5),
                 Item('New soap 100mL and 250mL', '—', 6)]
        batch = Batch('test.pdf', 'SPX', '', [Order('SPX', 'TEST', items)])
        rows = picking_rows([batch])
        self.assertEqual(len(rows), 5)
        self.assertIn(('New soap medical grade 1 L', 'Rose, 500mL', 5), rows)
        self.assertIn(('New soap medical grade 1 L', 'Lavender', 2), rows)
        self.assertIn(('New soap 100mL and 250mL', '—', 6), rows)

    def test_former_aliases_stay_separate_and_exact_repeats_combine(self):
        first = 'AK Fragrance Oil A-F 30mL & 100mL Perfume Collection for Men for Perfume Making'
        second = 'AK Fragrance Oil Best Seller 30mL & 100mL - Perfume Collection Men & Women for Perfume Making'
        items = [Item(first, 'Bgari Aqua (M),30 mL', 2), Item(first, 'Bgari Aqua (M),30 mL', 3),
                 Item(second, 'Bgari Aqua (M),30mL', 4)]
        rows = picking_rows([Batch('test.pdf', 'SPX', '', [Order('SPX', 'TEST', items)])])
        self.assertEqual(rows, [(first, 'Bgari Aqua (M),30 mL', 5), (second, 'Bgari Aqua (M),30mL', 4)])

    def test_summary_does_not_become_an_order_and_starts_new_page(self):
        batch = parse_pdf(ROOT / '20260926/SPX/20260926_SPX_PL_4.pdf')
        with tempfile.TemporaryDirectory(dir=APP / 'qa') as temp:
            path = Path(temp) / 'test.docx'
            build_document([batch], 'SPX', '20260926', path)
            d = Document(path)
            self.assertTrue(d.paragraphs[-3].paragraph_format.page_break_before)
            records = read_completed_doc(path, 'SPX')
            self.assertEqual(len(records), len(batch.orders))
            self.assertEqual(int(d.tables[-1].rows[-1].cells[3].text), batch.pieces)
            d.tables[-1].rows[-1].cells[3].text = '999'
            d.save(path)
            with self.assertRaises(PackingError):
                read_completed_doc(path, 'SPX')

    def test_automatic_import_accepts_unlabelled_shopee_for_courier_selection(self):
        import shutil
        import engine
        with tempfile.TemporaryDirectory(dir=APP / 'qa') as temp:
            temp = Path(temp)
            pdf = temp / 'packing.pdf'
            shutil.copyfile(ROOT / '20260926/SPX/20260926_SPX_PL_4.pdf', pdf)
            folder = temp / '20260926'
            result = engine.import_pdfs(folder, [str(pdf)])
            self.assertFalse(result['errors'])
            self.assertEqual(len(result['imported']), 1)
            unassigned = engine.scan(folder)['entries'][0]
            self.assertEqual((unassigned['platform'], unassigned['courier'], unassigned['status']), ('SHOPEE', 'Auto', 'COURIER'))
            labelled = temp / 'SPX_packing.pdf'
            pdf.rename(labelled)
            result = engine.import_pdfs(folder, [str(labelled)])
            self.assertEqual(len(result['imported']), 1)
            self.assertFalse(result['errors'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
