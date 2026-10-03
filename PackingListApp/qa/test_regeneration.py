"""History/selection regression checks with deterministic parsed source fixtures."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import sys

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
import engine
from generation_test_support import generate_approved
from extractors import Batch, Item, Order, PackingError, file_hash
from docx import Document


def layout(paths):
    return [{'path': str(p), 'pages': 2, 'orders': len(Document(p).tables)-1} for p in paths]


class RegenerationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.folder = Path(self.temp.name) / '20260926'
        (self.folder / 'SPX').mkdir(parents=True)
        self.items = {}
        self.parser = patch.object(engine, 'parse_pdf', side_effect=self.parse)
        self.parser.start()

    def tearDown(self):
        self.parser.stop()
        self.temp.cleanup()

    def parse(self, path, carrier=None, allow_unassigned=False):
        return Batch(str(path), 'SPX', file_hash(path), self.items[Path(path).name])

    def source(self, name, ids, quantity=2):
        path = self.folder / 'SPX' / name
        path.write_bytes(f'{name}:{ids}:{quantity}'.encode())
        self.items[name] = [Order('SPX', order, [Item('Test soap 1 L', 'Lemon', quantity)]) for order in ids]
        return path

    def scan(self):
        return engine.scan(self.folder)

    def generate(self, names=None, revision=False):
        entries = self.scan()['entries']
        selected = [{'path': e['path'], 'sha256': e['sha256'], 'revision': revision} for e in entries
                    if names is None or Path(e['path']).name in names]
        return generate_approved(self.folder, selected, verifier=layout)['outputs']

    def history(self):
        return json.loads((self.folder / '000 Output' / engine.HISTORY).read_text())

    def test_deleted_output_clears_repeat_memory_and_persists_reset(self):
        self.source('one.pdf', ['A'])
        first = self.generate()[0]
        original = self.history()['outputs'][first['filename']]
        Path(first['path']).unlink()
        entry = self.scan()['entries'][0]
        self.assertEqual(entry['status'], 'READY')
        self.assertEqual(entry['new_orders'], 1)
        self.assertNotIn('missing', entry['details'].lower())
        h = self.history()
        self.assertEqual(h['orders'], {})
        self.assertEqual(h['sources'], {})
        self.assertEqual(h['outputs'], {})
        self.assertEqual(h['archived_outputs'][-1]['order_snapshots'], original['order_snapshots'])
        selected = [{'path': entry['path'], 'sha256': entry['sha256']}]
        result = engine.generate(self.folder, selected, verifier=layout)
        self.assertFalse(result.get('confirmation_required', False))
        self.assertEqual(result['outputs'][0]['mode'], 'new')
        paragraphs = [p.text for p in Document(result['outputs'][0]['path']).paragraphs]
        self.assertIn('Total orders: 1', paragraphs)
        heading = next(p for p in paragraphs if p.startswith('Batch 1'))
        self.assertNotIn('order', heading)
        self.assertLess(paragraphs.index('Total orders: 1'), paragraphs.index(heading))

    def test_processed_file_can_generate_again_without_overwrite(self):
        self.source('one.pdf', ['A', 'B'])
        first = self.generate()[0]
        original_hash = file_hash(first['path'])
        self.assertEqual(self.scan()['entries'][0]['status'], 'DONE')
        second = self.generate()[0]
        self.assertTrue(second['filename'].endswith('_2.docx'))
        self.assertEqual(second['mode'], 'regenerated')
        self.assertEqual(file_hash(first['path']), original_hash)
        self.assertEqual({o['status'] for o in self.history()['outputs'].values()}, {'active'})
        self.assertEqual(len(engine.read_completed_doc(Path(second['path']), 'SPX')), 2)

    def test_deleting_output_marks_all_its_pdfs_missing_then_recovers(self):
        self.source('one.pdf', ['A'])
        self.source('two.pdf', ['B'])
        first = self.generate()[0]
        Path(first['path']).unlink()
        result = self.scan()
        self.assertEqual([e['status'] for e in result['entries']], ['READY', 'READY'])
        self.assertEqual(result['outputs'], [])
        second = self.generate()[0]
        self.assertTrue(second['filename'].endswith('_1.docx'))
        self.assertEqual([e['status'] for e in self.scan()['entries']], ['DONE', 'DONE'])
        links = self.history()['outputs'][second['filename']]['sources']
        self.assertEqual({s['path'] for s in links}, {'SPX/one.pdf', 'SPX/two.pdf'})
        self.assertTrue(all(s['sha256'] and s['order_keys'] for s in links))

    def test_surviving_copy_covers_orders_until_every_copy_is_deleted(self):
        self.source('one.pdf', ['A'])
        first, second = self.generate()[0], self.generate()[0]
        Path(second['path']).unlink()
        self.assertEqual(self.scan()['entries'][0]['status'], 'DONE')
        Path(first['path']).unlink()
        self.assertEqual(self.scan()['entries'][0]['status'], 'READY')
        third = self.generate()[0]
        self.assertTrue(third['filename'].endswith('_1.docx'))

    def test_overlapping_selected_pdfs_include_each_order_once_and_link_both(self):
        self.source('one.pdf', ['A', 'B'])
        first = self.generate()[0]
        self.source('two.pdf', ['B', 'C'])
        second = self.generate()[0]
        doc = engine.read_completed_doc(Path(second['path']), 'SPX')
        self.assertEqual(set(doc), {'SHOPEE:A', 'SHOPEE:B', 'SHOPEE:C'})
        self.assertEqual(int(Document(second['path']).tables[-1].rows[-1].cells[3].text), 6)
        self.assertEqual(len(self.history()['outputs'][second['filename']]['sources']), 2)
        Path(second['path']).unlink()
        entries = {e['path']: e for e in self.scan()['entries']}
        self.assertEqual(entries['SPX/one.pdf']['status'], 'DONE')
        self.assertEqual(entries['SPX/two.pdf']['status'], 'READY')
        self.assertEqual(entries['SPX/two.pdf']['new_orders'], 1)
        third = self.generate(['two.pdf'])[0]
        self.assertEqual(set(engine.read_completed_doc(Path(third['path']), 'SPX')), {'SHOPEE:B', 'SHOPEE:C'})

    def test_selected_duplicate_pdf_is_recorded_and_regeneratable(self):
        self.source('one.pdf', ['A'])
        self.source('two.pdf', ['A'])
        first = self.generate()[0]
        self.assertEqual(len(self.history()['outputs'][first['filename']]['sources']), 2)
        self.assertEqual(len(self.history()['sources']), 2)
        second = self.generate(['two.pdf'])[0]
        self.assertEqual(len(engine.read_completed_doc(Path(second['path']), 'SPX')), 1)

    def test_deleted_output_resets_changed_source_when_no_copy_survives(self):
        self.source('one.pdf', ['A'])
        first = self.generate()[0]
        Path(first['path']).unlink()
        self.source('one.pdf', ['A'], quantity=5)
        self.assertEqual(self.scan()['entries'][0]['status'], 'READY')
        second = self.generate()[0]
        self.assertEqual(int(Document(second['path']).tables[-1].rows[-1].cells[3].text), 5)

    def test_revision_invalidates_all_old_copies_for_missing_detection(self):
        self.source('one.pdf', ['A'])
        first, second = self.generate()[0], self.generate()[0]
        self.source('one.pdf', ['A'], quantity=7)
        third = self.generate(revision=True)[0]
        h = self.history()
        self.assertEqual(h['outputs'][first['filename']]['status'], 'superseded')
        self.assertEqual(h['outputs'][second['filename']]['status'], 'superseded')
        Path(third['path']).unlink()
        self.assertEqual(self.scan()['entries'][0]['status'], 'READY')

    def test_missing_legacy_doc_can_be_regenerated_from_logged_sources(self):
        self.source('one.pdf', ['A'])
        out = self.folder / '000 Output'
        out.mkdir()
        (out / 'Processed_Files.txt').write_text('PACKING LIST PROCESSING LOG\n\nSPX_Packing_Checklist_2026-09-26_1.docx\nStatus: Completed\n  - SPX/one.pdf | 1 orders\n')
        self.assertEqual(self.scan()['entries'][0]['status'], 'READY')
        second = self.generate()[0]
        self.assertTrue(second['filename'].endswith('_1.docx'))

    def test_old_app_history_upgrades_even_if_word_is_already_deleted(self):
        self.source('one.pdf', ['A'])
        first = self.generate()[0]
        h = self.history()
        for value in h['outputs'].values():
            value.pop('sources')
            value.pop('order_fingerprints')
        engine.atomic_json(self.folder / '000 Output' / engine.HISTORY, h)
        Path(first['path']).unlink()
        self.assertEqual(self.scan()['entries'][0]['status'], 'READY')
        self.assertTrue(self.history()['archived_outputs'][-1]['sources'])
        self.generate()

    def test_ignored_content_changes_require_review_before_regeneration(self):
        self.source('one.pdf', ['A'])
        self.generate()
        self.source('one.pdf', ['A'], quantity=3)
        e = self.scan()['entries'][0]
        engine.acknowledge(self.folder, {'path': e['path'], 'sha256': e['sha256']})
        self.assertTrue(self.scan()['entries'][0]['requires_revision'])
        with self.assertRaises(PackingError):
            self.generate()
        self.generate(revision=True)

    def test_deleted_highest_number_is_reused_and_old_record_archived(self):
        self.source('one.pdf', ['A'])
        first, second, third = self.generate()[0], self.generate()[0], self.generate()[0]
        old_id = self.history()['outputs'][third['filename']]['record_id']
        Path(third['path']).unlink()
        self.source('two.pdf', ['B'])
        replacement = self.generate(['two.pdf'])[0]
        self.assertEqual(replacement['filename'], third['filename'])
        h = self.history()
        self.assertNotEqual(h['outputs'][replacement['filename']]['record_id'], old_id)
        self.assertEqual(h['archived_outputs'][-1]['record_id'], old_id)
        shown = next(o for o in self.scan()['outputs'] if o['filename'] == replacement['filename'])
        self.assertEqual([s['path'] for s in shown['included_sources']], ['SPX/two.pdf'])
        self.assertEqual(shown['included_sources'][0]['orders'][0]['order_id'], 'B')
        self.assertTrue(Path(first['path']).is_file() and Path(second['path']).is_file())

    def test_missing_middle_number_is_not_filled(self):
        self.source('one.pdf', ['A'])
        first, second, third = self.generate()[0], self.generate()[0], self.generate()[0]
        Path(second['path']).unlink()
        fourth = self.generate()[0]
        self.assertTrue(fourth['filename'].endswith('_4.docx'))
        self.assertEqual([o['sequence'] for o in self.scan()['outputs']], [1, 3, 4])

    def test_reused_filename_does_not_make_old_pdf_look_processed(self):
        self.source('one.pdf', ['A'])
        first = self.generate()[0]
        Path(first['path']).unlink()
        self.source('two.pdf', ['B'])
        replacement = self.generate(['two.pdf'])[0]
        self.assertEqual(first['filename'], replacement['filename'])
        states = {e['path']: e['status'] for e in self.scan()['entries']}
        self.assertEqual(states, {'SPX/one.pdf': 'READY', 'SPX/two.pdf': 'DONE'})
        next_file = self.generate(['one.pdf'])[0]
        self.assertTrue(next_file['filename'].endswith('_2.docx'))

    def test_completed_contents_keep_the_generated_version(self):
        self.source('one.pdf', ['A'], quantity=2)
        first = self.generate()[0]
        self.source('one.pdf', ['A'], quantity=7)
        second = self.generate(revision=True)[0]
        shown = {o['filename']: o for o in self.scan()['outputs']}
        self.assertEqual(shown[first['filename']]['included_sources'][0]['orders'][0]['items'][0]['pieces'], 2)
        self.assertEqual(shown[second['filename']]['included_sources'][0]['orders'][0]['items'][0]['pieces'], 7)

    def test_sequences_are_numeric_per_platform_and_ignore_missing_records(self):
        out = self.folder / '000 Output'
        out.mkdir()
        history = {'outputs': {}}
        for group, number in [('FLASH', 10), ('SPX', 10), ('SPX', 2), ('SPX', 1), ('LAZADA_TIKTOK', 1), ('FLASH', 2)]:
            name = f'{group}_Packing_Checklist_2026-09-26_{number}.docx'
            (out / name).write_bytes(b'placeholder for filename/list tests')
            history['outputs'][name] = {'group': group, 'orders': [], 'order_snapshots': {}, 'record_id': name, 'status': 'active'}
        missing = 'SPX_Packing_Checklist_2026-09-26_99.docx'
        history['outputs'][missing] = {'group': 'SPX', 'orders': [], 'record_id': missing, 'status': 'active'}
        shown = engine.completed_outputs(self.folder, history)
        self.assertEqual([(o['group'], o['sequence']) for o in shown], [('SPX', 1), ('SPX', 2), ('SPX', 10), ('FLASH', 2), ('FLASH', 10), ('LAZADA_TIKTOK', 1)])
        self.assertEqual(engine.next_output_filename(out, 'SPX', '20260926'), 'SPX_Packing_Checklist_2026-09-26_11.docx')
        self.assertEqual(engine.next_output_filename(out, 'LAZADA_TIKTOK', '20260926'), 'LAZADA_TIKTOK_Packing_Checklist_2026-09-26_2.docx')

    def test_file_appearing_during_generation_is_not_overwritten(self):
        self.source('one.pdf', ['A'])
        selection = [{'path': e['path'], 'sha256': e['sha256']} for e in self.scan()['entries']]
        out = self.folder / '000 Output'
        def race(paths):
            (out / paths[0].name).write_bytes(b'Another file arrived')
            return layout(paths)
        with self.assertRaisesRegex(PackingError, 'appeared while generating'):
            generate_approved(self.folder, selection, verifier=race)
        self.assertEqual(next(out.glob('*.docx')).read_bytes(), b'Another file arrived')
        self.assertFalse((out / engine.JOURNAL).exists())


if __name__ == '__main__':
    unittest.main(verbosity=2)
