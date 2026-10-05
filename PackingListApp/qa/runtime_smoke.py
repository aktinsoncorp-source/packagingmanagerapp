"""Exercise the portable standard-runtime package through its public launcher."""
from pathlib import Path
import json
import shutil
import subprocess
import sys
import tempfile

import pdfplumber
from openpyxl import load_workbook
from extractors import parse_pdf, clean
from names import apply_names, load_catalog
from documents import item_prices

APP = Path(__file__).resolve().parents[1]
release = Path(sys.argv[1]).resolve()
work = Path(tempfile.mkdtemp(prefix='runtime-quantity-价格-', dir=APP / 'qa'))
folder = work / '20260926'
fixture = APP / 'qa/names-regression-fixtures/20260926'
for relative in ('SPX/20260926_SPX_PL_4.pdf', 'FLASH/20260926_Flash_PL_3.pdf',
                 'LAZADA & TIKTOK/20260926_Lazada_PL_1.pdf', 'LAZADA & TIKTOK/20260926_Tiktok_PL_1.pdf'):
    target = folder / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(fixture / relative, target)
catalog = work / 'Names 价格.json'
shutil.copy2(APP.parent / 'Packing List Manager/ProductNames.json', catalog)


def call(action, **fields):
    request = work / 'request with spaces 价格.json'
    response = work / 'response with spaces 价格.json'
    request.write_text(json.dumps(dict(action=action, folder=str(folder),
                                      catalog_path=str(catalog), **fields), ensure_ascii=False), encoding='utf-8')
    subprocess.run([str(release / 'App Files/PackingEngine.exe'), '--request', str(request),
                    '--response', str(response)], check=True, timeout=240,
                   creationflags=subprocess.CREATE_NO_WINDOW)
    result = json.loads(response.read_text(encoding='utf-8'))
    assert result['ok'], result
    return result


definitions = call('names_list')
book = work / 'Names 价格.xlsx'
call('names_export', path=str(book))
call('names_import', path=str(book), rows=definitions['rows'])
scan = call('scan')
assert len(scan['entries']) == 4, scan
sales = call('sales_export')
assert Path(sales['path']).name.startswith('2026-09-26_Item_Sales_')
assert load_workbook(sales['path']).sheetnames == ['Item Sales', 'Order Items', 'Export Info']
selected = [dict(path=entry['path'], sha256=entry['sha256']) for entry in scan['entries']]
generated = call('generate', selected=selected)
if generated.get('confirmation_required'):
    generated = call('generate', selected=selected, repeat_confirmation=generated['repeat_token'], repeat_action='skip')
assert len(generated['outputs']) == 3, generated
orders = 0
expected = {order.order_id: order for batch in apply_names(
    [parse_pdf(path) for path in folder.rglob('*.pdf') if '000 Output' not in path.parts],
    load_catalog(catalog)) for order in batch.orders}
for output in generated['outputs']:
    with pdfplumber.open(output['path']) as pdf:
        for page in pdf.pages:
            assert abs(page.width - 595.3) < 1 and abs(page.height - 841.9) < 1
            for table in page.extract_tables():
                if table and table[0] and table[0][0] and table[0][0].startswith(('Order ID:', 'Order Number:')):
                    orders += 1
                    identifier = table[0][0].split(':', 1)[1].split()[0]
                    order = expected[identifier]
                    rows = [[clean(cell) for cell in row if clean(cell)] for row in table[2:]]
                    wanted = [['[ ]', clean(item.product), clean(item.variation), str(item.pieces)] +
                              (item_prices(item) if order.platform in ('SPX', 'FLASH') else []) for item in order.items]
                    assert rows == wanted, (identifier, rows, wanted)
                    if output['filename'].startswith(('SPX_', 'FLASH_')):
                        headings = [clean(cell) for cell in table[1] if clean(cell)]
                        assert headings[-2:] == ['Unit Price', 'Total Price']
assert orders == 17, orders
rescan = call('scan')
assert all(entry['status'] == 'DONE' for entry in rescan['entries']), rescan
report = dict(ok=True, version='1.11.1', work=str(work), orders=orders,
              packing_pdfs=[output['path'] for output in generated['outputs']],
              sales=sales['path'], names_excel_roundtrip=True, unicode_paths=True,
              a4=True, source_price_columns=True)
(APP / 'qa/runtime-verification.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
print(json.dumps(report, indent=2))
