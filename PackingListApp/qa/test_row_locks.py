import copy
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
APP=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(APP/'src'))
import names
from extractors import Item, Order, Batch, PackingError, source_price

class RowLockTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=APP/'qa');self.path=Path(self.temp.name)/'names.json';self.book=Path(self.temp.name)/'names.xlsx'
        self.catalog=names.save_catalog(self.path,[{'original_product':'SOAP','original_variation':'100ML','alternative_product':'SOAP','alternative_variation':'100ML','locked':True,'aliases':[{'original_product':'LISTING','original_variation':'SMALL'}]}],0)
    def tearDown(self):self.temp.cleanup()
    def test_lock_persists_and_blocks_changes_and_deletion(self):
        self.assertTrue(names.load_catalog(self.path)['rows'][0]['locked'])
        rows=copy.deepcopy(self.catalog['rows']);rows[0]['alternative_product']='CHANGED'
        with self.assertRaisesRegex(PackingError,'Unlock'):names.save_catalog(self.path,rows,1)
        with self.assertRaisesRegex(PackingError,'Unlock'):names.save_catalog(self.path,[],1)
        rows[0]['locked']=False
        self.assertEqual(names.save_catalog(self.path,rows,1)['rows'][0]['alternative_product'],'CHANGED')
    def test_locked_linked_original_cannot_be_changed_in_excel(self):
        names.export_excel(self.book,self.catalog)
        book=load_workbook(self.book);sheet=book['Names'];self.assertTrue(sheet.protection.sheet);self.assertTrue(sheet['A3'].protection.locked)
        sheet['A3']='CHANGED ALIAS';book.save(self.book);book.close()
        with self.assertRaisesRegex(PackingError,'locked'):names.import_excel(self.book,self.catalog['rows'])
        self.assertEqual(names.load_catalog(self.path),self.catalog)
    def test_unchanged_locked_excel_round_trip(self):
        names.export_excel(self.book,self.catalog)
        self.assertEqual(names.import_excel(self.book,self.catalog['rows'])['rows'],self.catalog['rows'])
    def test_locked_linked_alternatives_cannot_be_split_in_excel(self):
        names.export_excel(self.book,self.catalog)
        book=load_workbook(self.book);sheet=book['Names'];sheet['C2']='FIRST';sheet['C3']='SECOND';book.save(self.book);book.close()
        with self.assertRaisesRegex(PackingError,'locked'):names.import_excel(self.book,self.catalog['rows'])
        self.assertEqual(names.load_catalog(self.path),self.catalog)
    def test_locked_rows_not_offered_for_merges(self):
        rows=copy.deepcopy(self.catalog['rows']);row=copy.deepcopy(rows[0]);row['id']='other';row['original_product']='OTHER';row['aliases']=[];row['locked']=False;rows.append(row)
        self.assertFalse(names.prepare_save(rows)['merges'])
        with self.assertRaisesRegex(PackingError,'Unlock'):names.merge_rows(rows,[[r['id'] for r in rows]])
    def test_source_price_and_name_mapping_keep_price_and_old_fingerprint(self):
        self.assertEqual(source_price('1,250.50','test'),'1250.50');self.assertIsNone(source_price('','test'))
        with self.assertRaises(PackingError):source_price('not a price','test')
        raw=Order('SPX','A',[Item('SOAP','100ML',2,'10.50')]);plain=Order('SPX','A',[Item('SOAP','100ML',2)])
        self.assertEqual(raw.fingerprint(),plain.fingerprint())
        printed=names.apply_names([Batch('x','SPX','hash',[raw])],self.catalog)
        self.assertEqual(printed[0].orders[0].items[0].unit_price,'10.50')

if __name__=='__main__':unittest.main()
