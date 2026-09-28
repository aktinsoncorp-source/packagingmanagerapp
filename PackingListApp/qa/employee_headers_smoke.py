"""Check employee fields in packaged outputs for all four supported platforms."""
from pathlib import Path
import json
import os
import shutil
import subprocess
import sys
import tempfile

APP=Path(__file__).resolve().parents[1]
ROOT=APP.parent
sys.path.insert(0,str(APP/'src'))
from docx import Document
import pdfplumber
from extractors import parse_pdf
from documents import verify_word
from engine import VERSION,read_completed_doc

exe=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else ROOT/'Packing List Manager/App Files/PackingEngine.exe'
fixture=Path(os.environ.get('PACKING_TEST_DATA',APP/'qa/names-regression-fixtures'))/'20260926'
work=Path(tempfile.mkdtemp(prefix='employee-headers-',dir=APP/'qa'))
folder=work/'20260926'
catalog=work/'ProductNames.json'
smallest={}
for path in fixture.rglob('*.pdf'):
    batch=parse_pdf(path)
    if batch.platform not in smallest or len(batch.orders)<len(smallest[batch.platform].orders):
        smallest[batch.platform]=batch
assert set(smallest)=={'SPX','FLASH','LAZADA','TIKTOK'}
for platform,batch in smallest.items():
    dest=folder/platform/Path(batch.path).name
    dest.parent.mkdir(parents=True,exist_ok=True)
    shutil.copyfile(batch.path,dest)

def call(action, **kwargs):
    result = _call(action, **kwargs)
    if action == 'generate' and result.get('confirmation_required'):
        result = _call(action, **kwargs, repeat_confirmation=result['repeat_token'])
    return result


def _call(action,**kwargs):
    request=work/'request.json';response=work/'response.json'
    request.write_text(json.dumps({'action':action,'folder':str(folder),'catalog_path':str(catalog),**kwargs}),encoding='utf-8')
    response.unlink(missing_ok=True)
    subprocess.run([str(exe),'--request',str(request),'--response',str(response)],check=True,timeout=240,creationflags=subprocess.CREATE_NO_WINDOW)
    result=json.loads(response.read_text(encoding='utf-8'))
    assert result['ok'],result
    return result

scan=call('scan')
assert scan['version']==VERSION
assert len(scan['entries'])==4 and all(e['status']=='READY' for e in scan['entries'])
result=call('generate', output_format='docx',selected=[{'path':e['path'],'sha256':e['sha256']} for e in scan['entries']])
paths=[Path(o['path']) for o in result['outputs']]
assert len(paths)==3
layout=verify_word(paths,export_pdf=True)
reports=[]
for path in paths:
    group=path.name.split('_Packing_')[0]
    batches=[b for b in smallest.values() if b.group==group]
    expected_orders={o.order_id:o for b in batches for o in b.orders}
    doc=Document(path)
    assert len(doc.tables)==len(expected_orders)+1
    for table in doc.tables[:-1]:
        head=table.cell(0,0).text
        order_id=head.split('\t')[0].split(': ')[1]
        order=expected_orders[order_id]
        label='Order Number' if order.platform=='LAZADA' else 'Order ID'
        assert head==f'{label}: {order_id}\tEmployee:__________________\tNumber __________\tTime __________',head
        assert sum(int(r.cells[3].text) for r in table.rows[2:])==order.pieces
    assert len(read_completed_doc(path,group))==len(expected_orders)
    assert int(doc.tables[-1].rows[-1].cells[3].text)==sum(b.pieces for b in batches)
    with pdfplumber.open(path.with_suffix('.pdf')) as pdf:
        header_count=0
        for page in pdf.pages:
            words=page.extract_words(x_tolerance=1,y_tolerance=2)
            for word in [w for w in words if w['text']=='Order']:
                line=[w for w in words if abs(w['top']-word['top'])<1.5]
                texts=[w['text'] for w in line]
                assert 'Employee:__________________' in texts and 'Number' in texts and 'Time' in texts,texts
                employee=next(w for w in line if w['text'].startswith('Employee:'))
                number=next(w for w in line if w['text']=='Number')
                time=next(w for w in line if w['text']=='Time')
                assert word['x0']<employee['x0']<employee['x1']<number['x0']<time['x0']
                assert max(w['x1'] for w in line)<page.width-22
                header_count+=1
        assert header_count==len(expected_orders),(header_count,len(expected_orders))
        reports.append({'group':group,'path':str(path),'orders':header_count,'pieces':sum(b.pieces for b in batches),'pages':len(pdf.pages),'all_fields_on_one_line':True})
assert all(e['status']=='DONE' for e in call('scan')['entries'])
report={'ok':True,'version':VERSION,'work':str(work),'platforms':sorted(smallest),'outputs':reports,'word_layout':layout,'source_fingerprints_preserved':True}
(APP/'qa/employee-1.6.1-verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
