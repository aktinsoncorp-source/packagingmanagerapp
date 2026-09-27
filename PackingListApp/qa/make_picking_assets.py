"""Distill approved September 26 picking names/layout without editing originals."""
from collections import Counter
from copy import deepcopy
from pathlib import Path
import hashlib
import json
import re
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'PackingListApp/src'))
from picking import sizes
from docx import Document
from docx.oxml.ns import qn

NAMES = {
    'AOS': 'AOS liquid (alpha olefin sulfonate)',
    'Anti foam': 'Anti foam (defoamer)',
    'Baking Soda': 'Baking soda (aluminum free)',
    'Caustic Soda': 'Caustic soda flakes',
    'CDEA': 'CDEA (cocodiethanolamide)',
    'CMC': 'CMC (technical grade)',
    'CTC': 'CTC (cetrimonium chloride)',
    'Citric Acid': 'Citric acid (food grade)',
    'Coco Betaine': 'Coco betaine / CAPB',
    'DOWANOL': 'Dowanol DPM',
    'DPG': 'DPG (dipropylene glycol)',
    'Essential Oil': 'Essential oil with dropper',
    'Fabric Conditioner': 'Fabric softener fragrance oil',
    'Fixative': 'Fixative',
    'Fragrance Oil': 'Perfume fragrance oil',
    'Household Scent': 'Household fragrance oil',
    'Industrial Salt': 'Industrial salt',
    'Mineral Oil': 'Mineral oil (USP, unscented)',
    'Oxygen Bleach': 'Oxygen bleach (sodium percarbonate)',
    'Petroleum Jelly': 'Petroleum jelly (USP)',
    'PG': 'PG (propylene glycol, food grade)',
    'Premix': 'Premix alcohol / perfume base',
    'Pure Isopropyl': 'Isopropyl alcohol 99.9%',
    'Silicone Emulsion': 'Silicone emulsion',
    'Soda Ash': 'Soda ash (sodium carbonate)',
    'Sodium Benzoate': 'Sodium benzoate (food grade)',
    'Stearic Acid': 'Stearic acid (triple pressed)',
    'Talcum': 'Talc powder',
    'Tergitol': 'Tergitol NP-10',
    'Vegetable Glycerine': 'Vegetable glycerine (USP)',
    'Zeolite': 'Zeolite powder',
}
assets = ROOT / 'PackingListApp/src/assets'
aliases = {}
evidence = []
for group in ['SPX', 'FLASH', 'LAZADA_TIKTOK']:
    source = ROOT / '20260926/000 Output' / f'{group}_Packing_Checklist_2026-09-26_1.docx'
    d = Document(source)
    expected = Counter({(r.cells[1].text, r.cells[2].text): int(r.cells[3].text) for r in d.tables[-1].rows[1:-1]})
    actual = Counter()
    for t in d.tables[:-1]:
        for r in t.rows[2:]:
            product, variation, qty = [c.text for c in r.cells[1:]]
            if product.startswith('AK EDTA'):
                name = 'EDTA (tetrasodium)' if 'Tetrasodium' in product else 'EDTA (neutralizer)'
            elif product.startswith('AK SLES'):
                name = 'SLES gel 70%' if '70%' in product else 'SLES gel'
            else:
                name = next(value for prefix, value in NAMES.items() if product.startswith('AK ' + prefix))
            size = sizes(variation) or sizes(product)
            assert len(size) == 1, (product, variation, size)
            if name in ('Perfume fragrance oil', 'Household fragrance oil', 'Fabric softener fragrance oil', 'Essential oil with dropper'):
                scent = re.sub(r',?\s*\d+(?:\.\d+)?\s*(?:ml|liters?|l)\b', '', variation, flags=re.I).strip(' ,')
                label = scent + ', ' + size[0]
            else:
                label = size[0]
            pair = (product, variation)
            value = {'product': product, 'variation': variation, 'name': name, 'size': label}
            assert pair not in aliases or aliases[pair] == value
            aliases[pair] = value
            actual[(name, label)] += int(qty)
    assert actual == expected, (group, actual - expected, expected - actual)
    evidence.append(f'{source}\nSHA256: {hashlib.sha256(source.read_bytes()).hexdigest()}\n{len(expected)} summary rows; {sum(expected.values())} pieces')
    if group == 'FLASH':
        continue
    keep = [deepcopy(p._p) for p in d.paragraphs[-3:]] + [deepcopy(d.tables[-1]._tbl)]
    for node in list(d._element.body):
        if node.tag != qn('w:sectPr'):
            d._element.body.remove(node)
    for node in keep:
        d._element.body.insert(len(d._element.body)-1, node)
    table = d.tables[0]
    for row in list(table.rows)[2:-1]:
        table._tbl.remove(row._tr)
    for row, texts in [(table.rows[1], ['[ ]', 'Product', 'Variation', '1']), (table.rows[-1], ['', 'TOTAL PIECES', '', '1'])]:
        for cell, text in zip(row.cells, texts):
            cell.paragraphs[0].runs[0].text = text
    d.paragraphs[0].runs[0].text = 'Product Picking Summary'
    d.paragraphs[1].runs[0].text = 'Date | Orders | Pieces'
    for p in d.sections[0].footer.paragraphs:
        p.clear()
    d.save(assets / ('summary_blue.docx' if group == 'LAZADA_TIKTOK' else 'summary_orange.docx'))
(assets / 'picking_aliases.json').write_text(json.dumps(list(aliases.values()), indent=2, ensure_ascii=False), encoding='utf-8')
(ROOT / 'PackingListApp/qa/summary-artifact.md').write_text(
    '# Product picking summary template\n\nApproved references, unchanged:\n\n' + '\n\n'.join(evidence) +
    '\n\nClone the final three paragraphs and four-column table. Title uses pageBreakBefore. '
    'Grid 680/5443/3402/1247 twips; Arial 9 pt; repeated header; gray borders; '
    'orange F8CBAD or blue BDD7EE. Each item row stays together; the summary may continue '
    'on additional A4 pages. The final item stays with TOTAL PIECES. '
    'Known product/variation aliases match the approved summaries exactly. Unlisted combinations retain source names.\n', encoding='utf-8')
print(f'Created summary templates and {len(aliases)} verified product/variation aliases.')
