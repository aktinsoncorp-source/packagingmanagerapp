"""Verify courier controls and continued Shopee orders with the packaged app."""
from pathlib import Path
import json, os, shutil, subprocess, sys, tempfile
APP=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(APP/'src'))
from docx import Document
import pdfplumber
from documents import verify_word
from extractors import parse_pdf
release=Path(sys.argv[1]).resolve()
source=Path(sys.argv[2]).resolve()
fixture=Path(os.environ.get('PACKING_TEST_DATA',APP/'qa/names-regression-fixtures'))/'20260926'
work=Path(tempfile.mkdtemp(prefix='couriers-1.7-',dir=APP/'qa'))
folder=work/'ui'/'20260928'
for src,relative in [(source,'SPX/PACKING LIST 1.pdf'),
 (fixture/'FLASH/20260926_Flash_PL_3.pdf','FLASH/SPX_old.pdf'),
 (fixture/'SPX/20260926_SPX_PL_4.pdf','packing.pdf'),
 (fixture/'LAZADA & TIKTOK/20260926_Lazada_PL_1.pdf','misc/lazada.pdf'),
 (fixture/'LAZADA & TIKTOK/20260926_Tiktok_PL_1.pdf','tiktok.pdf')]:
    dest=folder/relative;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(src,dest)
from pypdf import PdfWriter
writer=PdfWriter();writer.add_blank_page(width=595,height=842);writer.write(str(folder/'blank.pdf'))
(folder/'broken.pdf').write_bytes(b'%PDF-invalid')
subprocess.run([str(release/'Packing List Manager.exe'),'--self-test',str(folder),
 str(work/'main.png'),str(work/'history.png'),str(work/'names.png'),'--check-couriers'],
 check=True,timeout=240,creationflags=subprocess.CREATE_NO_WINDOW)
ui=json.loads((release/'SelfTest.json').read_text(encoding='utf-8'))
assert ui['ok'] and ui['rows']==5 and ui['selectable']==4 and ui['courierChangesWork'],ui
print(json.dumps({'ui':ui,'work':str(work)}),flush=True)
folder=work/'generation'/'20260928'
dest=folder/'SPX'/source.name;dest.parent.mkdir(parents=True);shutil.copyfile(source,dest)
def call(action,**kwargs):
    request=work/'request.json';response=work/'response.json'
    request.write_text(json.dumps({'action':action,'folder':str(folder),'catalog_path':str(work/'ProductNames.json'),**kwargs}),encoding='utf-8')
    response.unlink(missing_ok=True)
    subprocess.run([str(release/'App Files/PackingEngine.exe'),'--request',str(request),'--response',str(response)],check=True,timeout=240,creationflags=subprocess.CREATE_NO_WINDOW)
    data=json.loads(response.read_text(encoding='utf-8'));assert data['ok'],data;return data
scan=call('scan');entry=scan['entries'][0]
assert len(scan['entries'])==1 and entry['platform']=='SPX' and entry['courier']=='Auto',scan
result=call('generate',selected=[{'path':entry['path'],'sha256':entry['sha256'],'courier':entry['courier'],'platform':entry['platform']}])
path=Path(result['outputs'][0]['path'])
expected=parse_pdf(dest);doc=Document(path)
assert len(expected.orders)==127 and expected.pieces==300
assert len(doc.tables)==128
for order,table in zip(expected.orders,doc.tables[:-1]):
    assert order.order_id in table.cell(0,0).text
    assert 'Employee:__________________' in table.cell(0,0).text
    assert sum(int(r.cells[3].text) for r in table.rows[2:])==order.pieces
order=next(o for o in expected.orders if o.order_id=='2609263ADBUPHF')
assert len(order.items)==11 and order.pieces==14
assert int(doc.tables[-1].rows[-1].cells[3].text)==300
layout=verify_word([path],export_pdf=True)
with pdfplumber.open(path.with_suffix('.pdf')) as pdf:
    target_pages=[i+1 for i,p in enumerate(pdf.pages) if '2609263ADBUPHF' in (p.extract_text() or '')]
    assert len(target_pages)==1
    pages=len(pdf.pages)
assert call('scan')['entries'][0]['status']=='DONE'
report={'ok':True,'version':scan['version'],'work':str(work),'ui':ui,'docx':str(path),
 'orders':127,'pieces':300,'continued_order_lines':11,'continued_order_pieces':14,
 'word_layout':layout,'target_page':target_pages[0],'pages':pages}
(APP/'qa/couriers-1.7-verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
