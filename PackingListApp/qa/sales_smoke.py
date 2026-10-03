"""Verify the packaged quantity-export API and its actual Windows UI action."""
from pathlib import Path
import json
import shutil
import subprocess
import sys
import tempfile
from openpyxl import load_workbook
APP = Path(__file__).resolve().parents[1]
release = Path(sys.argv[1]).resolve()
work = Path(tempfile.mkdtemp(prefix='sales-smoke-', dir=APP/'qa'))
folder = work/'20260926'
fixture = APP/'qa/names-regression-fixtures/20260926'
for relative in ['SPX/20260926_SPX_PL_4.pdf','FLASH/20260926_Flash_PL_3.pdf','LAZADA & TIKTOK/20260926_Lazada_PL_1.pdf','LAZADA & TIKTOK/20260926_Tiktok_PL_1.pdf']:
    target=folder/relative
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(fixture/relative, target)
shutil.copyfile(fixture/'SPX/20260926_SPX_PL_4.pdf', folder/'duplicate.pdf')
def call(action):
    request=work/'request.json';response=work/'response.json'
    request.write_text(json.dumps({'action':action,'folder':str(folder),'catalog_path':str(work/'names.json')}),encoding='utf-8')
    response.unlink(missing_ok=True)
    subprocess.run([str(release/'App Files/PackingEngine.exe'),'--request',str(request),'--response',str(response)],check=True,timeout=90,creationflags=subprocess.CREATE_NO_WINDOW)
    data=json.loads(response.read_text(encoding='utf-8'))
    assert data['ok'],data
    return data
scan=call('scan');assert scan['version']=='1.11.1'
result=call('sales_export')
assert result['orders']==17 and result['quantity']==30,result
assert result['filename']=='2026-09-26_Item_Sales_1.xlsx'
book=load_workbook(result['path'],data_only=False)
assert book.sheetnames==['Item Sales','Order Items','Export Info']
assert sum(row[3] for row in list(book['Item Sales'].values)[1:])==30
assert sum(row[8] for row in list(book['Order Items'].values)[1:])==30
assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
book.close()
assert call('scan')['entries']==scan['entries']
subprocess.run([str(release/'Packing List Manager.exe'),'--self-test',str(folder),str(work/'main.png'),str(work/'history.png'),str(work/'names.png'),'--check-couriers','--select-for-preview','--check-sales'],check=True,timeout=180,creationflags=subprocess.CREATE_NO_WINDOW)
ui=json.loads((release/'SelfTest.json').read_text(encoding='utf-8'))
assert ui['ok'] and ui['selectionTotalsWork'] and ui['courierChangesWork'],ui
assert (folder/'000 Output/2026-09-26_Item_Sales_2.xlsx').exists()
assert not list((folder/'000 Output').glob('*.pdf'))
assert not (folder/'000 Output/.PackingData/Processing_History.json').exists()
report={'ok':True,'version':'1.11.1','work':str(work),'export':result,'ui':ui,'packing_history_unchanged':True}
(APP/'qa/sales-verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
