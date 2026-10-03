"""Sales reports must reconcile, deduplicate and leave packing state unchanged."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
import engine
import sales
from extractors import Batch, Item, Order, PackingError, file_hash
from openpyxl import load_workbook
from test_regeneration import layout


class SalesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.folder = Path(self.temp.name) / '20261003'
        self.folder.mkdir()
        self.catalog = Path(self.temp.name) / 'names.json'
        self.batches = {}
        self.parser = patch.object(engine, 'parse_pdf', side_effect=lambda path, **kwargs: self.batches[Path(path).name])
        self.parser.start()

    def tearDown(self):
        self.parser.stop()
        self.temp.cleanup()

    def source(self, name, platform, orders):
        path = self.folder / name
        path.write_text(name + str(orders), encoding='utf-8')
        self.batches[name] = Batch(str(path), platform, file_hash(path), orders)
        return path

    def export(self):
        return engine.handle({'action': 'sales_export', 'folder': str(self.folder), 'catalog_path': str(self.catalog)})

    def read(self, result, sheet='Item Sales'):
        book = load_workbook(result['path'], data_only=False)
        try:
            return list(book[sheet].values)
        finally:
            book.close()

    def test_platform_dedup_variations_order_counts_and_traceability(self):
        self.source('SPX.pdf', 'SPX', [Order('SPX', '00001234567890123456789', [Item('Soap', 'Lemon', 2), Item('Soap', 'Lemon', 1), Item('Soap', 'Rose', 3)])])
        self.source('FLASH.pdf', 'FLASH', [Order('FLASH', '00001234567890123456789', [Item('Soap', 'Lemon', 2), Item('Soap', 'Lemon', 1), Item('Soap', 'Rose', 3)])])
        self.source('LAZADA.pdf', 'LAZADA', [Order('LAZADA', '00001234567890123456789', [Item('Soap', 'Lemon', 4)])])
        result = self.export()
        self.assertEqual((result['orders'], result['quantity'], result['items']), (2, 10, 2))
        self.assertEqual([row[:8] for row in self.read(result)[1:]], [('2026-10-03', 'SOAP', 'LEMON', 7, 2, 3, 4, 0), ('2026-10-03', 'SOAP', 'ROSE', 3, 1, 3, 0, 0)])
        detail = self.read(result, 'Order Items')
        self.assertEqual(len(detail), 5)
        self.assertEqual(detail[1][2], '00001234567890123456789')
        self.assertTrue(all(len(row[10]) == 64 and len(row[11]) == 64 for row in detail[1:]))
        self.assertEqual(result['filename'], '2026-10-03_Item_Sales_1.xlsx')
        self.assertEqual(engine.scan(self.folder)['entries'][0]['status'], 'READY')
        self.assertFalse((self.folder / '000 Output' / engine.HISTORY).exists())

    def test_name_aliases_independent_fallback_and_literal_excel_text(self):
        self.source('SPX.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2), Item('Alias', 'L', 3), Item('=1+1', '', 1)])])
        catalog = {'revision': 3, 'rows': [{'original_product': 'SOAP', 'original_variation': 'LEMON',
                    'alternative_product': 'CLEANER', 'alternative_variation': 'CITRUS',
                    'aliases': [{'original_product': 'ALIAS', 'original_variation': 'L'}]}]}
        with patch.object(sales, 'load_catalog', return_value=catalog):
            result = self.export()
        rows = [row[:8] for row in self.read(result)]
        self.assertIn(('2026-10-03', 'CLEANER', 'CITRUS', 5, 1, 5, 0, 0), rows)
        book = load_workbook(result['path'])
        try:
            self.assertEqual(book['Item Sales']['B2'].value, '=1+1')
            self.assertEqual(book['Item Sales']['B2'].data_type, 's')
            self.assertEqual(book['Order Items']['E2'].value, 'Soap')
            self.assertEqual(book['Order Items']['F4'].value, None)
        finally:
            book.close()

    def test_reexport_is_snapshot_and_numbering_never_overwrites(self):
        self.source('SPX.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2)])])
        first = self.export()
        digest = file_hash(first['path'])
        second = self.export()
        third = self.export()
        Path(second['path']).unlink()
        self.assertTrue(self.export()['filename'].endswith('_4.xlsx'))
        self.assertEqual(file_hash(first['path']), digest)
        self.assertEqual(self.read(first), self.read(third))

    def test_conflicting_orders_refuse_partial_export(self):
        self.source('one.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2)])])
        self.source('two.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 3)])])
        with self.assertRaisesRegex(PackingError, 'Resolve'):
            self.export()
        self.assertFalse(list((self.folder / '000 Output').glob('*.xlsx')))

    def test_recognized_error_refuses_partial_export(self):
        self.source('one.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2)])])
        with patch.object(engine, 'parse_pdf', side_effect=PackingError('Missing product section')):
            with self.assertRaisesRegex(PackingError, 'Resolve'):
                self.export()

    def test_source_change_during_export_publishes_nothing(self):
        self.source('one.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2)])])
        with patch.object(sales, 'file_hash', return_value='changed'):
            with self.assertRaisesRegex(PackingError, 'changed during export'):
                self.export()
        self.assertFalse(list((self.folder / '000 Output').glob('*.xlsx')))
        self.assertFalse(list((self.folder / '000 Output' / engine.INTERNAL).glob('.sales-*')))

    def test_exact_mixed_missing_and_duplicate_price_values(self):
        self.source('one.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2, '12.50')])])
        result=self.export()
        self.assertEqual(self.read(result)[1][8:], (12.5, 'Source unit price'))
        self.assertEqual(self.read(result, 'Order Items')[1][-1], 12.5)
        self.source('two.pdf', 'SPX', [Order('SPX', 'B', [Item('Soap', 'Lemon', 1, '15.00')])])
        result=self.export()
        self.assertEqual(self.read(result)[1][8:], (None, 'Multiple source prices; see Order Items'))
        self.source('copy.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2, '13.00')])])
        with self.assertRaisesRegex(PackingError, 'Conflicting'):
            self.export()

    def test_empty_date_makes_no_workbook(self):
        with self.assertRaisesRegex(PackingError, 'No readable'):
            self.export()
        self.assertFalse(list((self.folder / '000 Output').glob('*.xlsx')))

    def test_packing_history_unchanged_and_reprint_does_not_add_sales(self):
        self.source('SPX.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 2)])])
        entry = engine.scan(self.folder)['entries'][0]
        selected = [{'path': entry['path'], 'sha256': entry['sha256']}]
        engine.generate(self.folder, selected, verifier=layout)
        warning = engine.generate(self.folder, selected, verifier=layout)
        engine.generate(self.folder, selected, verifier=layout, repeat_confirmation=warning['repeat_token'])
        before = {p: p.read_bytes() for p in (self.folder / '000 Output' / engine.INTERNAL).iterdir() if p.is_file()}
        result = self.export()
        self.assertEqual(result['quantity'], 2)
        self.assertEqual(before, {p: p.read_bytes() for p in before})
        self.assertEqual(engine.scan(self.folder)['entries'][0]['status'], 'DONE')
        self.source('SPX.pdf', 'SPX', [Order('SPX', 'A', [Item('Soap', 'Lemon', 5)])])
        with self.assertRaisesRegex(PackingError, 'Resolve'):
            self.export()


if __name__ == '__main__':
    unittest.main()
