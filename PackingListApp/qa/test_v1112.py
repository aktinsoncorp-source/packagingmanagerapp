"""Buyer comments and definition normalization without changing source identity."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
from docx import Document
from extractors import Batch, Item, Order, ShopeePart, PackingError, combine_shopee_parts, shopee_page
from documents import build_document, COMMENT_LABEL, order_cells
from engine import read_completed_doc
import names


class DefinitionNormalizationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.path = Path(self.temp.name) / 'ProductNames.json'

    def tearDown(self):
        self.temp.cleanup()

    def batch(self, product='AK DETERGENT SCENTS', variation="Oxyherbs',50 0 mL", platform='SPX'):
        return Batch('source.pdf', platform, 'hash', [Order(platform, '123', [Item(product, variation, 2, '45.00', '90.00')])])

    def row(self, variation="OXYHERBS',500 ML", **fields):
        return dict(original_product='AK DETERGENT SCENTS', original_variation=variation,
                    alternative_product='OXYHERBS SCENT', alternative_variation='500 ML', **fields)

    def test_space_insensitive_lookup_reuses_saved_alternatives_and_retains_source(self):
        catalog = names.save_catalog(self.path, [self.row(locked=True)], 0)
        batch = self.batch(' AK  DETER GENT SCENTS ')
        before = copy.deepcopy(batch.to_dict()); fingerprint = batch.orders[0].fingerprint()
        self.assertEqual(names.encounter(self.path, [batch]), catalog)
        item = names.apply_names([batch], catalog)[0].orders[0].items[0]
        self.assertEqual((item.product, item.variation, item.pieces, item.total_price), ('OXYHERBS SCENT', '500 ML', 2, '90.00'))
        self.assertEqual(batch.to_dict(), before)
        self.assertEqual(batch.orders[0].fingerprint(), fingerprint)

    def test_unfinished_spacing_duplicate_keeps_original_as_alias(self):
        unfinished = self.row("OXYHERBS',50 0 ML", locked=False)
        unfinished.update(alternative_product='', alternative_variation='')
        catalog = names.save_catalog(self.path, [unfinished, self.row(locked=True)], 0)
        saved_id = catalog['rows'][1]['id']
        repaired = names.encounter(self.path, [self.batch()])
        self.assertEqual(len(repaired['rows']), 1)
        row = repaired['rows'][0]
        self.assertEqual(row['id'], saved_id); self.assertTrue(row['locked'])
        self.assertEqual({s['original_variation'] for s in names.source_pairs(row)}, {"OXYHERBS',500 ML", "OXYHERBS',50 0 ML"})
        self.assertEqual(names.encounter(self.path, [self.batch()]), repaired)

    def test_locked_or_differently_linked_blank_definition_is_not_removed(self):
        for locked, aliases in [(True, []), (False, [dict(original_product='OTHER', original_variation='OTHER')])]:
            with self.subTest(locked=locked, aliases=aliases):
                self.path.unlink(missing_ok=True)
                blank = self.row("OXYHERBS',50 0 ML", locked=locked, aliases=aliases)
                blank.update(alternative_product='', alternative_variation='')
                catalog = names.save_catalog(self.path, [blank, self.row(locked=True)], 0)
                self.assertEqual(names.encounter(self.path, [self.batch()]), catalog)

    def test_conflicting_defined_spacing_variants_do_not_get_merged_or_guessed(self):
        other = self.row("OXYHERBS',50 0 ML"); other['alternative_product'] = 'DIFFERENT'
        catalog = names.save_catalog(self.path, [self.row(), other], 0)
        exact = names.apply_names([self.batch()], catalog)[0].orders[0].items[0]
        self.assertEqual(exact.product, 'DIFFERENT')
        with self.assertRaisesRegex(PackingError, 'Conflicting saved alternatives'):
            names.encounter(self.path, [self.batch(variation="Oxy herbs',5 00mL")])
        self.assertEqual(names.load_catalog(self.path), catalog)

    def test_lazada_labels_are_removed_from_new_definitions_only_and_source_is_preserved(self):
        for raw, expected in [('Scent: Cool Water Size: 30 mL', 'COOL WATER, 30 ML'),
                              ('Fragrance Oil:Eternity (w), Volume:100 mL', 'ETERNITY (W), 100 ML'),
                              ('Net_Content:1 kg', '1 KG'), ('Packs per Bundle:Single', 'SINGLE'),
                              ('Scent:Cool Water\nSize:30 mL', 'COOL WATER, 30 ML'), ('Cool Water, 30 mL', 'COOL WATER, 30 ML')]:
            with self.subTest(raw=raw):
                batch = self.batch(variation=raw, platform='LAZADA')
                catalog = names.encounter(self.path, [batch])
                self.assertIn(expected, [r['original_variation'] for r in catalog['rows']])
                self.assertEqual(names.apply_names([batch], catalog)[0].orders[0].items[0].variation, expected)
                self.assertEqual(batch.orders[0].items[0].variation, raw)
        self.assertEqual(names.definition_pair('A', 'Scent:Cool Water', 'TIKTOK')[1], 'SCENT:COOL WATER')

    def test_old_lazada_labelled_definition_and_independent_alternatives_still_work(self):
        row = self.row('Scent:Cool Water, Size:30 mL', locked=True); row.update(alternative_product='STANDARD', alternative_variation='')
        catalog = names.save_catalog(self.path, [row], 0)
        batch = self.batch(variation='Scent:Cool Water Size:30mL', platform='LAZADA')
        updated = names.encounter(self.path, [batch])
        self.assertEqual(updated['rows'][0]['id'], catalog['rows'][0]['id'])
        self.assertTrue(updated['rows'][0]['locked'])
        self.assertEqual(updated['rows'][0]['original_variation'], 'COOL WATER, 30 ML')
        self.assertEqual(updated['rows'][0]['aliases'][0]['original_variation'], 'SCENT:COOL WATER, SIZE:30 ML')
        self.assertEqual(names.encounter(self.path, [batch]), updated)
        item = names.apply_names([batch], catalog)[0].orders[0].items[0]
        self.assertEqual((item.product, item.variation), ('STANDARD', 'COOL WATER, 30ML'))

    def test_multiple_numeric_wraps_and_single_line_fields(self):
        for raw, expected in [('Scent,50 0 mL', 'SCENT,500 ML'), ('Scent,2\n50 mL', 'SCENT,250 ML'),
                              ('Scent,1 0 0 g', 'SCENT,100 G'), ('Scent,1 kg', 'SCENT,1 KG')]:
            self.assertEqual(names.definition_pair('A\n B', raw), ('A B', expected))


class BuyerCommentTests(unittest.TestCase):
    def test_comment_row_is_optional_only_for_spx_flash_and_roundtrips_history(self):
        with tempfile.TemporaryDirectory(dir=APP / 'qa') as directory:
            for platform in ('SPX', 'FLASH', 'LAZADA', 'TIKTOK'):
                for comment in ('', 'Send invoice. Please pack carefully.'):
                    with self.subTest(platform=platform, comment=comment):
                        group = 'LAZADA_TIKTOK' if platform in ('LAZADA', 'TIKTOK') else platform
                        order = Order(platform, 'QA123', [Item('A', '30 ML', 2, '45.00', '90.00')], comment)
                        path = Path(directory) / 'output.docx'
                        build_document([Batch('source.pdf', platform, 'hash', [order])], group, '20261005', path)
                        d = Document(path); expected = platform in ('SPX', 'FLASH') and bool(comment)
                        self.assertEqual(len(d.tables[0].rows), 4 if expected else 3)
                        if expected:
                            self.assertEqual([c.text for c in order_cells(d.tables[0].rows[-1])], [COMMENT_LABEL, comment])
                            self.assertTrue(d.tables[0].rows[-2].cells[0].paragraphs[0].paragraph_format.keep_with_next)
                        self.assertNotIn(COMMENT_LABEL, '\n'.join(c.text for r in d.tables[-1].rows for c in r.cells))
                        saved = read_completed_doc(path, group)[order.key]['order']
                        self.assertEqual(saved['buyer_comment'], comment if expected else '')
                        self.assertEqual(saved['items'][0]['pieces'], 2)
                        if expected:
                            # Earlier 1.11.2 documents used one merged comment cell.
                            cells = order_cells(d.tables[0].rows[-1])
                            merged = cells[0].merge(cells[1])
                            merged.text = COMMENT_LABEL + ' ' + comment
                            d.save(path)
                            old = read_completed_doc(path, group)[order.key]['order']
                            self.assertEqual(old['buyer_comment'], comment)
                            self.assertEqual(old['items'], saved['items'])

    def test_continuations_keep_comments_and_reject_conflicts(self):
        def part(first, comment):
            return ShopeePart(Order('SPX', 'QA1', [Item('A', '30 ML', 1)], comment), first, first)
        for comments in [('', 'Send invoice'), ('Send invoice', ''), ('Send invoice', 'Send invoice')]:
            order = combine_shopee_parts([part(1, comments[0]), part(2, comments[1])])[0]
            self.assertEqual(order.buyer_comment, 'Send invoice'); self.assertEqual(order.pieces, 2)
        with self.assertRaisesRegex(PackingError, 'different buyer comments'):
            combine_shopee_parts([part(1, 'One'), part(2, 'Two')])

    def test_comments_remain_in_original_and_printed_snapshots_without_changing_fingerprint(self):
        order = Order('SPX', 'QA1', [Item('A', '30 ML', 1)], 'Send invoice')
        plain = copy.deepcopy(order); plain.buyer_comment = ''
        self.assertEqual(order.fingerprint(), plain.fingerprint())
        batch = Batch('source.pdf', 'SPX', 'hash', [order])
        printed = names.apply_names([batch], {'rows': []})[0]
        self.assertEqual(batch.to_dict()['orders'][0]['buyer_comment'], 'Send invoice')
        self.assertEqual(printed.to_dict()['orders'][0]['buyer_comment'], 'Send invoice')

    def test_source_comment_is_trimmed_without_treating_its_order_words_as_label_metadata(self):
        class Page:
            width, height = 595, 842
            bbox = (0, 0, 297.5, 421)
            def crop(self, bbox):
                return self if bbox[0] == bbox[1] == 0 else Empty()
            def extract_text(self):
                return "Order ID: QA1 package 1\n# Product SKU Variation Qty Unit Total\n1 A V 1 45 45\nOrder ID:QA1 package 1\nBuyer’s Comments:\n Order ID: QA2 package 2\n Send invoice  "
            def extract_words(self, **kwargs):
                entries = [('#', 1, 40), ('Product', 20, 40), ('SKU', 80, 40), ('Variation', 110, 40), ('Qty', 170, 40), ('Unit', 200, 40), ('Total', 240, 40),
                           ('1', 1, 60), ('Order', 1, 90), ('ID:QA1', 30, 90), ('package', 100, 90), ('1', 150, 90), ('Order', 1, 120), ('ID:QA2', 30, 120)]
                return [dict(text=text, x0=x, top=y) for text, x, y in entries]
        class Empty:
            def extract_text(self): return ''
        with patch('extractors.region', side_effect=['A', 'V', '1', '45', '45']):
            order = shopee_page(Page(), 'SPX', 1)[0].order
        self.assertEqual(order.buyer_comment, 'Order ID: QA2 package 2 Send invoice')


if __name__ == '__main__':
    unittest.main()
