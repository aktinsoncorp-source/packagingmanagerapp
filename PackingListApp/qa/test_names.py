"""Definition persistence, Excel editing, and original-vs-printed history checks."""
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import copy
import json
import sys
import unittest

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
from docx import Document
from openpyxl import load_workbook
from extractors import Batch, Item, Order, PackingError, file_hash
import engine
from generation_test_support import generate_approved
import names
from picking import picking_rows


class NameTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(dir=APP / 'qa')
        self.root = Path(self.temp.name)
        self.path = self.root / 'ProductNames.json'
        self.book = self.root / 'Names.xlsx'
        self.batch = Batch('SPX/test.pdf', 'SPX', '', [Order('SPX', 'TEST1', [Item('Listing A', 'Lemon,1 L', 2)]),
                                                    Order('SPX', 'TEST2', [Item('Listing B', 'Lemon,1000mL', 3)])])

    def tearDown(self):
        self.temp.cleanup()

    def seed(self):
        return names.encounter(self.path, [self.batch])

    def defined(self):
        data = self.seed()
        for row in data['rows']:
            row.update(alternative_product='Dish Soap', alternative_variation='Lemon (1 L)')
        return names.save_catalog(self.path, data['rows'], data['revision'])

    def test_remembers_unique_exact_pairs_across_dates_and_platforms(self):
        data = self.seed()
        other = copy.deepcopy(self.batch)
        other.platform = 'FLASH'
        other.orders[0].items.append(Item('Listing A', 'Rose', 1))
        second = names.encounter(self.path, [self.batch, other])
        self.assertEqual(len(second['rows']), 3)
        self.assertEqual(names.load_catalog(self.path), second)
        self.assertEqual(names.encounter(self.path, [self.batch]), second)
        self.assertEqual(data['rows'][0]['id'], second['rows'][0]['id'])

    def test_each_alternative_has_independent_fallback_and_sources_unchanged(self):
        data = self.seed()
        for name, variation in [('', ''), ('Soap', ''), ('', 'Lemon 1L'), ('Soap', 'Lemon 1L')]:
            row = data['rows'][0]
            row.update(alternative_product=name, alternative_variation=variation)
            printed = names.apply_names([self.batch], data)[0].orders[0].items[0]
            self.assertEqual(printed.product, (name or 'Listing A').upper())
            self.assertEqual(printed.variation, (variation or 'Lemon,1 L').upper())
            self.assertEqual(printed.pieces, 2)
        self.assertEqual(self.batch.orders[0].items[0].product, 'Listing A')

    def test_final_names_merge_in_summary_and_quantities_stay_correct(self):
        catalog = self.defined()
        printed = names.apply_names([self.batch], catalog)
        self.assertEqual(picking_rows(printed), [('DISH SOAP', 'LEMON (1 L)', 5)])

    def test_validation_rejects_leading_spaces_symbols_and_controls(self):
        for value in [' Soap', '-Soap', ' Soap ', '=SUM(A1)', 'Soap\nNew', 'Soap\tNew', 'Soap@Home']:
            with self.subTest(value=value), self.assertRaises(PackingError):
                names.validate_alternative(value, 'Name')
        for value in ['', 'Soap - Lemon (1 L) / 50% + Oil', "Men's Soap, 1.5 L & Oil", '洗衣液 1L']:
            self.assertEqual(names.validate_alternative(value, 'Name'), value.upper())

    def test_delete_stays_deleted_until_next_encounter_then_blank(self):
        data = self.defined()
        removed = data['rows'].pop(0)
        names.save_catalog(self.path, data['rows'], data['revision'])
        self.assertEqual(len(names.load_catalog(self.path)['rows']), 1)
        names.export_excel(self.book, names.load_catalog(self.path))
        self.assertEqual(len(names.load_catalog(self.path)['rows']), 1)
        added = names.encounter(self.path, [self.batch])['rows'][0]
        self.assertNotEqual(added['id'], removed['id'])
        self.assertEqual((added['alternative_product'], added['alternative_variation']), ('', ''))

    def test_invalid_save_and_stale_revision_preserve_file(self):
        data = self.seed()
        before = self.path.read_bytes()
        data['rows'][0]['alternative_product'] = ' Bad'
        with self.assertRaises(PackingError):
            names.save_catalog(self.path, data['rows'], data['revision'])
        self.assertEqual(self.path.read_bytes(), before)
        data['rows'][0]['alternative_product'] = 'Good'
        with self.assertRaises(PackingError):
            names.save_catalog(self.path, data['rows'], data['revision'] - 1)
        self.assertEqual(self.path.read_bytes(), before)

    def test_corrupt_catalog_is_not_silently_replaced(self):
        self.path.write_text('broken')
        with self.assertRaises(PackingError):
            self.seed()
        self.assertEqual(self.path.read_text(), 'broken')

    def test_saved_backup_contains_previous_catalog(self):
        old = self.seed()
        self.defined()
        self.assertEqual(json.loads(self.path.with_suffix('.json.bak').read_text()), old)

    def test_excel_roundtrip_edits_originals_alternatives_and_adds(self):
        data = self.defined()
        names.export_excel(self.book, data)
        wb = load_workbook(self.book)
        ws = wb['Names']
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertTrue(ws.column_dimensions['E'].hidden)
        ws['A2'] = 'Renamed original'
        ws['B2'] = 'New original variation'
        ws['C2'] = 'New standard'
        ws['D2'] = None
        ws.append(['Added original', '250mL', 'Added standard', '250 mL'])
        wb.save(self.book)
        wb.close()
        before = self.path.read_bytes()
        preview = names.import_excel(self.book, data['rows'])
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual((preview['added'], preview['updated']), (1, 2))
        saved = names.save_catalog(self.path, preview['rows'], data['revision'])
        updated = next(r for r in saved['rows'] if r['original_product'] == 'RENAMED ORIGINAL')
        self.assertEqual(updated['id'], data['rows'][0]['id'])
        self.assertEqual(updated['alternative_variation'], '')
        self.assertFalse(any(r['original_product'] == 'LISTING A' for r in saved['rows']))

    def test_excel_missing_rows_are_retained(self):
        data = self.seed()
        names.export_excel(self.book, data)
        wb = load_workbook(self.book)
        wb['Names'].delete_rows(3)
        wb.save(self.book)
        wb.close()
        self.assertEqual(len(names.import_excel(self.book, data['rows'])['rows']), 2)

    def test_excel_duplicates_and_formulas_rejected_without_partial_save(self):
        data = self.seed()
        for mode in ['duplicate', 'formula', 'leading space', 'renamed collision']:
            names.export_excel(self.book, data)
            wb = load_workbook(self.book)
            ws = wb['Names']
            if mode == 'duplicate':
                ws.append([ws.cell(2, col).value for col in range(1, 6)])
            elif mode == 'formula':
                ws['C2'] = '=UPPER(A2)'
            elif mode == 'leading space':
                ws['C2'] = ' Invalid'
            else:
                ws['A2'], ws['B2'] = ws['A3'].value, ws['B3'].value
            wb.save(self.book)
            wb.close()
            before = self.path.read_bytes()
            with self.subTest(mode=mode), self.assertRaises(PackingError):
                names.import_excel(self.book, data['rows'])
            self.assertEqual(self.path.read_bytes(), before)

    def test_export_formula_like_original_is_literal_text(self):
        data = self.seed()
        data['rows'][0]['original_product'] = '=literal listing'
        names.export_excel(self.book, data)
        wb = load_workbook(self.book)
        self.assertEqual(wb['Names']['A2'].data_type, 's')
        self.assertEqual(wb['Names']['A2'].value, '=LITERAL LISTING')
        wb.close()

    def test_missing_variation_matches_dash_and_blank(self):
        self.batch.orders[0].items[0].variation = ''
        catalog = self.seed()
        catalog['rows'][0]['alternative_variation'] = 'Standard'
        self.assertEqual(names.apply_names([self.batch], catalog)[0].orders[0].items[0].variation, 'STANDARD')

    def test_generation_preserves_original_fingerprints_and_printed_history(self):
        folder = self.root / '20260926'
        (folder / 'SPX').mkdir(parents=True)
        source = folder / 'SPX/test.pdf'
        source.write_bytes(b'Test source')
        self.batch.path, self.batch.sha256 = str(source), file_hash(source)
        catalog = self.defined()
        with patch.object(engine, 'parse_pdf', return_value=self.batch):
            scan = engine.scan(folder, self.path)
            selection = [{'path': 'SPX/test.pdf', 'sha256': self.batch.sha256}]
            result = generate_approved(folder, selection, verifier=lambda paths: [{'path': str(p), 'pages': 2} for p in paths], catalog_path=self.path)
            doc = Document(result['outputs'][0]['path'])
            self.assertEqual(doc.tables[0].cell(2, 1).text, 'DISH SOAP')
            self.assertEqual(doc.tables[0].cell(2, 2).text, 'LEMON (1 L)')
            self.assertEqual(len(doc.tables[-1].rows), 3)
            self.assertEqual(doc.tables[-1].cell(1, 3).text, '5')
            after = engine.scan(folder, self.path)
            self.assertEqual(after['entries'][0]['status'], 'DONE')
            history = json.loads((folder / '000 Output' / engine.HISTORY).read_text())
            self.assertEqual(history['orders'][self.batch.orders[0].key]['fingerprint'], self.batch.orders[0].fingerprint())
            catalog['rows'][0]['alternative_product'] = 'Later name'
            names.save_catalog(self.path, catalog['rows'], catalog['revision'])
            old = engine.scan(folder, self.path)['outputs'][0]['included_sources'][0]['orders'][0]['items'][0]
            self.assertEqual(old['product'], 'DISH SOAP')

    def test_names_actions_work_without_a_selected_date(self):
        request = {'action': 'names_list', 'catalog_path': str(self.path)}
        self.assertEqual(engine.handle(request)['rows'], [])


if __name__ == '__main__':
    unittest.main(verbosity=2)
