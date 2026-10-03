"""SPX/FLASH packing prices must preserve source values and quantity checks."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal
from docx import Document
from docx.oxml.ns import qn

APP=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(APP/'src'))
from extractors import Batch,Item,Order,parse_pdf
from documents import build_document,item_prices
from engine import read_completed_doc
import names

class PackingPriceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=APP/'qa');self.path=Path(self.temp.name)/'packing.docx'
    def tearDown(self):self.temp.cleanup()
    def build(self,platform,items):
        batch=Batch('source.pdf',platform,'source-hash',[Order(platform,'TEST123',items)])
        build_document([batch],batch.group,'20261004',self.path)
        return Document(self.path)
    def test_decimal_line_total_is_exact(self):
        self.assertEqual(item_prices(Item('A','100ML',3,'0.10')),['0.10','0.30'])
        self.assertEqual(item_prices(Item('A','100ML',7,'1250.50')),['1250.50','8753.50'])
    def test_source_total_takes_precedence_over_computed_total(self):
        self.assertEqual(item_prices(Item('A','100ML',2,'10.00','18.50')),['10.00','18.50'])
    def test_missing_prices_are_not_guessed(self):
        self.assertEqual(item_prices(Item('A','100ML',5)),['—','—'])
        self.assertEqual(item_prices(Item('A','100ML',5,None,'9.00')),['—','9.00'])
    def test_spx_and_flash_have_prices_and_keep_table_width_and_style(self):
        original=Document(APP/'src/assets/orange.docx').tables[0]
        for platform in ('SPX','FLASH'):
            d=self.build(platform,[Item('A','100ML',2,'45.00','90.00'),Item('B','1L',1,'125.50','125.50')])
            table=d.tables[0]
            self.assertEqual([c.text for c in table.rows[1].cells],['Done','Product name','Variation','Pieces','Unit Price','Total Price'])
            self.assertEqual([c.text for c in table.rows[2].cells],['[     ]','A','100ML','2','45.00','90.00'])
            self.assertEqual(table.rows[3].cells[5].text,'125.50')
            self.assertEqual(table.rows[0]._tr.tc_lst[0].tcPr.gridSpan.val,6)
            self.assertEqual(sum(c.w for c in table._tbl.tblGrid.gridCol_lst),sum(c.w for c in original._tbl.tblGrid.gridCol_lst))
            for index,color in ((2,'FFFFFF'),(3,'F2F2F2')):
                self.assertTrue(all(c._tc.tcPr.find(qn('w:shd')).get(qn('w:fill'))==color for c in table.rows[index].cells))
            self.assertEqual(len(d.tables[-1].columns),4)
            self.assertEqual(d.tables[-1].rows[-1].cells[3].text,'3')
    def test_lazada_tiktok_order_layout_stays_four_columns(self):
        for platform in ('LAZADA','TIKTOK'):
            d=self.build(platform,[Item('A','100ML',2,'45.00','90.00')])
            self.assertEqual(len(d.tables[0].columns),4)
            self.assertNotIn('Unit Price',d.tables[0].rows[1].cells[-1].text)
    def test_names_and_snapshots_preserve_prices_and_fingerprints(self):
        item=Item('A','SMALL',2,'10.50','19.00');batch=Batch('a.pdf','SPX','hash',[Order('SPX','TEST',[item])])
        printed=names.apply_names([batch],{'rows':[dict(original_product='A',original_variation='SMALL',alternative_product='STANDARD',alternative_variation='100ML')]})
        display=printed[0].orders[0].items[0]
        self.assertEqual((display.unit_price,display.total_price),('10.50','19.00'))
        self.assertEqual(batch.to_dict()['orders'][0]['items'][0]['total_price'],'19.00')
        self.assertEqual(batch.orders[0].fingerprint(),Order('SPX','TEST',[Item('A','SMALL',2)]).fingerprint())
    def test_history_read_keeps_new_price_columns_and_quantities(self):
        self.build('SPX',[Item('A','100ML',3,'45.00','135.00')])
        item=read_completed_doc(self.path,'SPX')['SHOPEE:TEST123']['order']['items'][0]
        self.assertEqual((item['pieces'],item['unit_price'],item['total_price']),(3,'45.00','135.00'))
    def test_real_shopee_sources_read_both_printed_prices(self):
        root=Path(os.environ.get('PACKING_TEST_DATA',APP/'qa/names-regression-fixtures'))/'20260926'
        for platform in ('SPX','FLASH'):
            batch=parse_pdf(next((root/platform).glob('*.pdf')))
            self.assertTrue(batch.orders)
            for order in batch.orders:
                for item in order.items:
                    self.assertIsNotNone(item.unit_price);self.assertIsNotNone(item.total_price)
                    self.assertEqual(Decimal(item.unit_price)*item.pieces,Decimal(item.total_price))

if __name__=='__main__':unittest.main()
