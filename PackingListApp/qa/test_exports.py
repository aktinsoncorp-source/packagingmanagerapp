"""PDF/Word history, repeat-approval and printed checkbox regression checks."""
from pathlib import Path
import unittest
from unittest.mock import patch
import test_regeneration as fixtures
import engine
from extractors import PackingError
from docx import Document
from generation_test_support import generate_approved
from reportlab.pdfgen import canvas


def layout(paths, export_pdf=False):
    if export_pdf:
        for path in paths:
            c=canvas.Canvas(str(path.with_suffix('.pdf')))
            for _ in range(2):c.drawString(20,800,'Transaction test PDF');c.showPage()
            c.save()
    return [{'path':str(p),'pages':2} for p in paths]


class ExportTests(unittest.TestCase):
    setUp=fixtures.RegenerationTests.setUp
    tearDown=fixtures.RegenerationTests.tearDown
    parse=fixtures.RegenerationTests.parse
    source=fixtures.RegenerationTests.source
    scan=fixtures.RegenerationTests.scan
    history=fixtures.RegenerationTests.history
    def selected(self,names=None):
        return [{'path':e['path'],'sha256':e['sha256']} for e in self.scan()['entries'] if names is None or Path(e['path']).name in names]
    def make(self,format='pdf',names=None):
        return generate_approved(self.folder,self.selected(names),verifier=layout,output_format=format)['outputs'][0]
    def test_pdf_only_export_saved_and_removed_correctly(self):
        self.source('a.pdf',['A','B']);result=self.make();out=self.folder/'000 Output'
        self.assertTrue(result['filename'].endswith('_1.pdf'))
        self.assertFalse(list(out.glob('*.docx')));self.assertFalse(list((out/engine.INTERNAL).glob('.pending-*')))
        entry=self.scan()['entries'][0];self.assertEqual(entry['status'],'DONE')
        self.assertEqual(sum(o['pieces'] for o in entry['order_totals']),4)
        visible=self.scan()['outputs'][0];self.assertEqual(visible['sequence'],1)
        self.assertEqual(len(visible['included_sources'][0]['orders']),2)
        Path(result['path']).unlink();self.assertEqual(self.scan()['entries'][0]['status'],'READY')
        self.assertEqual(self.scan()['outputs'],[])
        self.assertTrue(self.make()['filename'].endswith('_1.pdf'))
    def test_mixed_formats_share_numbers_and_skip_lower_gaps(self):
        self.source('a.pdf',['A']);a=self.make('docx');b=self.make();c=self.make('docx')
        self.assertTrue(a['filename'].endswith('_1.docx'));self.assertTrue(b['filename'].endswith('_2.pdf'));self.assertTrue(c['filename'].endswith('_3.docx'))
        Path(b['path']).unlink();self.assertTrue(self.make()['filename'].endswith('_4.pdf'))
        self.assertEqual([v['sequence'] for v in self.scan()['outputs']],[1,3,4])
    def test_deleted_trailing_pdf_number_is_reused_in_word(self):
        self.source('a.pdf',['A']);self.make('docx');last=self.make();Path(last['path']).unlink()
        self.assertTrue(self.make('docx')['filename'].endswith('_2.docx'))
    def test_previously_generated_order_requires_confirmation(self):
        self.source('a.pdf',['A']);self.make();before=self.history()
        with patch.object(engine,'build_document') as builder:
            answer=engine.generate(self.folder,self.selected(),verifier=layout)
            self.assertTrue(answer['confirmation_required']);self.assertEqual(answer['repeat_count'],1);builder.assert_not_called()
        self.assertEqual(before,self.history())
    def test_overlap_with_unchecked_pdf_also_warns(self):
        self.source('a.pdf',['A','B']);self.source('b.pdf',['B','C'])
        answer=engine.generate(self.folder,self.selected(['b.pdf']),verifier=layout)
        self.assertEqual(answer['repeat_count'],1);self.assertEqual(answer['overlapping_orders'],1)
        self.assertEqual(answer['previously_generated'],0)
        self.assertFalse((self.folder/'000 Output'/engine.HISTORY).exists())
    def test_confirmed_overlap_is_printed_once(self):
        self.source('a.pdf',['A','B']);self.source('b.pdf',['B','C']);result=self.make('docx')
        self.assertEqual(len(Document(result['path']).tables),4)
        self.assertEqual(self.scan()['outputs'][0]['pieces'],6)
    def test_wrong_token_or_different_format_cannot_bypass_warning(self):
        self.source('a.pdf',['A']);self.make();selection=self.selected()
        first=engine.generate(self.folder,selection,verifier=layout)
        wrong=engine.generate(self.folder,selection,verifier=layout,repeat_confirmation='wrong')
        self.assertTrue(wrong['confirmation_required'])
        different=engine.generate(self.folder,selection,verifier=layout,output_format='pdf',repeat_confirmation=first['repeat_token'])
        self.assertTrue(different['confirmation_required'])
        self.assertEqual(len(self.history()['outputs']),1)
    def test_missing_pdf_export_records_nothing(self):
        self.source('a.pdf',['A'])
        with self.assertRaisesRegex(PackingError,'did not produce'):
            engine.generate(self.folder,self.selected(),output_format='pdf',verifier=lambda paths,export_pdf:[{'path':str(p),'pages':2} for p in paths])
        out=self.folder/'000 Output';self.assertFalse((out/engine.HISTORY).exists());self.assertFalse(list((out/engine.INTERNAL).glob('.pending-*')))
    def test_pdf_without_history_is_not_silently_ignored(self):
        self.source('a.pdf',['A']);result=self.make();out=self.folder/'000 Output'
        (out/engine.HISTORY).unlink();(out/engine.LOG).unlink()
        with self.assertRaisesRegex(PackingError,'missing'):self.scan()
    def test_wider_boxes_in_orders_and_summary(self):
        self.source('a.pdf',['A']);doc=Document(self.make('docx')['path'])
        self.assertEqual(doc.tables[0].rows[2].cells[0].text,'[     ]')
        self.assertEqual(doc.tables[-1].rows[1].cells[0].text,'[     ]')
    def test_public_generate_defaults_to_pdf(self):
        with patch.object(engine,'generate',return_value={'ok':True}) as generate:
            engine.handle({'action':'generate','folder':str(self.folder),'selected':[]})
            self.assertEqual(generate.call_args.kwargs['output_format'],'pdf')
    def test_invalid_format_rejected(self):
        self.source('a.pdf',['A'])
        with self.assertRaisesRegex(PackingError,'PDF or Word'):
            engine.generate(self.folder,self.selected(),output_format='png',verifier=layout)

    def test_status_tracks_pdf_word_both_and_deleted_files(self):
        self.source('a.pdf',['A']);pdf=self.make()
        self.assertEqual(self.scan()['entries'][0]['message'],'Process again')
        word=self.make('docx');self.assertEqual(self.scan()['entries'][0]['message'],'Process again')
        Path(pdf['path']).unlink();self.assertEqual(self.scan()['entries'][0]['message'],'Process again')
        Path(word['path']).unlink();self.assertEqual(self.scan()['entries'][0]['status'],'READY')
        self.assertEqual(self.scan()['entries'][0]['output_formats'],[])
    def test_old_word_record_without_format_field_is_recognized(self):
        self.source('a.pdf',['A']);word=self.make('docx');history=self.history()
        history['outputs'][word['filename']].pop('format')
        engine.atomic_json(self.folder/'000 Output'/engine.HISTORY,history)
        self.assertEqual(self.scan()['entries'][0]['message'],'Process again')
    def test_superseded_pdf_does_not_count_as_current_word_format(self):
        self.source('a.pdf',['A']);pdf=self.make();self.source('a.pdf',['A'],quantity=3)
        selected=self.selected();selected[0]['revision']=True
        generate_approved(self.folder,selected,verifier=layout,output_format='docx')
        self.assertTrue(Path(pdf['path']).exists())
        self.assertEqual(self.scan()['entries'][0]['message'],'Process again')
    def test_partial_format_coverage_is_labelled_as_some_orders(self):
        self.source('a.pdf',['A']);self.make('docx');self.source('b.pdf',['A','B'])
        entry=next(e for e in self.scan()['entries'] if e['path'].endswith('b.pdf'))
        self.assertEqual(entry['status'],'READY');self.assertIn('Some orders already processed',entry['message'])

    def test_public_generate_rejects_word_export(self):
        with patch.object(engine,'generate') as generate:
            with self.assertRaisesRegex(PackingError,'PDF only'):
                engine.handle({'action':'generate','folder':str(self.folder),'selected':[],'output_format':'docx'})
            generate.assert_not_called()


class SkipRepeatTests(unittest.TestCase):
    setUp=ExportTests.setUp
    tearDown=ExportTests.tearDown
    parse=ExportTests.parse
    source=ExportTests.source
    scan=ExportTests.scan
    history=ExportTests.history
    selected=ExportTests.selected
    make=ExportTests.make
    def skip(self, selection=None):
        selection=self.selected() if selection is None else selection
        warning=engine.generate(self.folder,selection,verifier=layout,repeat_action='skip')
        if warning.get('confirmation_required'):
            return engine.generate(self.folder,selection,verifier=layout,repeat_action='skip',repeat_confirmation=warning['repeat_token'])
        return warning
    def test_no_generates_new_orders_only_and_retains_earlier_output(self):
        self.source('old.pdf',['A']);old=self.make('docx')
        self.source('mixed.pdf',['A','B'])
        result=self.skip(self.selected(['mixed.pdf']))
        self.assertEqual(result['skipped_repeat_orders'],1)
        self.assertEqual(len(result['outputs']),1)
        saved=self.history()['outputs'][result['outputs'][0]['filename']]
        self.assertEqual(saved['orders'],['SHOPEE:B']);self.assertEqual(saved['sources'][0]['order_keys'],['SHOPEE:B'])
        self.assertEqual(saved['order_snapshots']['SHOPEE:B']['items'][0]['pieces'],2)
        self.assertEqual(self.history()['orders']['SHOPEE:A']['output'],old['filename'])
        self.assertTrue(Path(old['path']).exists())
        self.assertTrue(all(e['message']=='Process again' for e in self.scan()['entries']))
    def test_no_with_all_processed_produces_no_empty_document_or_history_change(self):
        self.source('a.pdf',['A']);self.make();before=self.history()
        with patch.object(engine,'build_document') as build:
            result=self.skip();build.assert_not_called()
        self.assertEqual(result['outputs'],[]);self.assertIn('No new orders',result['message'])
        self.assertEqual(before,self.history())
    def test_no_recovers_deleted_outputs(self):
        self.source('a.pdf',['A']);old=self.make();Path(old['path']).unlink()
        result=self.skip();self.assertEqual(len(result['outputs']),1)
        self.assertEqual(result['skipped_repeat_orders'],0)
        self.assertTrue(result['outputs'][0]['filename'].endswith('_1.docx'))
    def test_no_retains_new_order_shared_by_two_unprocessed_pdfs_once(self):
        self.source('a.pdf',['A','B']);self.source('b.pdf',['B','C'])
        result=self.skip();saved=self.history()['outputs'][result['outputs'][0]['filename']]
        self.assertEqual(set(saved['orders']),{'SHOPEE:A','SHOPEE:B','SHOPEE:C'})
        self.assertEqual(len(saved['orders']),3);self.assertEqual(result['skipped_repeat_orders'],0)
    def test_no_does_not_bypass_confirmation_token(self):
        self.source('a.pdf',['A']);self.make()
        answer=engine.generate(self.folder,self.selected(),verifier=layout,repeat_action='skip',repeat_confirmation='wrong')
        self.assertTrue(answer['confirmation_required'])
    def test_no_revision_preserves_unchanged_orders_and_real_removals(self):
        self.source('a.pdf',['A','B','C']);old=self.make('docx')
        self.source('a.pdf',['A','B'],quantity=3);self.items['a.pdf'][0].items[0].pieces=2
        selection=self.selected();selection[0]['revision']=True
        result=self.skip(selection);saved=self.history()['outputs'][result['outputs'][0]['filename']]
        self.assertEqual(saved['orders'],['SHOPEE:B']);self.assertEqual(saved['removed_orders'],['SHOPEE:C'])
        self.assertNotIn('removed_from_source_in',self.history()['orders']['SHOPEE:A'])
        self.assertEqual(self.history()['orders']['SHOPEE:A']['output'],old['filename'])
        self.assertEqual(set(self.history()['outputs'][old['filename']]['superseded_orders']),{'SHOPEE:B','SHOPEE:C'})
    def test_public_generate_forwards_no_decision(self):
        with patch.object(engine,'generate',return_value={'ok':True}) as generate:
            engine.handle({'action':'generate','folder':str(self.folder),'selected':[],'repeat_action':'skip'})
            self.assertEqual(generate.call_args.kwargs['repeat_action'],'skip')
    def test_alternating_white_gray_rows_keep_themed_headers(self):
        from docx.oxml.ns import qn
        from extractors import Item
        self.source('a.pdf',['A']);self.items['a.pdf'][0].items += [Item('Soap two','Lime',3),Item('Soap three','Mint',1)]
        doc=Document(self.make('docx')['path'])
        def fill(row):return [c._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill')) for c in row.cells]
        for table,first in [(doc.tables[0],2),(doc.tables[-1],1)]:
            self.assertEqual([fill(table.rows[first+i]) for i in range(3)],[['FFFFFF']*4,['F2F2F2']*4,['FFFFFF']*4])
            self.assertNotIn(fill(table.rows[first-1])[0],['FFFFFF','F2F2F2'])
