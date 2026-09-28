"""Packaged PDF/Word export and employee interface verification on isolated data."""
from pathlib import Path
import json, os, shutil, subprocess, sys, tempfile
APP=Path(__file__).resolve().parents[1];sys.path.insert(0,str(APP/'src'))
from docx import Document
import pdfplumber
from extractors import parse_pdf
release=Path(sys.argv[1]).resolve();fixture=APP/'qa/names-regression-fixtures/20260926'
work=Path(tempfile.mkdtemp(prefix='release18-',dir=APP/'qa'));folder=work/'20260926'
for relative in ['SPX/20260926_SPX_PL_4.pdf','FLASH/20260926_Flash_PL_3.pdf','LAZADA & TIKTOK/20260926_Lazada_PL_1.pdf','LAZADA & TIKTOK/20260926_Tiktok_PL_1.pdf']:
    target=folder/relative;target.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(fixture/relative,target)
shutil.copyfile(fixture/'SPX/20260926_SPX_PL_4.pdf',folder/'packing.pdf')
def call(action,**kwargs):
    request=work/'request.json';response=work/'response.json'
    request.write_text(json.dumps({'action':action,'folder':str(folder),'catalog_path':str(work/'ProductNames.json'),**kwargs}),encoding='utf-8')
    response.unlink(missing_ok=True)
    subprocess.run([str(release/'App Files/PackingEngine.exe'),'--request',str(request),'--response',str(response)],check=True,timeout=240,creationflags=subprocess.CREATE_NO_WINDOW)
    data=json.loads(response.read_text(encoding='utf-8'));assert data['ok'],data;return data
entries=call('scan')['entries'];selection=[{'path':e['path'],'sha256':e['sha256'],'platform':e['platform'],'courier':e['courier']} for e in entries if e['platform']!='SHOPEE']
assert len(selection)==4
first=call('generate',selected=selection)
assert first['confirmation_required'] and first['repeat_count']==7 and first['previously_generated']==0
assert not list((folder/'000 Output').glob('*_Packing_Checklist_*'))
pdfs=call('generate',selected=selection,repeat_confirmation=first['repeat_token'])['outputs']
assert len(pdfs)==3 and all(o['filename'].endswith('_1.pdf') for o in pdfs)
assert not list((folder/'000 Output').glob('*.docx'))
boxes=0
for output in pdfs:
    with pdfplumber.open(output['path']) as pdf:
        ids=[]
        for page in pdf.pages:
            assert abs(page.width-595.3)<1 and abs(page.height-841.9)<1
            for char in page.chars:
                if char['text']=='[' and char['x0']<55:
                    closing=[c for c in page.chars if c['text']==']' and c['x0']<62 and abs(c['top']-char['top'])<.8 and c['x0']>char['x1']]
                    assert closing and closing[0]['x0']-char['x1']>10,(output,char,closing)
                    boxes+=1
assert boxes>25
scan=call('scan');assert all(e['status']=='DONE' for e in scan['entries'] if e['platform']!='SHOPEE')
assert sum(o['order_count'] for o in scan['outputs'])==17 and sum(o['pieces'] for o in scan['outputs'])==30
repeat=call('generate',selected=selection,output_format='docx')
assert repeat['confirmation_required'] and repeat['previously_generated']==17
words=call('generate',selected=selection,output_format='docx',repeat_confirmation=repeat['repeat_token'])['outputs']
assert all(o['filename'].endswith('_2.docx') for o in words)
for output in words:
    doc=Document(output['path'])
    assert all(r.cells[0].text=='[     ]' for t in doc.tables[:-1] for r in t.rows[2:])
    assert all(r.cells[0].text=='[     ]' for r in doc.tables[-1].rows[1:-1])
subprocess.run([str(release/'Packing List Manager.exe'),'--self-test',str(folder),str(work/'main.png'),str(work/'history.png'),str(work/'names-locked.png'),'--check-couriers','--select-for-preview'],check=True,timeout=180,creationflags=subprocess.CREATE_NO_WINDOW)
ui=json.loads((release/'SelfTest.json').read_text(encoding='utf-8'));assert ui['ok'] and ui['selectionTotalsWork'],ui
report={'ok':True,'version':'1.8.0','work':str(work),'pdfs':pdfs,'words':words,'orders':17,'pieces':30,'wide_boxes_checked':boxes,'ui':ui,'repeat_warning_before_first_export':True,'regeneration_warning':True}
(APP/'qa/release-1.8-verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report,indent=2))
