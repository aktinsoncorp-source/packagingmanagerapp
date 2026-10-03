"""Packaged import of a pasted inventory snapshot, saving only an isolated catalog."""
import argparse
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
from openpyxl import load_workbook

APP=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(APP/'src'))
import names

def main():
    parser=argparse.ArgumentParser()
    for field in ('engine','workbook','catalog','output'):parser.add_argument(field,type=Path)
    args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=False)
    snapshot=args.catalog.read_bytes();catalog=json.loads(snapshot.decode('utf-8-sig'))
    clone=args.output/'Catalog.json';shutil.copy2(args.catalog,clone)
    expected={}
    wb=load_workbook(args.workbook,read_only=True,data_only=False)
    try:
        for cells in wb['Names'].iter_rows(min_row=2,max_col=4,values_only=True):
            if not any(cells):continue
            row=names.validate_rows([dict(zip(names.FIELDS,['' if c is None else str(c) for c in cells]))])[0]
            key=names.item_key(row['original_product'],row['original_variation'])
            assert key not in expected,'Fixture must contain distinct visible originals'
            expected[key]=tuple(row[f] for f in names.FIELDS[2:])
    finally:wb.close()
    def call(action,**fields):
        request=dict(action=action,catalog_path=str(clone),**fields)
        input_path=args.output/(action+'-request.json');output=args.output/(action+'-response.json')
        input_path.write_text(json.dumps(request,ensure_ascii=False),encoding='utf-8')
        subprocess.run([str(args.engine.resolve()),'--request',str(input_path.resolve()),'--response',str(output.resolve())],check=True,timeout=60)
        result=json.loads(output.read_text(encoding='utf-8'));assert result['ok'],result;return result
    imported=call('names_import',path=str(args.workbook.resolve()),rows=catalog['rows'])
    actual={names.item_key(s['original_product'],s['original_variation']):tuple(r[f] for f in names.FIELDS[2:]) for r in imported['rows'] for s in names.source_pairs(r)}
    assert actual==expected,'Import must keep all pasted originals and their exact alternative pairs'
    before_ids={r['id']:names.item_key(r['original_product'],r['original_variation']) for r in catalog['rows']}
    after_ids={r['id']:names.item_key(r['original_product'],r['original_variation']) for r in imported['rows']}
    assert all(after_ids.get(i)==k for i,k in before_ids.items()),'Every original base ID must stay with its base'
    prepared=call('names_prepare',rows=imported['rows'],priority=imported['priority'])
    saved=call('names_save',rows=prepared['rows'],revision=catalog['revision'],merge_groups=[])
    persisted={names.item_key(s['original_product'],s['original_variation']):tuple(r[f] for f in names.FIELDS[2:]) for r in saved['rows'] for s in names.source_pairs(r)}
    assert persisted==expected;assert args.catalog.read_bytes()==snapshot,'Source catalog must remain untouched'
    report=dict(ok=True,originals=len(expected),definitions=len(saved['rows']),matching_groups=len(prepared['merges']),source_catalog_unchanged=True,source_catalog_sha256=hashlib.sha256(snapshot).hexdigest(),**{k:imported[k] for k in ('added','updated','separated','repaired')})
    (args.output/'result.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report,indent=2))

if __name__=='__main__':main()
