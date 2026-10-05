from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import json
import os
import subprocess
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.shared import Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

from extractors import Batch, PackingError
from picking import picking_rows


ASSETS = Path(__file__).resolve().parent / 'assets'
COMMENT_LABEL = "BUYER'S COMMENT:"


def item_prices(item):
    unit = format(Decimal(item.unit_price), '.2f') if item.unit_price is not None else '—'
    total = item.total_price
    if total is None and item.unit_price is not None:
        total = Decimal(item.unit_price) * item.pieces
    return [unit, format(Decimal(total), '.2f') if total is not None else '—']


def add_price_columns(table):
    """Keep the A4 table width, shared order header, borders and row styles."""
    total = sum(int(column.w) // 635 for column in table._tbl.tblGrid.gridCol_lst)
    widths = [round(total * fraction) for fraction in (.07, .42, .23, .07, .105)]
    widths.append(total - sum(widths))
    for row in table.rows[1:]:
        for _ in range(2):
            row._tr.append(deepcopy(row._tr.tc_lst[-1]))
    grid = table._tbl.tblGrid
    for column in list(grid):
        grid.remove(column)
    for width in widths:
        column = OxmlElement('w:gridCol')
        column.set(qn('w:w'), str(width))
        grid.append(column)
    header = table.rows[0]._tr.tc_lst[0]
    header.get_or_add_tcPr().find(qn('w:gridSpan')).set(qn('w:val'), '6')
    table.cell(0, 0).width = Twips(total)
    for row in table.rows[1:]:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
    for cell, title in zip(table.rows[1].cells[4:], ('Unit Price', 'Total Price')):
        set_text(cell.paragraphs[0], title)
    for cell in table.rows[2].cells[4:]:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def set_text(paragraph, text):
    """Preserve the component's paragraph and run formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
    else:
        target = next((r for r in runs if r.text), runs[-1])
        for r in runs:
            r.text = text if r is target else ''


def set_order_header(paragraph, platform, order_id):
    """Keep all handwriting fields in the existing merged order-header cell."""
    paragraph.clear()
    tabs = paragraph.paragraph_format.tab_stops
    tabs.clear_all()
    for position in (185, 340, 440):
        tabs.add_tab_stop(Pt(position))
    label = 'Order Number' if platform == 'LAZADA' else 'Order ID'
    fields = [f'{label}: {order_id}', '\tEmployee:__________________',
              '\tNumber __________', '\tTime __________']
    for index, text in enumerate(fields):
        run = paragraph.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(9.5)
        run.bold = index == 0


def build_document(batches: list[Batch], group: str, date: str, destination: Path, revision=False):
    d = Document(ASSETS / ('blue.docx' if group == 'LAZADA_TIKTOK' else 'orange.docx'))
    body = d._element.body
    title, intro, heading, spacer = [deepcopy(p._p) for p in d.paragraphs]
    table_proto = deepcopy(d.tables[0]._tbl)
    prices = group in ('SPX', 'FLASH')
    if prices:
        add_price_columns(Table(table_proto, d._body))
    for node in list(body):
        if node.tag != qn('w:sectPr'):
            body.remove(node)

    def append(node):
        body.insert(len(body)-1, node)

    date_obj = datetime.strptime(date, '%Y%m%d')
    display_group = 'Lazada and TikTok' if group == 'LAZADA_TIKTOK' else group
    p = Paragraph(title, d._body)
    set_text(p, f'{display_group} Packing List {date_obj.day} {date_obj:%B %Y}')
    append(title)
    p = Paragraph(intro, d._body)
    intro_text = "One table is one order. Write the employee name, waybill number, and packing completion time in its header."
    if group == 'LAZADA_TIKTOK':
        intro_text += '  — = no variation provided.'
    if revision:
        intro_text += '  REVISION: check the previous packing record before packing these orders.'
    set_text(p, intro_text)
    append(intro)
    total_heading = deepcopy(heading)
    total_paragraph = Paragraph(total_heading, d._body)
    set_text(total_paragraph, f"Total orders: {sum(len(b.orders) for b in batches)}")
    total_paragraph.paragraph_format.keep_with_next = True
    append(total_heading)
    for batch_no, batch in enumerate(batches, 1):
        h = deepcopy(heading)
        p = Paragraph(h, d._body)
        label = batch.platform if group == 'LAZADA_TIKTOK' else f'Batch {batch_no}'
        p.runs[0].text = label
        if len(p.runs) > 1:
            p.runs[1].text = '   ' + Path(batch.path).name
        else:
            p.add_run('   ' + Path(batch.path).name)
        p.paragraph_format.keep_with_next = True
        append(h)
        for order in batch.orders:
            tbl = deepcopy(table_proto)
            table = Table(tbl, d._body)
            header = table.cell(0, 0).paragraphs[0]
            set_order_header(header, order.platform, order.order_id)
            row_proto = deepcopy(table.rows[2]._tr)
            tbl.remove(table.rows[2]._tr)
            for item_no, item in enumerate(order.items):
                row = deepcopy(row_proto)
                tbl.append(row)
                cells = table.rows[-1].cells
                values = ['[     ]', item.product, item.variation, str(item.pieces)]
                if prices:
                    values += item_prices(item)
                for cell, text in zip(cells, values):
                    set_text(cell.paragraphs[0], text)
                    fill = 'FFFFFF' if item_no % 2 == 0 else 'F2F2F2'
                    cell._tc.get_or_add_tcPr().find(qn('w:shd')).set(qn('w:fill'), fill)
            if prices and order.buyer_comment:
                tbl.append(deepcopy(row_proto))
                for cell in table.rows[-1].cells:
                    set_text(cell.paragraphs[0], '')
                cell = table.rows[-1].cells[0].merge(table.rows[-1].cells[-1])
                for paragraph in cell.paragraphs[1:]:
                    cell._tc.remove(paragraph._p)
                set_text(cell.paragraphs[0], COMMENT_LABEL + ' ' + order.buyer_comment)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                cell._tc.get_or_add_tcPr().find(qn('w:shd')).set(qn('w:fill'), 'FFFFFF')
            for row_no, row in enumerate(table.rows):
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_together = True
                        p.paragraph_format.keep_with_next = row_no < len(table.rows)-1
            append(tbl)
            append(deepcopy(spacer))
    summary_rows = picking_rows(batches)
    summary = Document(ASSETS / ('summary_blue.docx' if group == 'LAZADA_TIKTOK' else 'summary_orange.docx'))
    summary_texts = [f'{display_group} Product Picking Summary',
                     f'{date_obj.day} {date_obj:%B %Y}  |  {sum(len(b.orders) for b in batches)} orders  |  {sum(b.pieces for b in batches)} pieces',
                     summary.paragraphs[2].text]
    for proto, text in zip(summary.paragraphs, summary_texts):
        node = deepcopy(proto._p)
        set_text(Paragraph(node, d._body), text)
        append(node)
    tbl = deepcopy(summary.tables[0]._tbl)
    table = Table(tbl, d._body)
    row_proto, total_proto = deepcopy(table.rows[1]._tr), deepcopy(table.rows[-1]._tr)
    for row in list(table.rows)[1:]:
        tbl.remove(row._tr)
    for i, (name, size, pieces) in enumerate(summary_rows):
        tbl.append(deepcopy(row_proto))
        row = table.rows[-1]
        if row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is None:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, text in zip(row.cells, ['[     ]', name, size, str(pieces)]):
            set_text(cell.paragraphs[0], text)
            cell.paragraphs[0].paragraph_format.keep_with_next = i == len(summary_rows)-1
            cell._tc.get_or_add_tcPr().find(qn('w:shd')).set(qn('w:fill'),
                'FFFFFF' if i % 2 == 0 else 'F2F2F2')
    tbl.append(total_proto)
    set_text(table.rows[-1].cells[3].paragraphs[0], str(sum(r[2] for r in summary_rows)))
    append(tbl)
    footer_group = 'LAZADA + TIKTOK' if group == 'LAZADA_TIKTOK' else group
    for p in d.sections[0].footer.paragraphs:
        for run in p.runs:
            if 'Page' in run.text:
                run.text = f'{footer_group} {date_obj:%Y/%m/%d}  |  Page '
    d.core_properties.title = f'{display_group} Packing List {date_obj:%Y-%m-%d}'
    d.core_properties.author = 'Packing List Manager'
    d.core_properties.last_modified_by = 'Packing List Manager'
    d.core_properties.created = datetime.now()
    d.core_properties.modified = datetime.now()
    d.save(destination)
    actual = Document(destination)
    if len(actual.tables) != sum(len(b.orders) for b in batches) + 1:
        raise PackingError('Document verification failed: order count mismatch.')
    for table, order in zip(actual.tables, (o for b in batches for o in b.orders)):
        expected = [[i.product, i.variation, str(i.pieces)] + (item_prices(i) if prices else []) for i in order.items]
        has_comment = prices and bool(order.buyer_comment)
        item_rows = table.rows[2:-1] if has_comment else table.rows[2:]
        rows = [[c.text for c in row.cells[1:]] for row in item_rows]
        if has_comment and table.rows[-1].cells[0].text != COMMENT_LABEL + ' ' + order.buyer_comment:
            raise PackingError(f'Document verification failed for buyer comment on order {order.order_id}.')
        header = table.cell(0, 0).text
        expected_header = f"{'Order Number' if order.platform == 'LAZADA' else 'Order ID'}: {order.order_id}\tEmployee:__________________\tNumber __________\tTime __________"
        if rows != expected or header != expected_header:
            raise PackingError(f'Document verification failed for order {order.order_id}.')
    got = [(r.cells[1].text, r.cells[2].text, int(r.cells[3].text)) for r in actual.tables[-1].rows[1:-1]]
    total = sum(b.pieces for b in batches)
    if got != summary_rows or sum(r[2] for r in got) != total or int(actual.tables[-1].rows[-1].cells[3].text) != total:
        raise PackingError('Document verification failed: product picking totals mismatch.')


def verify_word(paths: list[Path], export_pdf=False):
    """Use the workstation's Word to verify real pagination before recording completion."""
    if os.name != 'nt':
        raise PackingError('Microsoft Word on Windows is required to verify the packing document layout.')
    with tempfile.TemporaryDirectory(prefix='PackingWord-') as work:
        request = Path(work) / 'request.json'
        result = Path(work) / 'result.json'
        request.write_text(json.dumps({'paths': [str(p.resolve()) for p in paths],
                                      'result': str(result), 'export_pdf': export_pdf}), encoding='utf-8')
        powershell = Path(os.environ.get('SystemRoot', r'C:\Windows')) / 'System32/WindowsPowerShell/v1.0/powershell.exe'
        try:
            proc = subprocess.run([str(powershell), '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                                   '-File', str(ASSETS / 'verify_word.ps1'), '-RequestFile', str(request)],
                                  capture_output=True, timeout=180, creationflags=subprocess.CREATE_NO_WINDOW)
        except subprocess.TimeoutExpired as exc:
            raise PackingError('Word did not finish the layout check. Close any Word dialog and try again. No completed batch was recorded.') from exc
        if not result.exists():
            error = proc.stderr.decode('utf-8', errors='replace').strip()
            raise PackingError('Microsoft Word could not check the layout. Open Word once, finish any setup prompts, then try again. ' + error[:300])
        data = json.loads(result.read_text(encoding='utf-8-sig'))
        if not data.get('ok'):
            raise PackingError('Word layout check: ' + data.get('error', 'unknown error'))
        return data['documents']
