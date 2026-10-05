"""Persistent name definitions, linked original pairs, and Excel round trips."""
from contextlib import contextmanager
from dataclasses import replace
from pathlib import Path
import copy
import json
import os
import re
import shutil
import uuid

from extractors import Item, PackingError, clean

FIELDS = ('original_product', 'original_variation', 'alternative_product', 'alternative_variation')
HEADERS = ('Original product name', 'Original variation', 'Alternative product name', 'Alternative variation', 'Entry ID', 'Original pair')
ALLOWED_SYMBOLS = " -()/%.+,&'"


def single_line(value, label):
    if not isinstance(value, str) or any(ord(c) < 32 or c in '\x7f\x85\u2028\u2029' for c in value):
        raise PackingError(f'{label} must be a single line of text; line breaks and paragraphs are not allowed.')
    return value.rstrip(' ').upper()


def item_key(product, variation):
    return product.rstrip(' ').upper(), variation.rstrip(' ').upper() or '—'


def definition_pair(product, variation, platform=''):
    """Readable definition fields; raw PDF items and fingerprints stay unchanged."""
    if platform == 'LAZADA':
        # Attribute labels delimit values; values may themselves contain spaces.
        labels = r'(?:^|[,;\n])\s*[^:,;\n]+:\s*|\b(?:Scent|Scents|Size|Volume|Color|Colour|Net_Content|Fragrance Oil|Packs per Bundle|Set)\s*:\s*'
        variation = re.sub(labels, ', ', variation, flags=re.I).strip(' ,;\n')
        variation = re.sub(r'\s*,\s*', ', ', variation)
    # A narrow Shopee column can wrap anywhere within a printed size's digits.
    variation = re.sub(r'\d(?:[\d\s]*\d)?(?=\s*(?:ML|L|G|KG|LITER|LITRE)\b)',
                       lambda m: re.sub(r'\s+', '', m[0]), variation, flags=re.I)
    return item_key(clean(product), clean(variation))


def lookup_key(product, variation, platform=''):
    return tuple(re.sub(r'\s+', '', value) for value in definition_pair(product, variation, platform))


def definition_mapping(rows, platform=''):
    mapping = {}
    for row in rows:
        for source in source_pairs(row):
            key = lookup_key(source['original_product'], source['original_variation'], platform)
            matches = mapping.setdefault(key, [])
            if not any(existing is row for existing in matches):
                matches.append(row)
    return mapping


def matched_definition(matches, product, variation):
    defined = [r for r in matches if r['alternative_product'] or r['alternative_variation']]
    if len({(r['alternative_product'], r['alternative_variation']) for r in defined}) > 1:
        exact = [r for r in defined if any(item_key(s['original_product'], s['original_variation']) == item_key(product, variation) for s in source_pairs(r))]
        if not exact:
            raise PackingError(f'Conflicting saved alternatives for spacing variants of {product} / {variation}. Review these definitions before generating.')
        return exact[0]
    return next(iter(defined or matches), {})


def validate_alternative(value, label):
    value = single_line(value, label)
    if not value:
        return value
    value = value[next((i for i, c in enumerate(value) if c.isalnum()), len(value)):]
    if not value:
        return ''
    if any(not (c.isalnum() or c in ALLOWED_SYMBOLS) for c in value):
        raise PackingError(f'{label} may contain letters, numbers, spaces, and these symbols: - ( ) / % . + , & apostrophe.')
    return value


def source_pairs(row):
    return [{f: row[f] for f in FIELDS[:2]}] + row.get('aliases', [])


def _source(value, label):
    if not isinstance(value, dict):
        raise PackingError(f'{label}: invalid original pair.')
    result = {f: single_line(value.get(f, ''), f'{label}: {f.replace("_", " ")}') for f in FIELDS[:2]}
    if not result['original_product'].strip():
        raise PackingError(f'{label}: Original product name is required.')
    result['original_variation'] = result['original_variation'] or '—'
    return result


def validate_rows(rows, allow_shared=False):
    if not isinstance(rows, list):
        raise PackingError('The names list is invalid.')
    result, keys, ids = [], {}, set()
    for number, row in enumerate(rows, 1):
        label = f'Row {number}'
        entry = _source(row, label)
        for field in FIELDS[2:]:
            entry[field] = validate_alternative(row.get(field, ''), f'{label}: {field.replace("_", " ")}')
        if not isinstance(row.get('locked', False), bool):
            raise PackingError(f'{label}: invalid row lock.')
        entry['locked'] = row.get('locked', False)
        entry['id'] = row.get('id') or uuid.uuid4().hex
        if not isinstance(entry['id'], str) or entry['id'] in ids:
            raise PackingError(f'{label}: duplicate or invalid entry identifier. Start with a fresh export.')
        aliases = row.get('aliases', [])
        if not isinstance(aliases, list):
            raise PackingError(f'{label}: invalid linked originals.')
        entry['aliases'] = []
        local = set()
        for index, source in enumerate([entry] + aliases):
            source = _source(source, label)
            key = item_key(source['original_product'], source['original_variation'])
            if key in local:
                continue  # The first original pair is the base, regardless of case.
            alternatives = tuple(entry[f] for f in FIELDS[2:])
            if key in keys and not (allow_shared and all(alternatives) and keys[key] == alternatives):
                raise PackingError(f'{label}: this original product and variation already belong to another row (ignoring case). Merge matching definitions or use distinct original pairs.')
            local.add(key)
            keys[key] = alternatives
            if index:
                entry['aliases'].append(source)
        ids.add(entry['id'])
        result.append(entry)
    return sorted(result, key=lambda r: (r['original_product'], r['original_variation']))


def prepare_save(rows, priority=()):
    # Assign IDs before sorting so unsorted/Excel order remains the fallback base.
    pending = copy.deepcopy(rows)
    for row in pending:
        if isinstance(row, dict) and not row.get('id'):
            row['id'] = uuid.uuid4().hex
    order = list(dict.fromkeys(list(priority) + [r.get('id') for r in pending if isinstance(r, dict)]))
    validated = validate_rows(pending, allow_shared=True)
    rank = {identifier: index for index, identifier in enumerate(order)}
    groups = {}
    for row in sorted(validated, key=lambda r: rank.get(r['id'], len(rank))):
        if not row.get('locked') and row['alternative_product'] and row['alternative_variation']:
            groups.setdefault((row['alternative_product'], row['alternative_variation']), []).append(row)
    matches = [{'ids': [r['id'] for r in group], 'alternative_product': key[0], 'alternative_variation': key[1],
                'base_product': group[0]['original_product'], 'base_variation': group[0]['original_variation'],
                'originals': [source for row in group for source in source_pairs(row)]}
               for key, group in groups.items() if len(group) > 1]
    return {'ok': True, 'rows': validated, 'merges': matches}


def merge_rows(rows, groups):
    rows = validate_rows(rows, allow_shared=True)
    by_id = {r['id']: r for r in rows}
    consumed = set()
    for identifiers in groups:
        if not isinstance(identifiers, list) or len(identifiers) < 2 or len(set(identifiers)) != len(identifiers):
            raise PackingError('Invalid merge selection.')
        if any(i not in by_id or i in consumed for i in identifiers):
            raise PackingError('The merge selection changed. Review the names again.')
        if any(by_id[i].get('locked') for i in identifiers):
            raise PackingError('Unlock each definition before merging its linked originals.')
        base = by_id[identifiers[0]]
        key = tuple(base[f] for f in FIELDS[2:])
        if not all(key) or any(tuple(by_id[i][f] for f in FIELDS[2:]) != key for i in identifiers):
            raise PackingError('Merge only rows with matching, nonblank alternative product names AND variations.')
        pairs = [source for identifier in identifiers for source in source_pairs(by_id[identifier])]
        base['aliases'] = pairs[1:]
        for identifier in identifiers[1:]:
            del by_id[identifier]
        consumed.update(identifiers)
    return validate_rows(list(by_id.values()))


def _load(path):
    if not path.exists():
        return {'schema': 2, 'revision': 0, 'rows': []}
    try:
        data = json.loads(path.read_text(encoding='utf-8-sig'))
        if data['schema'] not in (1, 2) or not isinstance(data['revision'], int):
            raise ValueError('unrecognized format')
        rows = validate_rows(data['rows'])
        if data['schema'] == 1:
            data = {'schema': 2, 'revision': data['revision'] + 1, 'rows': rows}
            _write(path, data)
        else:
            data['rows'] = rows
        return data
    except Exception as exc:
        raise PackingError(f'Cannot read saved product names: {exc}. Keep {path.name} and restore a valid backup.') from exc


@contextmanager
def catalog_lock(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.with_suffix('.lock').open('a+b') as lock:
        import msvcrt
        lock.seek(0)
        if not lock.read(1):
            lock.write(b'0')
            lock.flush()
        lock.seek(0)
        try:
            msvcrt.locking(lock.fileno(), msvcrt.LK_NBLCK, 1)
        except OSError:
            raise PackingError('Product names are being saved in another window. Try again shortly.')
        try:
            yield
        finally:
            lock.seek(0)
            msvcrt.locking(lock.fileno(), msvcrt.LK_UNLCK, 1)


def _write(path, data):
    temp = path.with_name(path.name + '.tmp-' + uuid.uuid4().hex)
    try:
        with temp.open('w', encoding='utf-8', newline='\n') as handle:
            json.dump(data, handle, ensure_ascii=False, indent=2)
            handle.flush()
            os.fsync(handle.fileno())
        if path.exists():
            backup = path.with_suffix('.json.bak')
            backup_temp = temp.with_suffix('.bak')
            shutil.copyfile(path, backup_temp)
            os.replace(backup_temp, backup)
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def load_catalog(path):
    with catalog_lock(path):
        return _load(path)


def save_catalog(path, rows, revision, merge_groups=(), lock_all=False, unlocked_ids=()):
    if not isinstance(lock_all, bool):
        raise PackingError('Invalid save row-lock option.')
    if not isinstance(unlocked_ids, (list, tuple)) or any(not isinstance(i, str) for i in unlocked_ids):
        raise PackingError('Invalid unlocked definition identifiers.')
    unlocked_ids = set(unlocked_ids)
    validated = merge_rows(rows, merge_groups)
    with catalog_lock(path):
        current = _load(path)
        if current['revision'] != revision:
            raise PackingError('The names list changed in another operation. Close Define Name, reopen it, and repeat your edits.')
        incoming = {r['id']: r for r in validated}
        for old in current['rows']:
            new = incoming.get(old['id'])
            if old.get('locked') and old['id'] not in unlocked_ids and (new is None or (new.get('locked') and new != old)):
                raise PackingError('Unlock the definition before changing or deleting its names.')
        # Validate edits and confirmed merges while their explicit unlocks still apply.
        # Relock the final definitions in the same atomic catalog write.
        if lock_all:
            for row in validated:
                row['locked'] = True
        if current['rows'] != validated:
            current = {'schema': 2, 'revision': revision + 1, 'rows': validated}
            _write(path, current)
        return current


def encounter(path, batches):
    """Remember source pairs on PDF scan, including all linked merged originals."""
    with catalog_lock(path):
        current = _load(path)
        rows = copy.deepcopy(current['rows'])
        for batch in batches:
            mapping = definition_mapping(rows, batch.platform)
            for order in batch.orders:
                for item in order.items:
                    key = lookup_key(item.product, item.variation, batch.platform)
                    matches = mapping.get(key, [])
                    row = matched_definition(matches, item.product, item.variation)
                    if not row:
                        product, variation = definition_pair(item.product, item.variation, batch.platform)
                        row = dict(zip(FIELDS, [product, variation, '', '']))
                        rows.append(row)
                        mapping[key] = [row]
                    elif row['alternative_product'] or row['alternative_variation']:
                        # Retire only unfinished, unlocked spacing duplicates. Keep
                        # their original pairs linked to the existing finished row.
                        for duplicate in list(matches):
                            if duplicate is row or duplicate.get('locked') or duplicate['alternative_product'] or duplicate['alternative_variation']:
                                continue
                            if any(lookup_key(s['original_product'], s['original_variation'], batch.platform) != key for s in source_pairs(duplicate)):
                                continue
                            row.setdefault('aliases', []).extend(source_pairs(duplicate))
                            rows.remove(duplicate)
                            matches.remove(duplicate)
                    if batch.platform == 'LAZADA' and ':' in row['original_variation']:
                        pair = definition_pair(row['original_product'], row['original_variation'], 'LAZADA')
                        occupied = any(item_key(s['original_product'], s['original_variation']) == pair
                                       for other in rows if other is not row for s in source_pairs(other))
                        if not occupied and pair != item_key(row['original_product'], row['original_variation']):
                            row.setdefault('aliases', []).append({f: row[f] for f in FIELDS[:2]})
                            row.update(zip(FIELDS[:2], pair))
        if rows != current['rows']:
            current = {'schema': 2, 'revision': current['revision'] + 1, 'rows': validate_rows(rows)}
            _write(path, current)
        return current


def apply_names(batches, catalog):
    def display(item, platform, mapping):
        row = matched_definition(mapping.get(lookup_key(item.product, item.variation, platform), []), item.product, item.variation)
        product, variation = definition_pair(item.product, item.variation, platform)
        # The PDF item/fingerprint stays untouched. Only the printed copy changes.
        return replace(item, product=(row.get('alternative_product') or product).rstrip(' ').upper(),
                       variation=(row.get('alternative_variation') or variation).rstrip(' ').upper())
    return [replace(b, orders=[replace(o, items=[display(i, b.platform, mapping) for i in o.items]) for o in b.orders])
            for b in batches for mapping in [definition_mapping(catalog['rows'], b.platform)]]


def export_excel(path, catalog):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.comments import Comment
    if path.suffix.lower() != '.xlsx':
        raise PackingError('Export names to an .xlsx Excel file.')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Names'
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='19579D')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for cell in (ws['C1'], ws['D1']):
        cell.fill = PatternFill('solid', fgColor='F8CBAD')
        cell.font = Font(name='Calibri', bold=True, color='202020')
    ws.row_dimensions[1].height = 32
    number = 1
    for row in validate_rows(catalog['rows']):
        for source in source_pairs(row):
            number += 1
            values = [source[f] for f in FIELDS[:2]] + [row[f] for f in FIELDS[2:]] + [row['id'], json.dumps(list(item_key(source['original_product'], source['original_variation'])), ensure_ascii=False)]
            for column, value in enumerate(values, 1):
                cell = ws.cell(number, column, value)
                cell.data_type = 's'
                cell.number_format = '@'
                cell.font = Font(name='Calibri', size=11, bold=column in (3, 4), color='19579D' if column in (3, 4) else '202020')
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.fill = PatternFill('solid', fgColor='FFF7E8' if column in (3, 4) else ('F2F5F9' if number % 2 == 0 else 'FFFFFF'))
            ws.row_dimensions[number].height = 60
    from openpyxl.styles import Protection
    for cells in ws.iter_rows(min_row=2):
        definition = next(r for r in catalog['rows'] if r['id'] == cells[4].value)
        for cell in cells[:4]:
            cell.protection = Protection(locked=bool(definition.get('locked')))
    if any(r.get('locked') for r in catalog['rows']):
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
    for column, width in zip('ABCDEF', [72, 34, 48, 34, 36, 36]):
        ws.column_dimensions[column].width = width
    ws.column_dimensions['E'].hidden = ws.column_dimensions['F'].hidden = True
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:F{max(1, ws.max_row)}'
    ws['A1'].comment = Comment('Each linked original pair has its own Excel row so all originals can be edited. Linked originals normally stay one definition; different new alternatives separate them on import. Hidden IDs preserve originals when names change. Missing rows are retained; delete definitions in Define Name. Names become UPPERCASE, trailing spaces are removed, and paragraphs/line breaks are rejected.', 'Packing List Manager')
    for cell in ('C1', 'D1'):
        ws[cell].comment = Comment('Blank uses the original field. Start with a letter or number. One changed alternative updates the linked definition. Different new alternatives separate the original pairs on import, preserving each Excel row. Matching definitions are merged only after confirmation in the app. Keep this workbook OPEN and click Import in Define Name.', 'Packing List Manager')
    temp = path.with_name(path.stem + '.tmp-' + uuid.uuid4().hex + '.xlsx')
    try:
        wb.save(temp)
        os.replace(temp, path)
    except PermissionError as exc:
        raise PackingError('Close this Excel file in Excel, then export again.') from exc
    finally:
        wb.close()
        temp.unlink(missing_ok=True)
    return {'ok': True, 'path': str(path), 'count': len(catalog['rows']), 'originals': number - 1}


def import_excel(path, current_rows):
    """Preview only: retain linked IDs and omitted originals; never save partial edits."""
    from openpyxl import load_workbook
    if path.suffix.lower() != '.xlsx':
        raise PackingError('Choose an .xlsx file exported from Define Name.')
    if path.stat().st_size > 20 * 1024 * 1024:
        raise PackingError('This workbook is too large. Import only the names worksheet.')
    rows = validate_rows(copy.deepcopy(current_rows))
    by_id = {r['id']: i for i, r in enumerate(rows)}
    by_key = {item_key(s['original_product'], s['original_variation']): i for i, r in enumerate(rows) for s in source_pairs(r)}
    replacements, alt_edits, alternatives, additions, targets = {}, {}, {}, [], {}
    priority = []
    updated = set()
    repaired = 0
    reordered = False
    try:
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if 'Names' not in wb.sheetnames:
                raise PackingError('The workbook needs a worksheet named Names.')
            ws = wb['Names']
            if ws.max_row > 20001:
                raise PackingError('Import at most 20,000 names at a time.')
            if [ws.cell(1, c).value for c in range(1, 5)] != list(HEADERS[:4]):
                raise PackingError('Keep the four column headings from a Define Name export.')
            has_ids = ws.cell(1, 5).value == HEADERS[4]
            has_pairs = ws.cell(1, 6).value == HEADERS[5]
            records, workbook_keys = [], set()
            for number, cells in enumerate(ws.iter_rows(min_row=2, max_col=6), 2):
                if not any(c.value is not None and c.value != '' for c in cells[:4]):
                    continue
                values = []
                for cell in cells[:4]:
                    if cell.data_type in ('f', 'e') or isinstance(cell.value, bool):
                        raise PackingError(f'Excel row {number}: use plain text, not formulas, errors, or TRUE/FALSE values.')
                    if cell.value is not None and not isinstance(cell.value, (str, int, float)):
                        raise PackingError(f'Excel row {number}: format names as text.')
                    values.append('' if cell.value is None else str(cell.value))
                row = validate_rows([dict(zip(FIELDS, values))])[0]
                identifier = str(cells[4].value or '') if has_ids else ''
                key = item_key(row['original_product'], row['original_variation'])
                records.append((number, row, identifier, key, cells[5].value if has_pairs else None))
                workbook_keys.add(key)
            for number, row, identifier, key, pair_id in records:
                target = by_id.get(identifier) if identifier in by_id else by_key.get(key)
                old_key = None
                record_repaired = False
                if target is not None:
                    existing = rows[target]
                    old_keys = {item_key(s['original_product'], s['original_variation']) for s in source_pairs(existing)}
                    if pair_id and identifier in by_id:
                        try:
                            old = json.loads(pair_id)
                            if not isinstance(old, list) or len(old) != 2 or not all(isinstance(v, str) for v in old):
                                raise ValueError('invalid pair')
                            old_key = item_key(*old)
                        except (TypeError, ValueError):
                            if key not in by_key:
                                raise PackingError(f'Excel row {number}: the original-pair ID is invalid and this new name cannot be matched safely. Keep your edits and correct this row from a fresh export.')
                            target, old_key = by_key[key], key
                            record_repaired = True
                    else:
                        old_key = key if key in old_keys else item_key(existing['original_product'], existing['original_variation'])
                    if old_key != key and old_key in workbook_keys:
                        # Pasting/sorting just the visible columns can leave an
                        # earlier row's IDs behind. Its original still has a
                        # visible row, so this is not a rename of that original.
                        target, old_key = by_key.get(key), key
                        record_repaired = True
                        reordered = True
                    elif old_key not in old_keys:
                        if key not in by_key:
                            raise PackingError(f'Excel row {number}: the hidden original ID does not match a saved original. Keep your edits and correct this row from a fresh export.')
                        target, old_key = by_key[key], key
                        record_repaired = True
                repaired += int(record_repaired)
                if target is None:
                    additions.append(row)
                    priority.append(row['id'])
                    continue
                existing = rows[target]
                if (target, old_key) in targets:
                    earlier = targets[(target, old_key)]
                    raise PackingError(f'Excel rows {earlier} and {number}: both update the same original product and variation. Keep one row for this original; merge different originals in the app.')
                targets[(target, old_key)] = number
                priority.append((target, old_key))
                if existing.get('locked'):
                    original = next(s for s in source_pairs(existing) if item_key(s['original_product'], s['original_variation']) == old_key)
                    if key != item_key(original['original_product'], original['original_variation']) or any(row[f] != existing[f] for f in FIELDS[2:]):
                        raise PackingError(f'Excel row {number}: this definition is locked. Unlock it in Define Name before editing.')
                replacements[(target, old_key)] = {f: row[f] for f in FIELDS[:2]}
                alternatives[(target, old_key)] = tuple(row[f] for f in FIELDS[2:])
                for field in FIELDS[2:]:
                    if row[field] != existing[field]:
                        changes = alt_edits.setdefault((target, field), set())
                        changes.add(row[field])
                updated.add(target)
        finally:
            wb.close()
    except PackingError:
        raise
    except Exception as exc:
        raise PackingError(f'Could not read the names workbook: {exc}') from exc
    imported, routed, separated = [], {}, 0
    for target, row in enumerate(rows):
        # A single changed value still updates a shared definition. If linked
        # originals receive distinct new values, keep their complete Excel
        # pairs together instead of combining fields or rejecting the import.
        split = any(len(alt_edits.get((target, field), ())) > 1 for field in FIELDS[2:])
        if reordered:
            # A pasted/reordered inventory list states complete per-original
            # alternatives. Do not turn its sizes/names into a shared edit.
            split |= len({alternatives.get((target, item_key(s['original_product'], s['original_variation'])), tuple(row[f] for f in FIELDS[2:])) for s in source_pairs(row)}) > 1
        shared = tuple(next(iter(alt_edits[(target, field)])) if not split and (target, field) in alt_edits else row[field] for field in FIELDS[2:])
        groups = {}
        for source in source_pairs(row):
            old_key = item_key(source['original_product'], source['original_variation'])
            values = alternatives.get((target, old_key), tuple(row[f] for f in FIELDS[2:])) if split else shared
            groups.setdefault(values, []).append((old_key, replacements.get((target, old_key), source)))
        for index, (values, originals) in enumerate(groups.items()):
            definition = dict(row, **originals[0][1])
            definition['aliases'] = [s for _, s in originals[1:]]
            definition.update(zip(FIELDS[2:], values))
            if index:
                definition['id'] = uuid.uuid4().hex
                separated += 1
            imported.append(definition)
            for old_key, _ in originals:
                routed[(target, old_key)] = definition['id']
    priority = [routed[p] if isinstance(p, tuple) else p for p in priority]
    return {'ok': True, 'rows': validate_rows(imported + additions, allow_shared=True), 'priority': list(dict.fromkeys(priority)), 'added': len(additions), 'updated': len(updated), 'separated': separated, 'repaired': repaired}
