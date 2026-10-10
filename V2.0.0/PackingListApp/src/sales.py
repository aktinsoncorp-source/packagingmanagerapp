"""Daily item quantities from current source PDFs, independent of packing/reprints."""
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import os
import re
import textwrap
import uuid

from extractors import PackingError, file_hash
from names import apply_names, load_catalog, encounter

SUMMARY_HEADERS = ['date', 'product_name', 'variation', 'quantity', 'order_count',
                   'shopee_quantity', 'lazada_quantity', 'tiktok_quantity', 'unit_price', 'price_status']
DETAIL_HEADERS = ['date', 'platform', 'order_id', 'line_number', 'original_product_name',
                  'original_variation', 'product_name', 'variation', 'quantity',
                  'source_pdf', 'source_sha256', 'order_fingerprint', 'unit_price']


def export_sales(folder, catalog_path):
    from engine import folder_lock, recover, scan_internal, preferred_order_sources, INTERNAL
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    if not folder.is_dir():
        raise PackingError('Choose an existing dated folder before exporting item sales.')
    from sku_registry import load_registry
    registry = load_registry(catalog_path)
    with folder_lock(folder) as out:
        recover(out)
        entries, batches, _ = scan_internal(folder, registry)
        problems = [e for e in entries if e['status'] in ('ERROR', 'REVIEW') or e.get('requires_revision')]
        if problems:
            raise PackingError('Resolve these source PDFs before exporting item sales: ' +
                               '; '.join(e['path'] + ': ' + e['message'] for e in problems))
        if not batches:
            raise PackingError('No readable packing orders to export for this date.')
        catalog = encounter(catalog_path, batches.values(), registry=registry) if registry is not None else load_catalog(catalog_path)
        keys = sorted(batches)
        named = dict(zip(keys, apply_names([batches[k] for k in keys], catalog, registry=registry)))
        preferred = preferred_order_sources({key: batches[key] for key in keys}, registry)
        seen, details, totals = {}, [], {}
        day = datetime.strptime(folder.name, '%Y%m%d').strftime('%Y-%m-%d')
        for key in keys:
            batch = batches[key]
            for original, order in zip(batch.orders, named[key].orders):
                fingerprint = original.fingerprint()
                priced_fingerprint = (fingerprint, sorted((i.product, i.variation, i.pieces, i.unit_price or '') for i in original.items))
                if original.key in seen:
                    if seen[original.key] != priced_fingerprint:
                        raise PackingError(f'Conflicting versions of order {original.order_id}; resolve its source PDFs first.')
                seen[original.key] = priced_fingerprint
                if preferred[original.key] != key:
                    continue
                platform = 'SHOPEE' if original.platform in ('SPX', 'FLASH', 'SHOPEE') else original.platform
                for line, (raw, item) in enumerate(zip(original.items, order.items), 1):
                    pair = (item.product, item.variation)
                    total = totals.setdefault(pair, {'quantity': 0, 'orders': set(), 'platforms': defaultdict(int), 'prices': set(), 'priced_quantity': 0})
                    total['quantity'] += item.pieces
                    total['orders'].add(original.key)
                    total['platforms'][platform] += item.pieces
                    if raw.unit_price is not None:
                        total['prices'].add(Decimal(raw.unit_price))
                        total['priced_quantity'] += item.pieces
                    details.append([day, platform, str(original.order_id), line, raw.product, raw.variation,
                                    item.product, item.variation, item.pieces, key, batch.sha256, fingerprint, float(raw.unit_price) if raw.unit_price is not None else None])
        if not details:
            raise PackingError('No readable packing orders to export for this date.')
        def price_status(total):
            if not total['prices']: return 'Not provided in source'
            if len(total['prices']) > 1: return 'Multiple source prices; see Order Items'
            if total['priced_quantity'] != total['quantity']: return 'Partial source price; see Order Items'
            return 'Source unit price'
        summary = [[day, product, variation, total['quantity'], len(total['orders']),
                    total['platforms']['SHOPEE'], total['platforms']['LAZADA'], total['platforms']['TIKTOK'],
                    float(next(iter(total['prices']))) if len(total['prices']) == 1 and total['priced_quantity'] == total['quantity'] else None, price_status(total)]
                   for (product, variation), total in sorted(totals.items())]
        quantity = sum(row[3] for row in summary)
        if quantity != sum(row[8] for row in details):
            raise PackingError('Item sales quantities do not reconcile with order details.')
        info = [
            ['date', day], ['total_orders', len(seen)], ['total_quantity', quantity],
            ['item_variation_pairs', len(summary)], ['source_pdf_count', len(batches)],
            ['name_catalog_revision', catalog['revision']],
            ['exported_at', datetime.now().astimezone().isoformat(timespec='seconds')],
            ['scope', 'All readable packing PDFs in the selected dated folder, regardless of Use checkboxes or packing output status.'],
            ['quantity_basis', 'Ordered units from packing PDFs; revenue, returns and cancellations are not inferred.'],
            ['price_basis', 'Unit prices are copied from source PDFs, not inferred. Blank means unavailable, mixed or partial; see Order Items for exact row prices. Currency is not assumed when absent from the source.'],
            ['date_basis', 'Selected source-folder date; not an inferred order/payment date.'],
            ['names', 'Current Define Name alternatives; blank fields retain original names. Original fields remain in Order Items.'],
            ['duplicates', 'Each platform/order ID counted once. SHOPEE combines SPX and FLASH. Matching source copies and reprints do not add sales.'],
            ['erp_import', 'This is a full-date snapshot. Replace the prior snapshot for that date; do not append repeated exports.'],
            ['order_counts', 'Item order_count is distinct per item; do not add these counts for the daily total.'],
        ]
        wb = Workbook()
        wb.remove(wb.active)
        def sheet(title, headers, rows, widths, table_name):
            ws = wb.create_sheet(title)
            for row in [headers] + rows:
                ws.append(row)
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        if len(cell.value) > 32767:
                            raise PackingError('A source field exceeds the Excel cell limit.')
                        cell.data_type = 's'  # Preserve IDs/leading zeros and literal formula-like names.
                    if headers[cell.column-1] == 'unit_price': cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                lines = max(len(textwrap.wrap(str(v), max(8, int(width * .85)))) or 1 for v, width in zip(row, widths))
                ws.row_dimensions[ws.max_row].height = max(26, 15 * lines + 8)
            ws.freeze_panes = 'A2'
            ws.sheet_view.showGridLines = False
            ws.row_dimensions[1].height = 30
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='234E70')
                cell.font = Font(color='FFFFFF', bold=True)
            for index, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(index)].width = width
            table = Table(displayName=table_name, ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(table)
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = '1:1'
            return ws
        sheet('Item Sales', SUMMARY_HEADERS, summary, [14, 52, 34, 14, 14, 20, 20, 20, 14, 36], 'ItemSales')
        sheet('Order Items', DETAIL_HEADERS, details, [14, 14, 24, 14, 52, 34, 52, 34, 14, 52, 68, 68, 14], 'OrderItems')
        sheet('Export Info', ['field', 'value'], info, [28, 110], 'ExportInfo')
        prefix = f'{day}_Item_Sales_'
        numbers = [int(m[1]) for p in out.iterdir()
                   if (m := re.fullmatch(re.escape(prefix) + r'(\d+)\.xlsx', p.name, re.I))]
        target = out / f'{prefix}{max(numbers, default=0)+1}.xlsx'
        temp = out / INTERNAL / f'.sales-{uuid.uuid4().hex}.xlsx'
        try:
            wb.save(temp)
            wb.close()
            check = load_workbook(temp, read_only=True, data_only=False)
            try:
                for title, headers, expected in [('Item Sales', SUMMARY_HEADERS, summary), ('Order Items', DETAIL_HEADERS, details)]:
                    actual = list(check[title].values)
                    normalized = [tuple(None if v == '' else v for v in row) for row in [headers] + expected]
                    if actual != normalized:
                        raise PackingError('Excel verification failed: exported data differs from source totals.')
            finally:
                check.close()
            for batch in batches.values():
                if file_hash(batch.path) != batch.sha256:
                    raise PackingError('A source PDF changed during export. Scan again and retry.')
            # Atomic no-overwrite publication on the supported Windows/NTFS workstation.
            os.link(temp, target)
        finally:
            wb.close()
            temp.unlink(missing_ok=True)
        return {'ok': True, 'path': str(target), 'filename': target.name,
                'orders': len(seen), 'quantity': quantity, 'items': len(summary),
                'message': f'Exported {len(summary)} items / variations, {quantity} units from {len(seen)} orders.'}
