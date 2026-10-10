"""Local batch processing with content verification and recoverable commits."""
from __future__ import annotations

from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import asdict, replace
from datetime import datetime
from pathlib import Path
import argparse
import copy
import hashlib
import json
import os
import re
import shutil
import sys
import uuid

from docx import Document

from extractors import Batch, Item, Order, PackingError, UnsupportedPdf, clean, file_hash, parse_pdf, source_price, normalize_sku
from documents import build_document, verify_word

VERSION = '2.0.0'
INTERNAL = '.PackingData'
HISTORY = INTERNAL + '/Processing_History.json'
LOG = INTERNAL + '/Processed_Files.txt'
JOURNAL = INTERNAL + '/.packing-transaction.json'
COURIERS = INTERNAL + '/Courier_Selections.json'
LOCK = INTERNAL + '/.packing.lock'


def natural(value):
    return [int(x) if x.isdigit() else x.lower() for x in re.split(r'(\d+)', str(value))]


def dated_folder(value):
    p = Path(value).resolve()
    try:
        datetime.strptime(p.name, '%Y%m%d')
    except ValueError:
        raise PackingError('Choose a dated folder named YYYYMMDD, such as 20260926.')
    if not re.fullmatch(r'\d{8}', p.name):
        raise PackingError('The dated folder name must contain exactly eight digits: YYYYMMDD.')
    return p


def relpath(path, folder):
    try:
        return Path(path).resolve().relative_to(folder.resolve()).as_posix()
    except ValueError:
        raise PackingError('The selected PDF must be inside the chosen dated folder. Use Add PDFs to copy it there.')


def atomic_json(path, value):
    temp = path.with_name(path.name + '.tmp-' + uuid.uuid4().hex)
    with temp.open('w', encoding='utf-8', newline='\n') as f:
        json.dump(value, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    os.replace(temp, path)


def internal_directory(out):
    """Persistent processing records, hidden in normal Windows Explorer views."""
    internal = safe_child(out, INTERNAL)
    internal.mkdir(parents=True, exist_ok=True)
    if os.name == 'nt':
        import ctypes
        import stat
        attributes = internal.stat().st_file_attributes
        if not attributes & stat.FILE_ATTRIBUTE_HIDDEN:
            kernel32 = ctypes.WinDLL('kernel32', use_last_error=True)
            kernel32.SetFileAttributesW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32]
            kernel32.SetFileAttributesW.restype = ctypes.c_int
            if not kernel32.SetFileAttributesW(str(internal), attributes | stat.FILE_ATTRIBUTE_HIDDEN):
                raise PackingError(f'Cannot hide the internal processing folder: {ctypes.WinError(ctypes.get_last_error())}')
    return internal


def migrate_internal_records(out):
    # Called while locked. Check every collision before moving anything. Atomic
    # renames also make a partly completed migration safe to resume after a crash.
    moves = []
    for relative in (HISTORY, LOG, JOURNAL, COURIERS):
        source, target = out / Path(relative).name, out / relative
        if source.exists():
            if target.exists() and file_hash(source) != file_hash(target):
                raise PackingError(f'Two different copies of {source.name} exist in 000 Output and {INTERNAL}. Keep both and ask your supervisor to review them.')
            moves.append((source, target))
    for source, target in moves:
        if target.exists():
            source.unlink()  # Identical duplicate already verified above.
        else:
            os.rename(source, target)


@contextmanager
def file_lock(path):
    lock = path.open('a+b')
    try:
        import msvcrt
        lock.seek(0)
        try:
            msvcrt.locking(lock.fileno(), msvcrt.LK_NBLCK, 1)
        except OSError:
            raise PackingError('This date is being processed in another window. Wait for it to finish.')
        try:
            if os.fstat(lock.fileno()).st_size == 0:
                lock.write(b'0')
                lock.flush()
            yield
        finally:
            lock.seek(0)
            msvcrt.locking(lock.fileno(), msvcrt.LK_UNLCK, 1)
    finally:
        lock.close()


@contextmanager
def folder_lock(folder):
    out = folder / '000 Output'
    internal_directory(out)
    with file_lock(out / LOCK):
        legacy_lock = out / '.packing.lock'
        if legacy_lock.exists():
            # Refuse migration while an older engine still owns the old lock.
            with file_lock(legacy_lock):
                migrate_internal_records(out)
            legacy_lock.unlink()
        else:
            migrate_internal_records(out)
        yield out


def safe_child(out, relative):
    p = (out / relative).resolve()
    if not p.is_relative_to(out.resolve()):
        raise PackingError('A saved processing record contains an invalid path. Ask your supervisor to review it.')
    return p


def recover(out):
    journal = out / JOURNAL
    if not journal.exists():
        return
    try:
        plan = json.loads(journal.read_text(encoding='utf-8'))
        for entry in plan['files']:
            target = safe_child(out, entry['target'])
            source = safe_child(out, entry['stage'])
            if target.exists():
                if file_hash(target) != entry['sha256']:
                    raise PackingError(f'Interrupted batch: {target.name} differs from the saved record. Keep all files and ask your supervisor to review.')
            elif source.exists() and file_hash(source) == entry['sha256']:
                os.rename(source, target)
            else:
                raise PackingError('Interrupted batch: a completed document is missing. Keep the processing files and ask your supervisor to review.')
        atomic_json(out / HISTORY, plan['history'])
        log = out / LOG
        text = log.read_text(encoding='utf-8-sig') if log.exists() else 'PACKING LIST PROCESSING LOG\n'
        marker = 'Transaction: ' + plan['id']
        if marker not in text:
            with log.open('a', encoding='utf-8', newline='\n') as f:
                f.write('\n' + plan['log'] + '\n' + marker + '\n')
                f.flush()
                os.fsync(f.fileno())
        journal.unlink()
        stage = safe_child(out, plan['stage_dir'])
        if stage.exists() and not any(stage.iterdir()):
            stage.rmdir()
    except PackingError:
        raise
    except Exception as exc:
        raise PackingError(f'Cannot recover the interrupted batch: {exc}. Keep the processing files for review.') from exc


def commit(out, history, files, stage, log):
    plan = {'id': uuid.uuid4().hex, 'history': history, 'files': files,
            'stage_dir': stage.relative_to(out).as_posix(), 'log': log}
    atomic_json(out / JOURNAL, plan)
    recover(out)


def read_completed_doc(path, group):
    from picking import is_picking_table
    records = {}
    d = Document(path)
    summary_total = None
    for table in d.tables:
        if is_picking_table(table):
            if summary_total is not None or table._tbl is not d.tables[-1]._tbl:
                raise PackingError(f'The earlier document {path.name} has an unexpected summary location.')
            try:
                quantities = [int(r.cells[3].text) for r in table.rows[1:-1]]
                summary_total = int(table.rows[-1].cells[3].text)
                if (table.rows[-1].cells[1].text != 'TOTAL PIECES' or not quantities
                        or any(q <= 0 for q in quantities) or sum(quantities) != summary_total):
                    raise ValueError('invalid total')
            except (ValueError, IndexError) as exc:
                raise PackingError(f'The earlier document {path.name} has an invalid picking summary.') from exc
            continue
        head = table.cell(0, 0).text
        m = re.search(r'Order (ID|Number):\s*(\w+)', head)
        if not m or len(table.rows) < 3:
            raise PackingError(f'The earlier document {path.name} has an unfamiliar order table.')
        platform = ('LAZADA' if m[1] == 'Number' else 'TIKTOK') if group == 'LAZADA_TIKTOK' else group
        if platform == 'LAZADA':
            platform = 'LAZADA'
        from documents import COMMENT_LABEL, order_cells
        priced = [c.text for c in order_cells(table.rows[1])[4:]] == ['Unit Price', 'Total Price']
        item_rows = [order_cells(row) for row in table.rows[2:]]
        comment = ''
        if item_rows[-1][0].text.startswith(COMMENT_LABEL):
            cells = item_rows.pop()
            comment = cells[1].text if len(cells) == 2 else cells[0].text[len(COMMENT_LABEL):].strip()
        items = [Item(clean(r[1].text), clean(r[2].text), int(r[3].text),
                      source_price(r[4].text, f'Order {m[2]}') if priced else None,
                      source_price(r[5].text, f'Order {m[2]}', 'total price') if priced else None) for r in item_rows]
        order = Order(platform, m[2], items, comment)
        records[order.key] = {'fingerprint': order.fingerprint(), 'order': asdict(order),
                              'output': path.name, 'source': 'Earlier completed document'}
    if summary_total is not None and summary_total != sum(i['pieces'] for r in records.values() for i in r['order']['items']):
        raise PackingError(f'The picking summary in {path.name} does not match its orders.')
    return records


def load_history(folder, batches):
    out = folder / '000 Output'
    p = out / HISTORY
    if p.exists():
        try:
            h = json.loads(p.read_text(encoding='utf-8'))
            if h.get('schema') != 1 or h.get('date') != folder.name:
                raise ValueError('version or date mismatch')
            for field in ('sources', 'orders', 'outputs'):
                if not isinstance(h.get(field), dict): raise ValueError('invalid ' + field)
            return h
        except Exception as exc:
            raise PackingError(f'Processing_History.json cannot be read ({exc}). Restore its backup before processing more orders.') from exc
    h = {'schema': 1, 'date': folder.name, 'sources': {}, 'orders': {}, 'outputs': {}, 'legacy': False}
    log_path = out / LOG
    docs = [p for p in out.glob('*_Packing_Checklist_*') if p.suffix.lower() in ('.docx', '.pdf')] if out.exists() else []
    if not log_path.exists():
        if docs:
            raise PackingError('Completed packing documents exist, but Processed_Files.txt is missing. Restore the log before making another batch.')
        return h
    text = log_path.read_text(encoding='utf-8-sig')
    if 'COMPLETED BY PACKING LIST MANAGER' in text or re.search(r'^Transaction:', text, re.M):
        raise PackingError('Processing_History.json is missing from a batch created by this app. Restore that file from backup before continuing.')
    h['legacy'] = True
    blocks = re.split(r'(?m)^(?=[A-Z_]+_Packing_Checklist_.*\.docx\s*$)', text)[1:]
    recorded_paths = set()
    missing_sources = set()
    for block in blocks:
        lines = block.splitlines()
        name = lines[0].strip()
        source_paths = [m.replace('\\','/') for m in re.findall(r'^\s+-\s+(.+?\.pdf)\s*\|', block, re.M)]
        recorded_paths.update(source_paths)
        if re.search(r'^Status:\s*SUPERSEDED', block, re.M):
            continue
        path = safe_child(out, name)
        group = name.split('_Packing_Checklist_')[0]
        if group not in ('SPX','FLASH','LAZADA','LAZADA_TIKTOK'):
            raise PackingError(f'Unrecognized earlier output group in {name}.')
        if path.is_file():
            records = read_completed_doc(path, group)
        else:
            # Pre-app logs identify the source PDFs even if their Word file was
            # deleted before the first migration. No old content is inferred.
            records = {}
            missing_sources.update(source_paths)
        h['orders'].update(records)
        h['outputs'][name] = {'group': group, 'orders': list(records), 'mode': 'legacy',
                              'created': 'Before Packing List Manager', 'status': 'active',
                              'order_fingerprints': {k: v['fingerprint'] for k, v in records.items()},
                              'sources': [{'path': key, 'sha256': None,
                                           'order_keys': [o.key for o in batches[key].orders if o.key in records] if key in batches else []}
                                          for key in source_paths]}
    if not blocks:
        raise PackingError('The existing processing log could not be imported safely. Ask your supervisor to review it.')
    for key, batch in batches.items():
        if key not in recorded_paths:
            continue
        same = all(o.key in h['orders'] and h['orders'][o.key]['fingerprint'] == o.fingerprint() for o in batch.orders)
        missing_only = key in missing_sources and not any(o.key in h['orders'] for o in batch.orders)
        h['sources'][key] = {'sha256': batch.sha256 if same or missing_only else None,
                             'platform': batch.platform, 'order_keys': [o.key for o in batch.orders],
                             'outputs': sorted(name for name, output in h['outputs'].items() if any(s['path'] == key for s in output['sources'])),
                             'legacy_missing': missing_only,
                             'note': 'Earlier output document missing; regenerate from the current PDF' if missing_only else
                                     ('Verified against earlier completed documents' if same else 'Earlier source differs from completed documents')}
    for key in recorded_paths - batches.keys():
        h['sources'][key] = {'sha256': None, 'order_keys': [], 'outputs': [], 'note': 'Source named in earlier processing log'}
    return h


def upgrade_output_links(folder, history):
    """Add immutable per-output source/order manifests to older app records."""
    changed = False
    for name, output in history['outputs'].items():
        records = None
        if 'order_fingerprints' not in output:
            path = safe_child(folder / '000 Output', name)
            records = read_completed_doc(path, output['group']) if path.is_file() else {}
            output['order_fingerprints'] = {key: value['fingerprint'] for key, value in records.items()}
            for key, value in history['orders'].items():
                if value['output'] == name:
                    output['order_fingerprints'].setdefault(key, value['fingerprint'])
            changed = True
        if 'order_snapshots' not in output:
            if records is None:
                path = safe_child(folder / '000 Output', name)
                records = read_completed_doc(path, output['group']) if path.is_file() else {}
            output['order_snapshots'] = {}
            for key, fingerprint in output.get('order_fingerprints', {}).items():
                saved = records.get(key) or history['orders'].get(key)
                if saved and saved['fingerprint'] == fingerprint:
                    output['order_snapshots'][key] = saved['order']
            changed = True
        if 'record_id' not in output:
            output['record_id'] = uuid.uuid4().hex
            changed = True
        if 'sources' not in output:
            output['sources'] = []
            for key, source in history['sources'].items():
                order_keys = sorted(set(source.get('order_keys', [])) & set(output['orders']))
                if name in source.get('outputs', []) or order_keys:
                    output['sources'].append({'path': key, 'sha256': source.get('sha256'), 'order_keys': order_keys})
            changed = True
    return changed


def completed_outputs(folder, history):
    """Visible files in platform/numeric order, with their saved included data."""
    result = []
    for name, output in history['outputs'].items():
        if not safe_child(folder / '000 Output', name).is_file():
            continue
        snapshots = output.get('display_order_snapshots', output.get('order_snapshots', {}))
        sources = []
        for source in sorted(output.get('sources', []), key=lambda s: natural(s['path'])):
            included = []
            for key in source.get('order_keys', []):
                if key not in output['orders']:
                    continue
                order = snapshots.get(key)
                included.append(copy.deepcopy(order) if order else {'order_id': key.split(':', 1)[-1], 'items': [], 'details_available': False})
            sources.append({'path': source['path'], 'orders': included,
                            'pieces': sum(i['pieces'] for order in included for i in order['items'])})
        match = re.search(r'_(\d+)\.(?:docx|pdf)$', name, re.I)
        result.append({'filename': name, 'group': output['group'], 'sequence': int(match[1]) if match else 0,
                       'status': output['status'], 'exists': True, 'created': output.get('created', ''),
                       'order_count': len(output['orders']), 'pieces': sum(i['pieces'] for o in snapshots.values() for i in o['items']),
                       'included_sources': sources, 'record_id': output['record_id']})
    platform_order = {'SPX': 0, 'FLASH': 1, 'LAZADA_TIKTOK': 2, 'LAZADA': 3}
    return sorted(result, key=lambda o: (platform_order.get(o['group'], 4), o['sequence'], o['filename']))


def next_output_filename(out, group, date, reserved=(), output_format='docx'):
    prefix = f'{group}_Packing_Checklist_{datetime.strptime(date,"%Y%m%d"):%Y-%m-%d}_'
    numbers = []
    for name in [p.name for p in out.iterdir() if p.is_file() and p.suffix.lower() in ('.docx', '.pdf')] + list(reserved):
        match = re.fullmatch(re.escape(prefix) + r'(\d+)\.(?:docx|pdf)', name, re.I)
        if match:
            numbers.append(int(match[1]))
    return f'{prefix}{max(numbers, default=0)+1}.{output_format}'


def available_order_outputs(folder, history):
    """All surviving copies of the current order version, excluding revisions."""
    available = defaultdict(list)
    for name, output in history['outputs'].items():
        if output.get('status') == 'superseded' or not safe_child(folder / '000 Output', name).is_file():
            continue
        for key, fingerprint in output.get('order_fingerprints', {}).items():
            if key not in output.get('superseded_orders', []):
                available[(key, fingerprint)].append(name)
    return available


def forget_deleted_outputs(folder, history):
    """Remove deleted-file coverage; retain immutable manifests only for audit."""
    missing = [name for name in history['outputs']
               if not safe_child(folder / '000 Output', name).is_file()]
    if not missing:
        return
    for name in missing:
        output = history['outputs'].pop(name)
        history.setdefault('archived_outputs', []).append(
            {'filename': name, **output, 'archived_at': datetime.now().isoformat(),
             'reason': 'Output deleted; processing coverage cleared'})
    available = available_order_outputs(folder, history)
    for key, order in list(history['orders'].items()):
        if order.get('output') not in missing:
            continue
        copies = available.get((key, order['fingerprint']), [])
        if not copies:
            del history['orders'][key]
        elif order.get('output') not in copies:
            order['output'] = copies[0]
    for key, source in list(history['sources'].items()):
        source['order_keys'] = [k for k in source.get('order_keys', []) if k in history['orders']]
        source['outputs'] = [n for n in source.get('outputs', []) if n in history['outputs']]
        if not source['order_keys']:
            del history['sources'][key]


def batch_details(batch, old=None):
    parts = [f'{batch.platform} | {len(batch.orders)} orders | {batch.pieces} pieces', '']
    if old:
        parts.extend(['PREVIOUSLY RECORDED', old, '', 'CURRENT PDF', ''])
    for order in batch.orders:
        parts.append(f'Order {order.order_id}')
        for i in order.items:
            parts.append(f'  {i.pieces} pcs | {i.product} | {i.variation}' + (f' | SKU: {i.seller_sku}' if i.seller_sku else ''))
        parts.append('')
    return '\r\n'.join(parts)


def load_couriers(folder):
    path = folder / '000 Output' / COURIERS
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding='utf-8'))
        if data.get('schema') != 1 or not isinstance(data.get('choices'), dict):
            raise ValueError('invalid format')
        for key, value in data['choices'].items():
            if not isinstance(value, dict) or value.get('carrier') not in ('SPX', 'FLASH') or not re.fullmatch(r'[a-f0-9]{64}', value.get('sha256', '')):
                raise ValueError('invalid saved courier')
        return data['choices']
    except Exception as exc:
        raise PackingError(f'{COURIERS} cannot be read ({exc}). Restore its backup or correct the courier settings before generating.') from exc


def store_courier(out, choices, key, sha256, carrier):
    if carrier == 'Auto':
        choices.pop(key, None)
    else:
        choices[key] = {'sha256': sha256, 'carrier': carrier}
    atomic_json(out / COURIERS, {'schema': 1, 'choices': choices})


def set_courier(folder, selection, carrier):
    if carrier not in ('Auto', 'SPX', 'FLASH'):
        raise PackingError('Choose Auto, SPX or FLASH.')
    with folder_lock(folder) as out:
        recover(out)
        path = safe_child(folder, selection['path'])
        if path.suffix.lower() != '.pdf' or '000 Output' in path.relative_to(folder).parts:
            raise PackingError('Select a source packing PDF.')
        digest = file_hash(path)
        if digest != selection.get('sha256'):
            raise PackingError('This PDF changed after the scan. Scan again before choosing its courier.')
        try:
            batch = parse_pdf(path, allow_unassigned=True)
            platform = batch.platform
        except PackingError as exc:
            platform = getattr(exc, 'platform', '')
            if platform not in ('SHOPEE', 'SPX', 'FLASH'):
                raise
        if platform not in ('SHOPEE', 'SPX', 'FLASH'):
            raise PackingError('Courier selection applies only to Shopee packing PDFs.')
        if file_hash(path) != digest:
            raise PackingError('The PDF changed while selecting its courier. Scan again.')
        store_courier(out, load_couriers(folder), relpath(path, folder), digest, carrier)
    return {'ok': True}


def order_sku_evidence(order, registry=None):
    """Separate SKU evidence; historical order fingerprints stay unchanged."""
    def field(value, key, default=''):
        return value.get(key, default) if isinstance(value, dict) else getattr(value, key, default)
    platform = field(order, 'platform')
    platform = 'SHOPEE' if platform in ('SPX', 'FLASH', 'SHOPEE') else platform
    groups = {}
    for item in field(order, 'items', []):
        key = (clean(field(item, 'product')), clean(field(item, 'variation')), field(item, 'pieces', 0))
        evidence = groups.setdefault(key, {'count': 0, 'skus': Counter()})
        evidence['count'] += 1
        sku = normalize_sku(field(item, 'seller_sku'))
        if not sku or sku in ('-', '—', 'N/A', 'NA', 'NONE', 'NULL'):
            continue
        resolved = registry.resolve(platform, sku) if registry is not None else {}
        identity = ('CATALOG', resolved['sku']) if resolved else (platform, sku)
        evidence['skus'][identity] += 1
    return groups


def conflicting_order_skus(orders, registry=None):
    """Blank legacy cells may be filled, but contradictory printed codes cannot."""
    evidence = [order_sku_evidence(order, registry) for order in orders]
    if len(evidence) < 2:
        return False
    for key in set.intersection(*(set(value) for value in evidence)):
        required = Counter()
        for value in evidence:
            required |= value[key]['skus']
        if sum(required.values()) > min(value[key]['count'] for value in evidence):
            return True
    return False


def complete_sku_sources(orders, registry=None):
    """Indices of real copies containing all noncontradictory SKU evidence."""
    evidence = [order_sku_evidence(order, registry) for order in orders]
    required = {}
    for value in evidence:
        for key, group in value.items():
            required[key] = required.get(key, Counter()) | group['skus']
    return [index for index, value in enumerate(evidence)
            if all(skus <= value.get(key, {}).get('skus', Counter()) for key, skus in required.items())]


def preferred_order_sources(batches, registry=None, priorities=None):
    """Choose an actual source copy with the most SKU evidence, preserving ties."""
    priorities = priorities or {}
    preferred, scores, candidates = {}, {}, defaultdict(list)
    for key, batch in batches.items():
        for order in batch.orders:
            candidates[order.key].append((key, order))
            evidence = order_sku_evidence(order, registry)
            score = (priorities.get(key, 0), sum(sum(value['skus'].values()) for value in evidence.values()))
            if order.key not in scores or score > scores[order.key]:
                preferred[order.key], scores[order.key] = key, score
    for order_key, copies in candidates.items():
        orders = [order for _, order in copies]
        if len(copies) > 1 and len({order.fingerprint() for order in orders}) == 1 and not conflicting_order_skus(orders, registry):
            complete = complete_sku_sources(orders, registry)
            if not complete:
                raise PackingError(f'Order {orders[0].order_id}: its PDF copies show different partial SKU information. Keep or select a complete copy with all seller SKUs before exporting.')
            eligible = {copies[index][0] for index in complete}
            if preferred[order_key] not in eligible:
                # Explicit revisions remain authoritative; do not silently
                # replace a reviewed new source with an older complete copy.
                priority = scores[order_key][0]
                eligible = [copies[index][0] for index in complete if priorities.get(copies[index][0], 0) == priority]
                if not eligible:
                    raise PackingError(f'Order {orders[0].order_id}: the reviewed source is missing SKU information present in another copy. Use a complete current source PDF.')
                preferred[order_key] = eligible[0]
    return preferred


def scan_internal(folder, registry=None):
    batches, errors, modes = {}, {}, {}
    choices = load_couriers(folder)
    if folder.exists():
        paths = sorted((p for p in folder.rglob('*') if p.is_file() and p.suffix.lower() == '.pdf'
                        and '000 Output' not in p.relative_to(folder).parts), key=natural)
        for path in paths:
            key = relpath(path, folder)
            try:
                choice = choices.get(key, {})
                mode = choice.get('carrier') if choice and file_hash(path) == choice['sha256'] else None
                modes[key] = mode or 'Auto'
                batches[key] = parse_pdf(path, carrier=mode, allow_unassigned=True)
            except UnsupportedPdf:
                continue
            except Exception as exc:
                try:
                    digest = file_hash(path)
                except OSError:
                    digest = ''
                platform = getattr(exc, 'platform', '')
                errors[key] = {'path': key, 'platform': platform, 'is_shopee': platform in ('SHOPEE', 'SPX', 'FLASH'),
                               'courier': modes.get(key, 'Auto'), 'status': 'ERROR', 'message': str(exc),
                               'details': str(exc), 'orders': 0, 'pieces': 0, 'new_orders': 0, 'sha256': digest, 'outputs': []}
    history = load_history(folder, batches)
    upgrade_output_links(folder, history)
    forget_deleted_outputs(folder, history)
    available = available_order_outputs(folder, history)
    occurrences = defaultdict(list)
    for key, batch in batches.items():
        for order in batch.orders:
            occurrences[order.key].append((key, order))
    conflicts = {key for values in occurrences.values()
                 if len({order.fingerprint() for _, order in values}) > 1 or
                 conflicting_order_skus([order for _, order in values], registry) or
                 not complete_sku_sources([order for _, order in values], registry)
                 for key, _ in values}
    seen = {}
    entries = []
    for key in sorted(set(batches) | set(errors), key=natural):
        if key in errors:
            entries.append(errors[key])
            continue
        batch = batches[key]
        source = history['sources'].get(key)
        prior = history['orders']
        sku_changes = {o.key for o in batch.orders if o.key in prior and
                       conflicting_order_skus([prior[o.key].get('order', {}), o], registry)}
        changed_orders = [o for o in batch.orders if o.key in prior and
                          (prior[o.key]['fingerprint'] != o.fingerprint() or o.key in sku_changes)]
        uncovered = [o for o in batch.orders if (not available.get((o.key, o.fingerprint())) or o.key in sku_changes)
                     and not prior.get(o.key, {}).get('removed_from_source_in')]
        fresh = [o for o in uncovered if o.key not in seen]
        old_outputs = sorted(name for name, output in history['outputs'].items()
                             if any(o.key in output['orders'] for o in batch.orders)
                             or any(s['path'] == key for s in output.get('sources', [])))
        missing_outputs = [name for name in old_outputs if not safe_child(folder / '000 Output', name).is_file()]
        ignored = bool(source and source.get('ignored_changes') and source.get('sha256') == batch.sha256)
        current_outputs = set()
        for order in batch.orders:
            fingerprint = prior.get(order.key, {}).get('fingerprint', order.fingerprint()) if ignored else order.fingerprint()
            current_outputs.update(available.get((order.key, fingerprint), []))
        output_formats = [fmt for fmt in ('pdf', 'docx') if any(Path(name).suffix.lower() == '.' + fmt for name in current_outputs)]
        format_label = ' + '.join('PDF' if fmt == 'pdf' else 'Word' for fmt in output_formats)
        old_text = None
        if batch.platform == 'SHOPEE':
            status, message = 'COURIER', 'Choose SPX or FLASH in Courier; Auto found no unique courier label in the filename or folder.'
        elif key in conflicts:
            status, message = 'ERROR', 'Conflicting versions of the same order are present in this dated folder. Keep the correct PDF and scan again.'
        elif (source and source.get('sha256') != batch.sha256) or (changed_orders and not ignored) or (ignored and changed_orders and missing_outputs):
            status, message = 'REVIEW', 'File or order details changed. Review the previous and current versions before continuing.'
            prior_lines = []
            keys = set(source.get('order_keys', []) if source else []) | {o.key for o in changed_orders}
            for oid in sorted(keys):
                if oid in prior:
                    saved = prior[oid]
                    prior_lines.append(f"Order {saved['order']['order_id']} — {saved['output']}")
                    for item in saved['order']['items']:
                        prior_lines.append(f"  {item['pieces']} pcs | {item['product']} | {item['variation']}" +
                                           (f" | SKU: {item['seller_sku']}" if item.get('seller_sku') else ''))
            old_text = '\r\n'.join(prior_lines) or 'An earlier source file is recorded in the processing log.'
        elif ignored and changed_orders:
            status, message = 'DONE', 'Earlier output kept. Review this changed PDF to generate its current contents.'
        elif fresh and (missing_outputs or any(o.key in prior for o in fresh) or (source and source.get('legacy_missing'))):
            status, message = 'MISSING', f'Output missing — {len(fresh)} orders without a PDF or Word file; ready to generate again'
        elif source and source.get('sha256') == batch.sha256 and not uncovered:
            status, message = 'DONE', 'Process again' if format_label else 'Earlier processing recorded — no current output'
        elif fresh:
            status, message = 'READY', f'{len(fresh)} new orders ready'
            skipped = len(batch.orders) - len(fresh)
            if skipped: message += f'; {skipped} other orders also included if selected'
        else:
            status, message = 'DUPLICATE', 'Repeated orders — select to generate again'
        if status == 'DUPLICATE' and not uncovered and format_label:
            message = 'Process again'
        elif status in ('READY', 'MISSING', 'DUPLICATE') and format_label:
            message += ' | Some orders already processed'
        if status in ('READY','MISSING','DONE','DUPLICATE'):
            for o in batch.orders:
                seen.setdefault(o.key, key)
        links = '\r\n'.join(('MISSING: ' if name in missing_outputs else 'output document: ') + name for name in old_outputs)
        entries.append({'path': key, 'platform': batch.platform, 'is_shopee': batch.platform in ('SHOPEE', 'SPX', 'FLASH'),
                        'courier': modes.get(key, 'Auto'), 'status': status, 'message': message,
                        'order_totals': [{'key': o.key, 'pieces': o.pieces} for o in batch.orders],
                        'output_formats': output_formats,
                        'details': (links + '\r\n\r\n' if links else '') + batch_details(batch, old_text), 'orders': len(batch.orders),
                        'pieces': batch.pieces, 'new_orders': len(fresh) if status == 'READY' else 0,
                        'missing_orders': len(fresh) if status == 'MISSING' else 0,
                        'requires_revision': status == 'REVIEW' or bool(ignored and changed_orders),
                        'sha256': batch.sha256, 'outputs': old_outputs, 'missing_outputs': missing_outputs})
    return entries, batches, history


def scan(folder, catalog_path=None):
    from sku_registry import load_registry
    registry = load_registry(catalog_path) if catalog_path is not None else None
    out = folder / '000 Output'
    if out.exists():
        with folder_lock(folder):
            recover(out)
            entries, batches, history = scan_internal(folder, registry)
            # Save associations while legacy documents still exist, so a later
            # deletion never destroys the PDF-to-output mapping.
            if (history['outputs'] or history.get('archived_outputs')) and (not (out / HISTORY).exists() or
                    json.loads((out / HISTORY).read_text(encoding='utf-8')) != history):
                atomic_json(out / HISTORY, history)
    else:
        entries, batches, history = scan_internal(folder, registry)
    from product_attention import catalog_attention, source_attention, source_definition_ids
    catalog, source_issues, scan_ids = {'rows': []}, [], []
    if catalog_path is not None:
        from names import encounter
        catalog = encounter(catalog_path, batches.values(), registry=registry)
        source_issues = source_attention(catalog, batches.values(), registry)
        scan_ids = source_definition_ids(catalog, batches.values(), registry)
    scan_id_set = set(scan_ids)
    scan_catalog = {'rows': [row for row in catalog['rows'] if row['id'] in scan_id_set]}
    return {'ok': True, 'entries': entries, 'folder': str(folder), 'output_folder': str(out),
            'outputs': completed_outputs(folder, history), 'version': VERSION,
            'name_attention': catalog_attention(catalog, source_issues), 'name_source_issues': source_issues,
            'name_scan_attention': catalog_attention(scan_catalog, source_issues), 'name_scan_ids': scan_ids}


def repeat_warning(selected, batches, history, output_format):
    selected_keys = {o.key for s in selected for o in batches[s['path'].replace('\\', '/')].orders}
    recorded = set(history['orders'])
    for source in history['sources'].values():
        recorded.update(source.get('order_keys', []))
    occurrences = defaultdict(int)
    for batch in batches.values():
        for order in batch.orders:
            occurrences[order.key] += 1
    previous = sorted(selected_keys & recorded)
    overlapping = sorted(k for k in selected_keys if occurrences[k] > 1)
    repeated = sorted(set(previous) | set(overlapping))
    payload = {'selected': selected, 'repeated': repeated, 'previous': previous,
               'overlapping': overlapping, 'format': output_format}
    token = hashlib.sha256(json.dumps(payload, sort_keys=True).encode('utf-8')).hexdigest()
    return {'ok': True, 'confirmation_required': True, 'repeat_count': len(repeated),
            'previously_generated': len(previous), 'overlapping_orders': len(overlapping),
            'repeat_token': token}


def generate(folder, selected, verifier=verify_word, catalog_path=None, output_format='docx', repeat_confirmation=None, repeat_action='include'):
    if output_format not in ('pdf', 'docx'):
        raise PackingError('Choose PDF or Word output.')
    if repeat_action not in ('include', 'skip'):
        raise PackingError('Choose whether to include or skip already-generated orders.')
    if not selected:
        raise PackingError('Select at least one PDF to generate.')
    from sku_registry import load_registry
    registry = load_registry(catalog_path) if catalog_path is not None else None
    with folder_lock(folder) as out:
        recover(out)
        entries, batches, history = scan_internal(folder, registry)
        from names import encounter, apply_names
        entry_map = {e['path']: e for e in entries}
        chosen, seen = [], set()
        modes, skipped = {}, {}
        selected_batches = {s['path'].replace('\\', '/'): batches[s['path'].replace('\\', '/')]
                            for s in selected if s['path'].replace('\\', '/') in batches}
        priorities = {s['path'].replace('\\', '/'): int(bool(s.get('revision'))) for s in selected}
        preferred = preferred_order_sources(selected_batches, registry, priorities)
        for selection in sorted(selected, key=lambda s: not bool(s.get('revision'))):
            key = selection['path'].replace('\\','/')
            entry = entry_map.get(key)
            if not entry or selection.get('sha256') != entry['sha256']:
                raise PackingError(f'{key} changed after the scan. Scan again before generating.')
            if selection.get('courier', entry['courier']) != entry['courier'] or selection.get('platform', entry['platform']) != entry['platform']:
                raise PackingError(f'{key}: courier changed after the scan. Scan again before generating.')
            revision = bool(selection.get('revision'))
            if entry.get('requires_revision') and not revision:
                raise PackingError(f'{key} needs an explicit revision review.')
            if entry['status'] not in ('READY','MISSING','DONE','REVIEW','DUPLICATE'):
                raise PackingError(f'{key} is not ready to generate: {entry["message"]}')
            batch = batches[key]
            orders = []
            skipped[key] = []
            for o in batch.orders:
                if o.key in seen or preferred.get(o.key) != key:
                    skipped[key].append(o.order_id)
                else:
                    orders.append(o)
                    seen.add(o.key)
            if orders:
                chosen.append((key, replace(batch, orders=orders)))
                modes[key] = 'revision' if revision else 'standard'
        if not chosen:
            raise PackingError('No readable orders remain in the selected PDFs.')
        warning = repeat_warning(selected, batches, history, output_format)
        if warning['repeat_count'] and repeat_confirmation != warning['repeat_token']:
            return warning
        skipped_existing = set()
        if repeat_action == 'skip':
            available = available_order_outputs(folder, history)
            filtered = []
            for key, batch in chosen:
                remaining = []
                for order in batch.orders:
                    saved = history['orders'].get(order.key, {}).get('order', {})
                    if available.get((order.key, order.fingerprint())) and not conflicting_order_skus([saved, order], registry):
                        skipped_existing.add(order.key)
                    else:
                        remaining.append(order)
                if remaining:
                    filtered.append((key, replace(batch, orders=remaining)))
            chosen = filtered
            if not chosen:
                return {'ok': True, 'outputs': [], 'skipped_repeat_orders': len(skipped_existing),
                        'message': 'No new orders to generate. All selected orders already have generated documents.'}
        catalog = encounter(catalog_path, batches.values(), registry=registry) if catalog_path is not None else {'revision': 0, 'rows': []}
        grouped = defaultdict(list)
        for key, batch in chosen:
            # Revisions have a distinct file from genuinely new orders.
            grouped[(batch.group, modes[key])].append((key, batch))
        stage = out / INTERNAL / ('.pending-' + uuid.uuid4().hex)
        stage.mkdir()
        files, plans, reserved = [], [], set()
        new_history = copy.deepcopy(history)
        try:
            for (group, mode), values in sorted(grouped.items()):
                if mode == 'standard':
                    mode = 'regenerated' if any(entry_map[key]['status'] == 'MISSING' or any(o.key in history['orders'] for o in batch.orders)
                                                for key, batch in values) else 'new'
                filename = next_output_filename(out, group, folder.name, reserved, output_format)
                reserved.add(filename)
                dest = stage / filename
                selected_batches = apply_names([batch for _, batch in values], catalog, registry=registry) if catalog_path is not None else [batch for _, batch in values]
                build_dest = dest.with_suffix('.docx')
                build_document(selected_batches, group, folder.name, build_dest, revision=(mode=='revision'))
                plans.append({'filename': filename, 'group': group, 'mode': mode, 'values': values, 'dest': dest,
                              'display_batches': selected_batches, 'build_dest': build_dest})
            paths = [p['build_dest'] for p in plans]
            verification = verifier(paths, export_pdf=True) if output_format == 'pdf' else verifier(paths)
            page_counts = {Path(v['path']).with_suffix('.' + output_format).name: v['pages'] for v in verification}
            if output_format == 'pdf':
                import pdfplumber
                for plan in plans:
                    if not plan['dest'].is_file():
                        raise PackingError('PDF export did not produce a file. No completed batch was recorded.')
                    with pdfplumber.open(plan['dest']) as pdf:
                        if not pdf.pages or len(pdf.pages) != page_counts.get(plan['filename']):
                            raise PackingError('PDF export page count does not match the checked document.')
                    plan['build_dest'].unlink()

            for selection in selected:
                key = selection['path'].replace('\\','/')
                if file_hash(batches[key].path) != batches[key].sha256:
                    raise PackingError(f'{key} changed while the documents were generated. Scan again.')
            when = datetime.now().isoformat(timespec='seconds')
            lines = [f'COMPLETED BY PACKING LIST MANAGER {VERSION}', f'Time: {when}',
                     f'Product name definitions revision: {catalog["revision"]}', f'Output format: {output_format.upper()}',
                     'Repeated orders within this run appear once.',
                     f'Already-generated orders skipped: {len(skipped_existing)}',
                     'Selected PDFs: ' + ', '.join(s['path'] for s in selected)]
            if history.get('legacy'):
                lines.append('Earlier PDF-to-document associations retained from the processing log and available documents; source fingerprints saved in Processing_History.json.')
                new_history['legacy'] = False
            for plan in plans:
                name, mode, values = plan['filename'], plan['mode'], plan['values']
                record_id = uuid.uuid4().hex
                if name in new_history['outputs']:
                    # A deleted suffix may be reused for unrelated PDFs. Retain
                    # its old manifest separately so it cannot cover new files.
                    previous = new_history['outputs'].pop(name)
                    new_history.setdefault('archived_outputs', []).append(
                        {'filename': name, **previous, 'archived_at': when, 'reason': 'Missing filename reused'})
                    lines.append(f'{name}: previous missing-file record archived; filename reused for this generation.')
                keys = [o.key for _, b in values for o in b.orders]
                replaced = defaultdict(set)
                removed = set()
                def record_replacements(order_key):
                    for old_name, old in new_history['outputs'].items():
                        if order_key in old['orders'] and (mode == 'revision' or
                                (plan['group'] == 'LAZADA_TIKTOK' and old['group'] == 'LAZADA')):
                            replaced[old_name].add(order_key)
                for key, batch in values:
                    if mode == 'revision':
                        previous_keys = set(history['sources'].get(key, {}).get('order_keys', []))
                        removed.update(previous_keys - {o.key for o in batches[key].orders})
                        for removed_key in previous_keys - {o.key for o in batches[key].orders}:
                            if removed_key in new_history['orders']:
                                prior = new_history['orders'][removed_key]
                                record_replacements(removed_key)
                                prior['removed_from_source_in'] = name
                    for o in batch.orders:
                        record_replacements(o.key)
                        new_history['orders'][o.key] = {'fingerprint': o.fingerprint(), 'order': asdict(o), 'output': name,
                                                      'output_record_id': record_id, 'source': key}
                    original = batches[key]
                    new_history['sources'][key] = {'sha256': original.sha256, 'platform': original.platform,
                                                  'order_keys': [o.key for o in original.orders],
                                                  'outputs': sorted(set(history['sources'].get(key, {}).get('outputs', [])) | {name}),
                                                  'note': 'Completed revision' if mode=='revision' else 'Process again', 'created': when}
                for old_name, replaced_keys in replaced.items():
                    old = new_history['outputs'].get(old_name)
                    if old:
                        all_replaced = set(old.get('superseded_orders', [])) | replaced_keys
                        old['superseded_orders'] = sorted(all_replaced)
                        old['status'] = 'superseded' if set(old['orders']) <= all_replaced else 'partially superseded'
                        old.setdefault('replacements', []).append(name)
                        lines.append(f'{old_name}: {old["status"]}; replacement {name}; affected orders ' + ', '.join(k.split(':',1)[1] for k in sorted(replaced_keys)))
                fingerprints = {o.key: o.fingerprint() for _, b in values for o in b.orders}
                sources = []
                for selection in selected:
                    key = selection['path'].replace('\\','/')
                    batch = batches[key]
                    if batch.group != plan['group'] or bool(selection.get('revision')) != (mode == 'revision'):
                        continue
                    linked_keys = [o.key for o in batch.orders if fingerprints.get(o.key) == o.fingerprint()]
                    if linked_keys:
                        sources.append({'path': key, 'sha256': batch.sha256, 'order_keys': linked_keys})
                new_history['outputs'][name] = {'group': plan['group'], 'orders': keys, 'mode': mode, 'created': when,
                                               'status': 'active', 'format': output_format, 'pages': page_counts.get(name), 'removed_orders': sorted(removed),
                                               'order_fingerprints': fingerprints, 'sources': sources, 'record_id': record_id,
                                               'name_catalog_revision': catalog['revision'],
                                               'display_order_snapshots': {o.key: asdict(o) for b in plan['display_batches'] for o in b.orders},
                                               'order_snapshots': {o.key: asdict(o) for _, b in values for o in b.orders}}
                lines += ['', name, f'Generation record: {record_id}', f'Status: Completed and checked ({mode})',
                          f'{len(keys)} orders | {sum(len(o.items) for _,b in values for o in b.orders)} product rows | '
                          f'{sum(b.pieces for _,b in values)} pieces | {page_counts.get(name)} A4 pages', 'Source PDFs:']
                if removed:
                    lines.append('    Removed from revised source (retained in history for duplicate prevention): ' + ', '.join(k.split(':',1)[1] for k in sorted(removed)))
                for source in sources:
                    key = source['path']
                    lines.append(f'  - {key} | {len(source["order_keys"])} linked orders | SHA256 {source["sha256"]}')
                    lines.append('    Order IDs: ' + ', '.join(k.split(':', 1)[1] for k in source['order_keys']))
                    if skipped[key]: lines.append('    Shared orders printed once in this run: ' + ', '.join(skipped[key]))
                files.append({'target': name, 'stage': plan['dest'].relative_to(out).as_posix(), 'sha256': file_hash(plan['dest'])})
            for selection in selected:
                key = selection['path'].replace('\\','/')
                batch = batches[key]
                linked_outputs = [name for name, output in new_history['outputs'].items()
                                  if any(s['path'] == key for s in output.get('sources', []))]
                new_history['sources'][key] = {'sha256': batch.sha256, 'platform': batch.platform,
                                              'order_keys': [o.key for o in batch.orders], 'outputs': linked_outputs,
                                              'note': 'Process again', 'created': when}
            for key, source in new_history['sources'].items():
                source['outputs'] = [name for name, output in new_history['outputs'].items()
                                     if any(s['path'] == key for s in output.get('sources', []))]
            for entry in files:
                if safe_child(out, entry['target']).exists():
                    raise PackingError(f'{entry["target"]} appeared while generating. Scan again to choose the next available number.')
            commit(out, new_history, files, stage, '\n'.join(lines))
        except Exception:
            # Leave an existing commit journal intact for deterministic recovery.
            if not (out / JOURNAL).exists() and stage.exists():
                for p in stage.iterdir():
                    if p.is_file(): p.unlink()
                stage.rmdir()
            raise
        return {'ok': True, 'message': f'Created {len(plans)} packing document(s).',
                'skipped_repeat_orders': len(skipped_existing),
                'outputs': [{'path': str(out/p['filename']), 'filename': p['filename'],
                             'pages': page_counts.get(p['filename']), 'mode': p['mode']} for p in plans]}


def acknowledge(folder, selection, catalog_path=None):
    from sku_registry import load_registry
    registry = load_registry(catalog_path) if catalog_path is not None else None
    with folder_lock(folder) as out:
        recover(out)
        entries, batches, history = scan_internal(folder, registry)
        key = selection['path']
        entry = next((e for e in entries if e['path']==key), None)
        if not entry or not entry.get('requires_revision') or entry['sha256'] != selection.get('sha256'):
            raise PackingError('This file changed after review. Scan again.')
        batch = batches[key]
        history['sources'][key] = {'sha256': batch.sha256, 'platform': batch.platform,
                                  'order_keys': [o.key for o in batch.orders], 'outputs': entry['outputs'],
                                  'ignored_changes': True,
                                  'note': 'Changed file reviewed and ignored; use the earlier output'}
        stage = out / INTERNAL / ('.pending-' + uuid.uuid4().hex)
        stage.mkdir()
        commit(out, history, [], stage,
               f'CHANGED FILE REVIEW {datetime.now().isoformat(timespec="seconds")}\n'
               f'{key}\nSHA256 {batch.sha256}\nDecision: Keep earlier completed output; ignore this changed version. No new document generated.')
    return {'ok': True, 'message': 'Review saved. The earlier output remains the packing record.'}


def import_pdfs(folder, paths, carrier=None):
    folder.mkdir(parents=True, exist_ok=True)
    imported, skipped, errors = [], [], []
    for value in paths:
        source = Path(value).resolve()
        try:
            if source.suffix.lower() != '.pdf': raise PackingError('Only PDF files can be added.')
            batch = parse_pdf(source, carrier, allow_unassigned=True)
            target_folder = folder if batch.platform == 'SHOPEE' else folder / ('LAZADA & TIKTOK' if batch.group=='LAZADA_TIKTOK' else batch.group)
            target_folder.mkdir(exist_ok=True)
            matches = [p for p in target_folder.iterdir() if p.is_file() and p.suffix.lower()=='.pdf' and file_hash(p)==batch.sha256]
            if matches:
                skipped.append(source.name)
                continue
            target = target_folder / source.name
            n = 2
            while target.exists():
                target = target_folder / f'{source.stem}_import_{n}.pdf'
                n += 1
            temp = target_folder / ('.import-' + uuid.uuid4().hex)
            try:
                shutil.copyfile(source, temp)
                if file_hash(temp) != batch.sha256: raise PackingError('The source changed during import. Wait for the download to finish.')
                os.rename(temp, target)
            finally:
                if temp.exists(): temp.unlink()
            if carrier in ('SPX', 'FLASH') and batch.platform in ('SPX', 'FLASH'):
                with folder_lock(folder) as out:
                    store_courier(out, load_couriers(folder), relpath(target, folder), batch.sha256, carrier)
            imported.append(str(target))
        except Exception as exc:
            errors.append({'file': source.name, 'message': str(exc)})
    return {'ok': True, 'imported': imported, 'skipped': skipped, 'errors': errors}


def handle(request):
    action = request['action']
    from names import load_catalog, save_catalog, export_excel, import_excel, prepare_save
    source = Path(__file__).resolve()
    if source.parent.name == 'Source' and source.parent.parent.name == 'App Files':
        default_root = source.parents[2]
    else:
        default_root = Path(sys.executable).parent.parent if getattr(sys, 'frozen', False) else source.parents[2] / 'Packing List Manager'
    catalog_path = Path(request.get('catalog_path', default_root / 'ProductNames.json')).resolve()
    if action == 'names_list':
        from names import migrate_registry
        from product_attention import catalog_attention
        from sku_registry import load_registry
        registry = load_registry(catalog_path)
        catalog = migrate_registry(catalog_path, registry) if registry is not None else load_catalog(catalog_path)
        return {'ok': True, **catalog, 'name_attention': catalog_attention(catalog)}
    if action == 'names_prepare': return prepare_save(request['rows'], request.get('priority', []))
    if action == 'names_save': return {'ok': True, **save_catalog(catalog_path, request['rows'], request['revision'], request.get('merge_groups', []), lock_all=request.get('lock_all', False), unlocked_ids=request.get('unlocked_ids', []), review_source_issues=request.get('review_source_issues'))}
    if action == 'names_export': return export_excel(Path(request['path']), load_catalog(catalog_path))
    if action == 'names_import': return import_excel(Path(request['path']), request['rows'])
    folder = dated_folder(request['folder'])
    if action == 'courier': return set_courier(folder, request['selected'], request['carrier'])
    if action == 'scan': return scan(folder, catalog_path)
    if action == 'sales_export':
        from sales import export_sales
        return export_sales(folder, catalog_path)
    if action == 'generate':
        if request.get('output_format', 'pdf') != 'pdf':
            raise PackingError('Packing List Manager exports PDF only.')
        return generate(folder, request['selected'], catalog_path=catalog_path, output_format='pdf', repeat_confirmation=request.get('repeat_confirmation'), repeat_action=request.get('repeat_action', 'include'))
    if action == 'import': return import_pdfs(folder, request['files'], request.get('carrier'))
    if action == 'acknowledge': return acknowledge(folder, request['selected'], catalog_path)
    raise PackingError('Unknown action.')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--request', required=True)
    parser.add_argument('--response', required=True)
    args = parser.parse_args()
    try:
        request = json.loads(Path(args.request).read_text(encoding='utf-8-sig'))
        result = handle(request)
    except Exception as exc:
        result = {'ok': False, 'error': str(exc)}
    Path(args.response).write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')


if __name__ == '__main__':
    main()
