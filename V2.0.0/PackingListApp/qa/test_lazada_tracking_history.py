"""Tracking-aware packing without guessing or rewriting earlier audit records."""
import copy
from dataclasses import asdict
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
import engine
from extractors import Batch, Item, Order, PackingError, file_hash
from docx import Document
from generation_test_support import generate_approved
from test_regeneration import layout


class LazadaTrackingHistoryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.folder = Path(self.temp.name) / '20261007'
        (self.folder / 'LAZADA').mkdir(parents=True)
        self.batches = {}
        self.parser = patch.object(engine, 'parse_pdf', side_effect=lambda path, **kwargs: self.batches[Path(path).name])
        self.parser.start()

    def tearDown(self):
        self.parser.stop()
        self.temp.cleanup()

    def order(self, tracking='TRACK-A', pieces=2):
        return Order('LAZADA', '123456789', [Item('SOAP', 'LEMON', pieces)], tracking_id=tracking)

    def source(self, name, orders):
        path = self.folder / 'LAZADA' / name
        path.write_text(name + repr(orders), encoding='utf-8')
        self.batches[name] = Batch(str(path), 'LAZADA', file_hash(path), orders)
        return 'LAZADA/' + name

    def old_output(self, paths, order=None, source_hash=True):
        order = order or self.order('')
        key = 'LAZADA:' + order.order_id
        snapshot = asdict(order)
        snapshot.pop('tracking_id')
        sources = [{'path': path, 'sha256': self.batches[Path(path).name].sha256 if source_hash else None,
                    'order_keys': [key]} for path in paths]
        filename = 'LAZADA_TIKTOK_Packing_Checklist_2026-10-07_1.pdf'
        output = {'group': 'LAZADA_TIKTOK', 'orders': [key], 'mode': 'new', 'created': 'Earlier',
                  'status': 'active', 'format': 'pdf', 'record_id': 'legacy-test',
                  'order_fingerprints': {key: order.fingerprint()}, 'sources': sources,
                  'order_snapshots': {key: snapshot}, 'display_order_snapshots': {key: copy.deepcopy(snapshot)}}
        history = {'schema': 1, 'date': self.folder.name, 'legacy': False,
                   'orders': {key: {'fingerprint': order.fingerprint(), 'order': snapshot,
                                    'output': filename, 'source': paths[0]}},
                   'sources': {s['path']: {**copy.deepcopy(s), 'platform': 'LAZADA', 'outputs': [filename]} for s in sources},
                   'outputs': {filename: output}}
        out = self.folder / '000 Output'
        (out / engine.INTERNAL).mkdir(parents=True, exist_ok=True)
        (out / filename).write_bytes(b'previously completed document; snapshots already available')
        engine.atomic_json(out / engine.HISTORY, history)
        return filename, copy.deepcopy(output)

    def history(self):
        return json.loads((self.folder / '000 Output' / engine.HISTORY).read_text(encoding='utf-8'))

    def selection(self, entries, revision=False):
        return [{'path': e['path'], 'sha256': e['sha256'], 'revision': revision} for e in entries]

    def test_unchanged_source_proves_legacy_package_and_preserves_all_snapshots(self):
        path = self.source('one.pdf', [self.order()])
        filename, old = self.old_output([path])
        scanned = engine.scan(self.folder)
        self.assertEqual(scanned['entries'][0]['status'], 'DONE')
        current = self.history()['outputs'][filename]
        self.assertEqual(current.pop('tracking_order_keys'), {'LAZADA:123456789': self.order().key})
        self.assertEqual(current, old)
        entry = scanned['entries'][0]
        self.assertEqual(entry['order_totals'][0]['tracking_id'], 'TRACK-A')
        self.assertIn('Tracking ID: TRACK-A', entry['details'])

    def test_no_skips_legacy_covered_package_without_creating_output(self):
        path = self.source('one.pdf', [self.order()])
        self.old_output([path])
        selected = self.selection(engine.scan(self.folder)['entries'])
        warning = engine.generate(self.folder, selected, verifier=layout)
        self.assertEqual(warning['previously_generated'], 1)
        result = engine.generate(self.folder, selected, verifier=layout,
                                 repeat_confirmation=warning['repeat_token'], repeat_action='skip')
        self.assertEqual(result['outputs'], [])
        self.assertEqual(result['skipped_repeat_orders'], 1)

    def test_different_tracking_is_new_and_never_hidden_by_legacy_order_number(self):
        path = self.source('one.pdf', [self.order()])
        self.old_output([path])
        engine.scan(self.folder)
        self.source('two.pdf', [self.order('TRACK-B')])
        entries, batches, history = engine.scan_internal(self.folder)
        by_path = {e['path']: e for e in entries}
        self.assertEqual(by_path[path]['status'], 'DONE')
        new = by_path['LAZADA/two.pdf']
        self.assertEqual((new['status'], new['new_orders']), ('READY', 1))
        warning = engine.repeat_warning(self.selection([new]), batches, history, 'pdf')
        self.assertEqual(warning['repeat_count'], 0)

    def test_unverified_legacy_source_requires_review_even_with_identical_items(self):
        path = self.source('one.pdf', [self.order()])
        filename, _ = self.old_output([path], source_hash=False)
        entry = engine.scan(self.folder)['entries'][0]
        self.assertEqual(entry['status'], 'REVIEW')
        self.assertTrue(entry['requires_revision'])
        self.assertIn('Tracking ID', entry['message'])
        self.assertNotIn('tracking_order_keys', self.history()['outputs'][filename])

    def test_ambiguous_old_output_does_not_cover_multiple_tracking_numbers(self):
        first = self.source('one.pdf', [self.order()])
        second = self.source('two.pdf', [self.order('TRACK-B')])
        filename, _ = self.old_output([first, second])
        entries = engine.scan(self.folder)['entries']
        self.assertEqual([e['status'] for e in entries], ['REVIEW', 'REVIEW'])
        self.assertNotIn('tracking_order_keys', self.history()['outputs'][filename])

    def test_missing_or_changed_linked_copy_cannot_prove_legacy_tracking(self):
        first = self.source('one.pdf', [self.order()])
        second = self.source('two.pdf', [self.order('TRACK-B')])
        filename, _ = self.old_output([first, second])
        (self.folder / second).unlink()
        entries = engine.scan(self.folder)['entries']
        self.assertEqual(entries[0]['status'], 'REVIEW')
        self.assertNotIn('tracking_order_keys', self.history()['outputs'][filename])
        self.source('two.pdf', [self.order('TRACK-C')])
        entries = engine.scan(self.folder)['entries']
        self.assertEqual([e['status'] for e in entries], ['REVIEW', 'REVIEW'])
        self.assertNotIn('tracking_order_keys', self.history()['outputs'][filename])

    def test_explicit_revision_resolves_ambiguous_legacy_without_erasing_old_snapshots(self):
        first = self.source('one.pdf', [self.order()])
        second = self.source('two.pdf', [self.order('TRACK-B')])
        filename, original = self.old_output([first, second])
        entries = engine.scan(self.folder)['entries']
        result = generate_approved(self.folder, self.selection(entries, revision=True), verifier=layout)
        history = self.history()
        old = history['outputs'][filename]
        self.assertEqual(old['status'], 'superseded')
        self.assertEqual(old['order_snapshots'], original['order_snapshots'])
        self.assertEqual(old['display_order_snapshots'], original['display_order_snapshots'])
        new = history['outputs'][result['outputs'][0]['filename']]
        self.assertEqual(set(new['orders']), {self.order().key, self.order('TRACK-B').key})
        self.assertEqual([e['status'] for e in engine.scan(self.folder)['entries']], ['DONE', 'DONE'])

    def test_deleted_legacy_output_clears_verified_tracking_coverage(self):
        path = self.source('one.pdf', [self.order()])
        filename, original = self.old_output([path])
        engine.scan(self.folder)
        (self.folder / '000 Output' / filename).unlink()
        entry = engine.scan(self.folder)['entries'][0]
        self.assertEqual((entry['status'], entry['new_orders']), ('READY', 1))
        history = self.history()
        self.assertEqual(history['orders'], {})
        self.assertEqual(history['archived_outputs'][0]['order_snapshots'], original['order_snapshots'])

    def test_revision_removes_old_verified_tracking_coverage(self):
        path = self.source('one.pdf', [self.order()])
        filename, _ = self.old_output([path])
        engine.scan(self.folder)
        self.source('one.pdf', [self.order('TRACK-B')])
        entries = engine.scan(self.folder)['entries']
        self.assertEqual(entries[0]['status'], 'REVIEW')
        result = generate_approved(self.folder, self.selection(entries, revision=True), verifier=layout)
        history = self.history()
        self.assertEqual(history['outputs'][filename]['status'], 'superseded')
        self.assertEqual(history['orders'][self.order().key]['removed_from_source_in'], result['outputs'][0]['filename'])
        self.assertNotIn((self.order().key, self.order().fingerprint()), engine.available_order_outputs(self.folder, history))

    def test_same_order_different_tracking_are_both_generated_and_counted(self):
        self.source('split.pdf', [self.order(), self.order('TRACK-B', 3)])
        entry = engine.scan(self.folder)['entries'][0]
        self.assertEqual((entry['orders'], entry['shipments'], entry['pieces']), (1, 2, 5))
        self.assertIn('1 orders | 2 shipments', entry['details'])
        result = engine.generate(self.folder, self.selection([entry]), verifier=layout)
        output = self.history()['outputs'][result['outputs'][0]['filename']]
        self.assertEqual(len(output['orders']), 2)
        self.assertEqual({o['tracking_id'] for o in output['order_snapshots'].values()}, {'TRACK-A', 'TRACK-B'})
        self.assertEqual({o['tracking_id'] for o in output['display_order_snapshots'].values()}, {'TRACK-A', 'TRACK-B'})
        generated = Path(result['outputs'][0]['path'])
        self.assertEqual(set(engine.read_completed_doc(generated, 'LAZADA_TIKTOK')), {self.order().key, self.order('TRACK-B').key})
        duplicated = Document(generated)
        duplicated.tables[0]._tbl.addnext(copy.deepcopy(duplicated.tables[0]._tbl))
        invalid = generated.with_name('duplicate-tracking.docx')
        duplicated.save(invalid)
        with self.assertRaisesRegex(PackingError, 'repeats order'):
            engine.read_completed_doc(invalid, 'LAZADA_TIKTOK')

    def test_mixed_missing_and_known_tracking_is_rejected_without_double_counting(self):
        self.source('old.pdf', [self.order('')])
        self.source('current.pdf', [self.order()])
        entries = engine.scan(self.folder)['entries']
        self.assertEqual([e['status'] for e in entries], ['ERROR', 'ERROR'])
        self.assertTrue(all('with and without a Tracking ID' in e['message'] for e in entries))


if __name__ == '__main__':
    unittest.main()
