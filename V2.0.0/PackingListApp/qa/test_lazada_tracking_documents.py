"""Lazada split shipments retain distinct tracking numbers in packing tables."""
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
from documents import build_document
from extractors import Batch, Item, Order


class LazadaTrackingDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.path = Path(self.temp.name) / 'tracking.docx'

    def tearDown(self):
        self.temp.cleanup()

    def build(self, platform, orders):
        batch = Batch('source.pdf', platform, 'source-hash', orders)
        build_document([batch], batch.group, '20261007', self.path)
        return Document(self.path)

    def test_same_order_number_has_separate_tracking_headers_and_all_pieces(self):
        orders = [Order('LAZADA', '123456789012345', [Item('PRODUCT', '50 ML', qty)],
                        tracking_id=tracking)
                  for tracking, qty in (('LAZTRACK001', 2), ('LAZTRACK002', 3))]
        document = self.build('LAZADA', orders)
        self.assertEqual(len(document.tables), 3)
        original = Document(APP / 'src/assets/blue.docx').tables[0]
        for table, order in zip(document.tables, orders):
            header = table.cell(0, 0).text
            self.assertEqual(header.count('Order Number: 123456789012345'), 1)
            self.assertIn('\tEmployee:__________________\tNumber __________\tTime __________', header)
            self.assertTrue(header.endswith('\nTracking Number: ' + order.tracking_id))
            self.assertEqual(len(table.columns), 4)
            self.assertEqual(sum(c.width for c in table.columns),
                             sum(c.width for c in original.columns))
            self.assertEqual(table.cell(0, 0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')),
                             original.cell(0, 0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')))
            self.assertTrue(table.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next)
            self.assertTrue(table.cell(0, 0).paragraphs[0].paragraph_format.keep_together)
            self.assertEqual(table.rows[2].cells[3].text, str(order.pieces))
        summary = document.tables[-1]
        self.assertEqual([c.text for c in summary.rows[1].cells], ['[     ]', 'PRODUCT', '50 ML', '5'])
        self.assertEqual(summary.rows[-1].cells[3].text, '5')
        paragraphs = [p.text for p in document.paragraphs]
        self.assertTrue(any(p.startswith('One table is one shipment.') for p in paragraphs))
        self.assertIn('Total orders: 1 | Shipments: 2', paragraphs)
        self.assertTrue(any('1 orders  |  2 shipments  |  5 pieces' in p for p in paragraphs))

    def test_legacy_lazada_without_tracking_keeps_existing_header(self):
        document = self.build('LAZADA', [Order('LAZADA', 'LEGACY123', [Item('A', '1 L', 1)])])
        self.assertTrue(any(p.text.startswith('One table is one order.') for p in document.paragraphs))
        self.assertEqual(document.tables[0].cell(0, 0).text,
                         'Order Number: LEGACY123\tEmployee:__________________\tNumber __________\tTime __________')

    def test_tracking_header_is_lazada_only(self):
        for platform in ('TIKTOK', 'SPX', 'FLASH'):
            with self.subTest(platform=platform):
                document = self.build(platform, [Order(platform, 'ORDER123', [Item('A', '1 L', 1)],
                                                       tracking_id='TRACK123')])
                self.assertEqual(document.tables[0].cell(0, 0).text,
                                 'Order ID: ORDER123\tEmployee:__________________\tNumber __________\tTime __________')


if __name__ == '__main__':
    unittest.main(verbosity=2)
