"""Product reminders remain advisory and preserve independent print fallbacks."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
import engine
import names
from extractors import Batch, Item, Order, file_hash
from product_attention import catalog_attention, remap_attention_ids, row_attention_reasons, source_attention, source_definition_ids
from sku_registry import SkuRegistry


def definition(identity='known', sku='PERF-A-30ML', product='INTERNAL A', size='30 ML'):
    return {'id': identity, 'sku': sku, 'sku_scope': 'CATALOG', 'sku_aliases': [],
            'original_product': 'ECOMMERCE A', 'original_variation': 'SMALL',
            'alternative_product': product, 'alternative_variation': size, 'locked': True}


def batch(sku='', product='ECOMMERCE A', variation='SMALL', platform='SPX'):
    return Batch('packing.pdf', platform, 'hash', [Order(platform, 'ORDER-1', [
        Item(product, variation, 2, '45.00', '89.00', sku, sku)])])


def registry_data():
    return {'schema': 1, 'records': [{
        'sku': 'PERF-A-30ML', 'original_product': 'ECOMMERCE A', 'original_variation': 'SMALL',
        'sku_aliases': [{'platform': 'SHOPEE', 'sku': 'PERF-A-30ML'},
                        {'platform': 'LAZADA', 'sku': '123456'}],
        'ecommerce_sources': [{'platform': 'SHOPEE', 'sku': 'PERF-A-30ML',
                              'product': 'ECOMMERCE A', 'variation': 'SMALL'}]}]}


class ProductAttentionTests(unittest.TestCase):
    def test_complete_partial_blank_and_unusable_skus(self):
        self.assertEqual(row_attention_reasons(definition()), [])
        self.assertEqual(row_attention_reasons(definition(size=' ')), ['missing_size'])
        self.assertEqual(row_attention_reasons(definition(product='')), ['missing_product_name'])
        self.assertEqual(row_attention_reasons(definition(sku='', product='', size='')),
                         ['missing_sku', 'missing_product_name', 'missing_size'])
        for sku in ('30 ML', 'N/A', '—', '123456'):
            with self.subTest(sku=sku):
                self.assertEqual(row_attention_reasons(definition(sku=sku)), ['missing_sku'])

    def test_numeric_source_provenance_and_unknown_structured_codes_remain_valid(self):
        row = definition(sku='123456')
        row['sku_aliases'] = [{'platform': 'LAZADA', 'sku': '123456'}]
        self.assertEqual(row_attention_reasons(row), [])
        self.assertEqual(row_attention_reasons(definition(sku='NEW-CODE', product='')),
                         ['missing_product_name'])

    def test_counts_are_distinct_rows_and_reports_do_not_modify_catalog(self):
        rows = [definition('one', sku='', product='', size=''), definition('two', size=''), definition('three')]
        catalog = {'rows': rows}
        before = copy.deepcopy(catalog)
        result = catalog_attention(catalog, [{'id': 'one', 'reasons': ['missing_sku'], 'sources': []}] * 2)
        self.assertEqual((result['count'], result['missing_sku_count'], result['missing_name_count']), (2, 1, 2))
        self.assertEqual(catalog, before)

    def test_missing_pdf_sku_is_reported_even_when_registry_and_saved_names_are_complete(self):
        registry = SkuRegistry(registry_data())
        catalog = {'rows': [definition()]}
        value = batch()
        original = copy.deepcopy(value.to_dict())
        issues = source_attention(catalog, [value], registry)
        self.assertEqual([(entry['id'], entry['reasons']) for entry in issues], [('known', ['missing_sku'])])
        result = catalog_attention(catalog, issues)
        self.assertEqual((result['count'], result['missing_sku_count'], result['missing_name_count']), (1, 1, 0))
        self.assertEqual(result['items'][0]['sources'][0]['order_id'], 'ORDER-1')
        self.assertEqual(names.apply_names([value], catalog, registry)[0].orders[0].items[0].product, 'INTERNAL A')
        self.assertEqual(value.to_dict(), original)

    def test_generic_size_is_reported_but_registered_numeric_and_unknown_pdf_codes_are_valid(self):
        registry = SkuRegistry(registry_data())
        row = definition()
        row['sku_aliases'] = registry.entries[0]['sku_aliases']
        catalog = {'rows': [row]}
        self.assertEqual(source_attention(catalog, [batch('123456', platform='LAZADA')], registry), [])
        self.assertEqual(source_attention(catalog, [batch('998877', platform='LAZADA')], registry), [])
        self.assertEqual(source_attention(catalog, [batch('UNKNOWN-CODE')], registry), [])
        generic = source_attention(catalog, [batch('30ML', platform='TIKTOK')], registry)
        self.assertEqual(generic[0]['id'], 'known')

    def test_current_source_report_clears_when_pdf_has_usable_sku(self):
        catalog = {'rows': [definition()]}
        registry = SkuRegistry(registry_data())
        self.assertEqual(catalog_attention(catalog, source_attention(catalog, [batch()], registry))['count'], 1)
        self.assertEqual(catalog_attention(catalog, source_attention(catalog, [batch('PERF-A-30ML')], registry))['count'], 0)

    def test_saved_names_clear_only_name_reasons_and_keep_current_missing_sku_reminder(self):
        catalog = {'rows': [definition(product='', size='')]}
        issues = source_attention(catalog, [batch()], SkuRegistry(registry_data()))
        self.assertEqual(catalog_attention(catalog, issues)['missing_name_count'], 1)
        catalog['rows'][0].update(alternative_product='STANDARD', alternative_variation='30 ML')
        report = catalog_attention(catalog, issues)
        self.assertEqual(report['items'][0]['reasons'], ['missing_sku'])

    def test_confirmed_merge_retains_source_issue_but_deleted_or_other_platform_row_does_not(self):
        old = definition('removed')
        issue = dict(old, reasons=['missing_sku'], sources=[])
        kept = definition('kept')
        self.assertEqual(catalog_attention({'rows': [kept]}, [issue])['items'][0]['id'], 'kept')
        self.assertEqual(catalog_attention({'rows': []}, [issue])['count'], 0)
        other = definition('other')
        other['sku_scope'] = 'TIKTOK'
        self.assertEqual(catalog_attention({'rows': [other]}, [issue])['count'], 0)

    def test_attention_scope_follows_merged_distinct_sku_and_linked_name_without_mutating_catalogs(self):
        base = definition('base', sku='PERF-BASE')
        sku_row = definition('sku-row', sku='PERF-MERGED')
        legacy = definition('legacy', sku='')
        legacy.update(original_product='OLD NAME', original_variation='OLD SIZE')
        previous = {'rows': [base, sku_row, legacy]}
        kept = copy.deepcopy(base)
        kept['sku_aliases'] = [{'platform': '', 'sku': 'PERF-MERGED', 'canonical_sku': 'PERF-MERGED'}]
        kept['aliases'] = [{'original_product': 'OLD NAME', 'original_variation': 'OLD SIZE'}]
        current = {'rows': [kept]}
        originals = copy.deepcopy((previous, current))
        self.assertEqual(remap_attention_ids(previous, current, {'sku-row', 'legacy'}), {'base'})
        self.assertEqual((previous, current), originals)

    def test_attention_scope_follows_platform_alias_but_not_same_code_on_unrelated_platform(self):
        removed = definition('removed', sku='PLATFORM-CODE')
        removed['sku_scope'] = 'TIKTOK'
        previous = {'rows': [removed]}
        unrelated = definition('unrelated', sku='PLATFORM-CODE')
        unrelated['sku_scope'] = 'LAZADA'
        self.assertEqual(remap_attention_ids(previous, {'rows': [unrelated]}, {'removed'}), set())
        kept = definition('base')
        kept['sku_aliases'] = [{'platform': 'TIKTOK', 'sku': 'PLATFORM-CODE'}]
        self.assertEqual(remap_attention_ids(previous, {'rows': [unrelated, kept]}, {'removed'}), {'base'})

    def test_attention_scope_drops_deleted_or_ambiguous_rows_and_keeps_surviving_ids(self):
        removed = definition('removed', sku='')
        survivor = definition('survivor')
        survivor['original_product'] = 'SURVIVING TITLE'
        previous = {'rows': [removed, survivor]}
        replacement = definition('unrelated', sku='OTHER')
        replacement.update(original_product='OTHER NAME', original_variation='OTHER SIZE')
        self.assertEqual(remap_attention_ids(previous, {'rows': [replacement, survivor]}, {'removed', 'survivor'}), {'survivor'})
        self.assertEqual(remap_attention_ids(previous, {'rows': [replacement, survivor]}, {'removed'}), set())
        duplicates = [definition('duplicate-one'), definition('duplicate-two')]
        self.assertEqual(remap_attention_ids(previous, {'rows': duplicates}, {'removed'}), set())
        self.assertEqual(remap_attention_ids(previous, {'rows': []}, {'removed'}), set())

    def test_source_matching_retains_linked_original_and_whitespace_semantics(self):
        row = definition(sku='')
        row['aliases'] = [{'original_product': 'OLDER TITLE', 'original_variation': '30 ML'}]
        issues = source_attention({'rows': [row]}, [batch(product='OLDER\n TITLE', variation='3 0 ML')])
        self.assertEqual(issues[0]['id'], 'known')
        self.assertEqual(source_definition_ids({'rows': [row]}, [batch(product='OLDER\n TITLE', variation='3 0 ML')]), ['known'])

    def test_blank_fields_fall_back_independently_without_mutating_source_or_fingerprint(self):
        source = batch('PERF-A-30ML')
        original = copy.deepcopy(source.to_dict())
        fingerprint = source.orders[0].fingerprint()
        for product, size, expected in (('', '', ('ECOMMERCE A', 'SMALL')),
                                        ('STANDARD', '', ('STANDARD', 'SMALL')),
                                        ('', '30 ML', ('ECOMMERCE A', '30 ML'))):
            catalog = {'rows': [definition(product=product, size=size)]}
            catalog_attention(catalog)
            item = names.apply_names([source], catalog, SkuRegistry(registry_data()))[0].orders[0].items[0]
            self.assertEqual((item.product, item.variation), expected)
        self.assertEqual(source.to_dict(), original)
        self.assertEqual(source.orders[0].fingerprint(), fingerprint)


class ProductAttentionWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.root = Path(self.temp.name)
        self.folder = self.root / '20261010'
        (self.folder / 'SPX').mkdir(parents=True)
        self.catalog_path = self.root / 'ProductNames.json'
        self.source = self.folder / 'SPX' / 'packing.pdf'
        self.source.write_bytes(b'stable-source-fixture')
        self.value = batch()
        self.value.path, self.value.sha256 = str(self.source), file_hash(self.source)
        self.parser = patch.object(engine, 'parse_pdf', return_value=self.value)
        self.parser.start()

    def tearDown(self):
        self.parser.stop()
        self.temp.cleanup()

    def test_scan_lists_missing_names_without_changing_ready_status_or_generation(self):
        from docx import Document
        from reportlab.pdfgen.canvas import Canvas

        def export_fixture(paths, export_pdf=False):
            self.assertTrue(export_pdf)
            for path in paths:
                document = Document(path)
                self.assertEqual(document.tables[0].cell(2, 1).text, 'ECOMMERCE A')
                self.assertEqual(document.tables[0].cell(2, 2).text, 'SMALL')
                pdf = Canvas(str(Path(path).with_suffix('.pdf')), pagesize=(595.276, 841.890))
                pdf.drawString(30, 800, 'Test export; real renderer has separate layout checks.')
                pdf.save()
            return [{'path': str(path), 'pages': 1} for path in paths]

        scanned = engine.scan(self.folder, self.catalog_path)
        self.assertEqual(scanned['entries'][0]['status'], 'READY')
        self.assertEqual(scanned['name_attention']['count'], 1)
        self.assertEqual(scanned['name_attention']['items'][0]['reasons'],
                         ['missing_sku', 'missing_product_name', 'missing_size'])
        self.assertEqual(len(scanned['name_source_issues']), 1)
        listed = engine.handle({'action': 'names_list', 'catalog_path': str(self.catalog_path)})
        self.assertEqual(listed['name_attention']['count'], 1)
        selected = [{'path': entry['path'], 'sha256': entry['sha256']} for entry in scanned['entries']]
        result = engine.generate(self.folder, selected, catalog_path=self.catalog_path,
                                 output_format='pdf', verifier=export_fixture)
        self.assertTrue(result['ok'])
        self.assertEqual(len(result['outputs']), 1)
        self.assertTrue(Path(result['outputs'][0]['path']).exists())
        history = json.loads((self.folder / '000 Output' / engine.HISTORY).read_text())
        self.assertEqual(history['orders']['SHOPEE:ORDER-1']['order']['items'][0]['seller_sku'], '')
        printed = next(iter(history['outputs'].values()))['display_order_snapshots']['SHOPEE:ORDER-1']['items'][0]
        self.assertEqual((printed['product'], printed['variation']), ('ECOMMERCE A', 'SMALL'))

    def test_scan_reports_missing_pdf_sku_after_catalog_resolves_complete_definition(self):
        self.catalog_path.with_name('SkuCatalog.json').write_text(json.dumps(registry_data()))
        names.save_catalog(self.catalog_path, [definition()], 0)
        scanned = engine.scan(self.folder, self.catalog_path)
        self.assertEqual(scanned['name_attention']['items'][0]['reasons'], ['missing_sku'])
        self.assertEqual(scanned['entries'][0]['status'], 'READY')
        self.assertEqual(engine.handle({'action': 'names_list', 'catalog_path': str(self.catalog_path)})['name_attention']['count'], 0)

    def test_unknown_sku_with_existing_ecommerce_name_gets_its_own_undefined_report(self):
        self.catalog_path.with_name('SkuCatalog.json').write_text(json.dumps(registry_data()))
        names.save_catalog(self.catalog_path, [definition()], 0)
        self.value.orders[0].items[0].seller_sku = 'NEW-CODE'
        scanned = engine.scan(self.folder, self.catalog_path)
        report = scanned['name_attention']
        self.assertEqual(report['count'], 1)
        self.assertEqual(report['items'][0]['sku'], 'NEW-CODE')
        self.assertEqual(report['items'][0]['reasons'], ['missing_product_name', 'missing_size'])
        self.assertEqual(scanned['name_source_issues'], [])

    def test_immediate_scan_attention_omits_unencountered_catalog_deficiencies(self):
        self.catalog_path.with_name('SkuCatalog.json').write_text(json.dumps(registry_data()))
        unused = definition('unused', sku='UNUSED-SKU', product='', size='')
        unused.update(original_product='UNUSED TITLE', original_variation='LARGE')
        names.save_catalog(self.catalog_path, [definition(), unused], 0)
        self.value.orders[0].items[0].seller_sku = 'PERF-A-30ML'
        scanned = engine.scan(self.folder, self.catalog_path)
        self.assertEqual(scanned['name_attention']['count'], 1)
        self.assertEqual(scanned['name_scan_attention']['count'], 0)
        self.assertEqual(scanned['name_scan_ids'], ['known'])

    def test_review_locks_persist_through_save_completion_and_excel_export(self):
        from openpyxl import load_workbook
        from products_ui import ProductsModel

        rows = [definition('complete', 'PERF-COMPLETE-30ML', 'COMPLETE', '30 ML'),
                definition('blank', 'TEST-AURORA-50ML', '', ''),
                definition('partial', 'PERF-PARTIAL-30ML', 'PARTIAL', ''),
                definition('no-sku', '', '', ''),
                definition('source', 'PERF-SOURCE-30ML', 'SOURCE', '30 ML')]
        for row in rows:
            row['original_product'] = 'ECOMMERCE ' + row['id'].upper()
        initial = names.save_catalog(self.catalog_path, rows, 0, lock_all=True)
        issues = [dict(rows[-1], reasons=['missing_sku'], sources=[])]
        model = ProductsModel(initial, issues)
        self.assertEqual(model.unlocked_ids, {'blank', 'partial', 'no-sku', 'source'})

        def save(lock_all=True):
            result = engine.handle({'action': 'names_save', 'catalog_path': str(self.catalog_path),
                                    'rows': model.payload(), 'revision': model.revision,
                                    'lock_all': lock_all, 'unlocked_ids': sorted(model.unlocked_ids),
                                    'review_source_issues': issues})
            self.assertTrue(result['ok'])
            persisted = names.load_catalog(self.catalog_path)
            self.assertEqual(result['rows'], persisted['rows'])
            model.load(result)
            return {row['id']: row for row in persisted['rows']}

        persisted = save()
        self.assertEqual({key: row['locked'] for key, row in persisted.items()},
                         {'complete': True, 'blank': False, 'partial': False,
                          'no-sku': False, 'source': False})
        model.edit('blank', {'alternative_product': 'AURORA PEARL', 'alternative_variation': '50 ML'})
        persisted = save()
        self.assertTrue(persisted['blank']['locked'])
        self.assertNotIn('blank', {item['id'] for item in catalog_attention(
            {'rows': list(persisted.values())}, issues)['items']})
        self.assertFalse(persisted['source']['locked'])

        model.set_lock('complete', False)
        persisted = save(lock_all=False)
        self.assertFalse(persisted['complete']['locked'])
        self.assertTrue(persisted['blank']['locked'])
        workbook_path = self.root / 'review-locks.xlsx'
        engine.handle({'action': 'names_export', 'catalog_path': str(self.catalog_path),
                       'path': str(workbook_path)})
        workbook = load_workbook(workbook_path)
        try:
            worksheet = workbook['Names']
            self.assertTrue(worksheet.protection.sheet)
            for cells in worksheet.iter_rows(min_row=2):
                expected = persisted[cells[4].value]['locked']
                self.assertEqual([cell.protection.locked for cell in cells[:4]], [expected] * 4)
        finally:
            workbook.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
