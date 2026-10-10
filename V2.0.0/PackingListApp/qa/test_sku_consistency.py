"""SKU evidence supplements immutable fingerprints before deduplication."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
import engine
import names
import sales
from extractors import Batch, Item, Order, PackingError, file_hash
from generation_test_support import generate_approved
from sku_registry import SkuRegistry
from test_regeneration import layout


def order(sku='', product='Same ecommerce title', quantity=2, platform='SPX'):
    return Order(platform, 'SAME-ORDER', [Item(product, 'Same variant', quantity, '10.00', '20.00', sku, sku)])


class SkuEvidenceTests(unittest.TestCase):
    def registry(self):
        return SkuRegistry({'schema': 1, 'records': [{
            'sku': 'MASTER-A', 'original_product': 'MASTER', 'original_variation': '100ML',
            'sku_aliases': [{'platform': 'LAZADA', 'sku': 'OLD-123'}, {'platform': 'LAZADA', 'sku': 'MASTER-A'}],
            'ecommerce_sources': []}]})

    def test_different_codes_conflict_without_changing_fingerprint(self):
        a, b = order('SKU-A'), order('SKU-B')
        self.assertEqual(a.fingerprint(), b.fingerprint())
        self.assertTrue(engine.conflicting_order_skus([a, b]))

    def test_approved_alias_and_wrapped_case_codes_agree(self):
        registry = self.registry()
        self.assertFalse(engine.conflicting_order_skus([order('old-\n123', platform='LAZADA'),
                                                       order('MASTER-A', platform='LAZADA')], registry))

    def test_blank_legacy_does_not_hide_conflict_between_two_new_copies(self):
        self.assertFalse(engine.conflicting_order_skus([order(), order('SKU-A')]))
        self.assertTrue(engine.conflicting_order_skus([order(), order('SKU-A'), order('SKU-B')]))

    def test_reordered_rows_use_multisets_not_line_positions(self):
        first = order('SKU-A')
        first.items.append(Item('OTHER', 'BIG', 4, seller_sku='SKU-B'))
        second = copy.deepcopy(first)
        second.items.reverse()
        self.assertFalse(engine.conflicting_order_skus([first, second]))
        self.assertEqual(first.fingerprint(), second.fingerprint())

    def test_repeated_identical_descriptions_keep_sku_multiplicity(self):
        first = order('SKU-A')
        first.items.append(copy.deepcopy(first.items[0]))
        second = copy.deepcopy(first)
        second.items[1].seller_sku = 'SKU-B'
        self.assertTrue(engine.conflicting_order_skus([first, second]))
        first.items[1].seller_sku = ''
        self.assertFalse(engine.conflicting_order_skus([first, second]))

    def test_partial_complementary_copies_require_a_complete_real_source(self):
        first = order('SKU-A')
        first.items.append(Item('OTHER', 'BIG', 4))
        second = copy.deepcopy(first)
        second.items[0].seller_sku = ''
        second.items[1].seller_sku = 'SKU-B'
        self.assertFalse(engine.conflicting_order_skus([first, second]))
        self.assertEqual(engine.complete_sku_sources([first, second]), [])
        batches = {'one': Batch('one', 'SPX', '1', [first]), 'two': Batch('two', 'SPX', '2', [second])}
        with self.assertRaisesRegex(PackingError, 'partial SKU'):
            engine.preferred_order_sources(batches)
        complete = copy.deepcopy(first)
        complete.items[1].seller_sku = 'SKU-B'
        batches['complete'] = Batch('complete', 'SPX', '3', [complete])
        self.assertEqual(engine.preferred_order_sources(batches), {'SHOPEE:SAME-ORDER': 'complete'})


class SkuWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.root = Path(self.temp.name)
        self.folder = self.root / '20261010'
        self.folder.mkdir()
        self.catalog = self.root / 'ProductNames.json'
        self.batches = {}
        self.parser = patch.object(engine, 'parse_pdf', side_effect=lambda path, **kwargs: self.batches[Path(path).name])
        self.parser.start()

    def tearDown(self):
        self.parser.stop()
        self.temp.cleanup()

    def source(self, filename, value):
        path = self.folder / filename
        path.write_text(filename + repr(value), encoding='utf-8')
        self.batches[filename] = Batch(str(path), value.platform, file_hash(path), [value])
        return path

    def selected(self, revision=False):
        return [{'path': entry['path'], 'sha256': entry['sha256'], 'revision': revision}
                for entry in engine.scan(self.folder, self.catalog)['entries']]

    def define(self):
        names.save_catalog(self.catalog, [{'sku': 'SKU-A', 'alternative_product': 'INTERNAL A',
                                          'alternative_variation': '100 ML'}], 0)

    def test_scan_rejects_same_order_with_contradictory_sku_before_any_export(self):
        self.source('one.pdf', order('SKU-A'))
        self.source('two.pdf', order('SKU-B'))
        entries = engine.scan(self.folder, self.catalog)['entries']
        self.assertEqual([entry['status'] for entry in entries], ['ERROR', 'ERROR'])
        with self.assertRaisesRegex(PackingError, 'not ready'):
            engine.generate(self.folder, self.selected(), verifier=layout, catalog_path=self.catalog)
        with self.assertRaisesRegex(PackingError, 'Resolve'):
            sales.export_sales(self.folder, self.catalog)
        self.assertFalse(list((self.folder / '000 Output').glob('*.xlsx')))

    def test_packing_prefers_selected_sku_bearing_copy_and_keeps_original_audit(self):
        self.define()
        self.source('a-legacy.pdf', order())
        self.source('z-sku.pdf', order('SKU-A'))
        result = generate_approved(self.folder, self.selected(), verifier=layout, catalog_path=self.catalog)
        doc = Document(result['outputs'][0]['path'])
        self.assertEqual(doc.tables[0].cell(2, 1).text, 'INTERNAL A')
        history = json.loads((self.folder / '000 Output' / engine.HISTORY).read_text())
        self.assertEqual(history['orders']['SHOPEE:SAME-ORDER']['source'], 'z-sku.pdf')
        self.assertEqual(history['orders']['SHOPEE:SAME-ORDER']['order']['items'][0]['seller_sku'], 'SKU-A')
        self.assertEqual(self.batches['a-legacy.pdf'].orders[0].items[0].seller_sku, '')

    def test_sales_prefers_sku_copy_but_counts_order_once(self):
        self.define()
        self.source('a-legacy.pdf', order())
        self.source('z-sku.pdf', order('SKU-A'))
        result = sales.export_sales(self.folder, self.catalog)
        self.assertEqual((result['orders'], result['quantity']), (1, 2))
        book = load_workbook(result['path'], read_only=True, data_only=True)
        try:
            row = list(book['Order Items'].values)[1]
            self.assertEqual(row[6], 'INTERNAL A')
            self.assertEqual(row[9], 'z-sku.pdf')
        finally:
            book.close()

    def test_new_source_sku_conflicting_with_history_requires_review_and_revision(self):
        old_path = self.source('old.pdf', order('SKU-A'))
        first = generate_approved(self.folder, self.selected(), verifier=layout, catalog_path=self.catalog)
        old_path.unlink()
        self.source('new.pdf', order('SKU-B'))
        entry = engine.scan(self.folder, self.catalog)['entries'][0]
        self.assertEqual(entry['status'], 'REVIEW')
        self.assertIn('SKU: SKU-A', entry['details'])
        self.assertIn('SKU: SKU-B', entry['details'])
        with self.assertRaisesRegex(PackingError, 'explicit revision'):
            engine.generate(self.folder, self.selected(), verifier=layout, catalog_path=self.catalog)
        # No excludes matching prior outputs, not a reviewed SKU correction.
        second = generate_approved(self.folder, self.selected(revision=True), verifier=layout,
                                   catalog_path=self.catalog, repeat_action='skip')
        self.assertEqual(len(second['outputs']), 1)
        history = json.loads((self.folder / '000 Output' / engine.HISTORY).read_text())
        self.assertEqual(history['outputs'][first['outputs'][0]['filename']]['status'], 'superseded')
        self.assertEqual(history['orders']['SHOPEE:SAME-ORDER']['order']['items'][0]['seller_sku'], 'SKU-B')
        self.assertEqual(history['orders']['SHOPEE:SAME-ORDER']['fingerprint'], order('SKU-A').fingerprint())

    def test_legacy_history_without_sku_still_covers_matching_new_download(self):
        old_path = self.source('old.pdf', order())
        generate_approved(self.folder, self.selected(), verifier=layout, catalog_path=self.catalog)
        old_path.unlink()
        self.source('new.pdf', order('SKU-A'))
        entry = engine.scan(self.folder, self.catalog)['entries'][0]
        self.assertFalse(entry['requires_revision'])
        result = generate_approved(self.folder, self.selected(), verifier=layout,
                                   catalog_path=self.catalog, repeat_action='skip')
        self.assertEqual(result['outputs'], [])


if __name__ == '__main__':
    unittest.main(verbosity=2)
