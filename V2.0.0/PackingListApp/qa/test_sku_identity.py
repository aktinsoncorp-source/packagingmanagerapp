"""V2 SKU identity, safe legacy migration and protected Excel provenance."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / 'src'))
import names
from extractors import Batch, Item, Order, PackingError, lazada_page, normalize_sku
from sku_registry import SkuRegistry


def record(sku='PERF-A-30ML', product='SHOPEE PERFUME A', variation='30 ML', lazada='OLD LISTING A', old_sku='123456'):
    return {'sku': sku, 'original_product': product, 'original_variation': variation,
            'sku_aliases': [{'platform': 'SHOPEE', 'sku': sku},
                            {'platform': 'LAZADA', 'sku': sku}, {'platform': 'LAZADA', 'sku': old_sku}],
            'ecommerce_sources': [{'platform': 'SHOPEE', 'sku': sku, 'product': product, 'variation': variation},
                                  {'platform': 'LAZADA', 'sku': sku, 'product': lazada, 'variation': 'SIZE:30 ML'}]}


class SkuIdentityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP / 'qa')
        self.path = Path(self.temp.name) / 'ProductNames.json'
        self.book = Path(self.temp.name) / 'Names.xlsx'
        self.registry = SkuRegistry({'schema': 1, 'records': [record()]})

    def tearDown(self):
        self.temp.cleanup()

    def batch(self, sku='PERF-A-30ML', product='RENAMED SHOP TITLE', variation='NEW VARIANT NAME', platform='SPX'):
        item = Item(product, variation, 2, '45.00', '89.00', normalize_sku(sku), sku)
        return Batch('test.pdf', platform, 'source-hash', [Order(platform, 'ORDER-1', [item], 'Carefully please')])

    def defined(self):
        catalog = names.migrate_registry(self.path, self.registry)
        catalog['rows'][0].update(alternative_product='INTERNAL PERFUME', alternative_variation='30 ML', locked=True)
        return names.save_catalog(self.path, catalog['rows'], catalog['revision'])

    def edit(self, action):
        book = load_workbook(self.book)
        action(book['Names'])
        book.save(self.book)
        book.close()

    def test_changed_ecommerce_names_reuse_sku_and_preserve_source_fingerprint_and_prices(self):
        catalog = self.defined()
        batch = self.batch('perf-a-\n30ml')
        before = copy.deepcopy(batch.to_dict())
        fingerprint = batch.orders[0].fingerprint()
        catalog = names.encounter(self.path, [batch], self.registry)
        self.assertEqual(len(catalog['rows']), 1)
        self.assertTrue(catalog['rows'][0]['locked'])
        item = names.apply_names([batch], catalog, self.registry)[0].orders[0].items[0]
        self.assertEqual((item.product, item.variation), ('INTERNAL PERFUME', '30 ML'))
        self.assertEqual((item.unit_price, item.total_price, item.seller_sku), ('45.00', '89.00', 'PERF-A-30ML'))
        self.assertEqual(batch.to_dict(), before)
        self.assertEqual(batch.orders[0].fingerprint(), fingerprint)
        legacy = Item(batch.orders[0].items[0].product, batch.orders[0].items[0].variation, 2)
        self.assertEqual(Order('SPX', 'ORDER-1', [legacy]).fingerprint(), fingerprint)
        self.assertIn('RENAMED SHOP TITLE', [s['product'] for s in catalog['rows'][0]['ecommerce_sources']])
        self.assertEqual(names.encounter(self.path, [batch], self.registry), catalog)

    def test_old_numeric_lazada_alias_uses_approved_shopee_identity(self):
        catalog = self.defined()
        batch = self.batch('123456', platform='LAZADA')
        shown = names.apply_names([batch], catalog, self.registry)[0].orders[0].items[0]
        self.assertEqual((shown.product, shown.variation), ('INTERNAL PERFUME', '30 ML'))
        self.assertEqual(shown.seller_sku, '123456')

    def test_migration_preserves_existing_ids_alternatives_locks_and_old_original_links(self):
        old = names.save_catalog(self.path, [{'id': 'legacy', 'original_product': 'OLD LISTING A',
            'original_variation': 'SIZE:30 ML', 'alternative_product': 'REAL NAME', 'alternative_variation': '',
            'locked': True, 'aliases': [{'original_product': 'ANCIENT TITLE', 'original_variation': 'SMALL'}]}], 0)
        old_bytes = self.path.read_bytes()
        migrated = names.migrate_registry(self.path, self.registry)
        self.assertEqual(len(migrated['rows']), 1)
        row = migrated['rows'][0]
        self.assertEqual((row['id'], row['alternative_product'], row['alternative_variation'], row['locked']),
                         ('legacy', 'REAL NAME', '', True))
        self.assertEqual(row['sku'], 'PERF-A-30ML')
        self.assertIn({'original_product': 'ANCIENT TITLE', 'original_variation': 'SMALL'}, names.source_pairs(row))
        self.assertEqual(self.path.with_suffix('.json.bak').read_bytes(), old_bytes)
        self.assertEqual(names.migrate_registry(self.path, self.registry), migrated)

    def test_conflicting_existing_definitions_remain_reviewable_and_do_not_guess(self):
        names.save_catalog(self.path, [
            {'id': 'one', 'original_product': 'OLD LISTING A', 'original_variation': 'SIZE:30 ML',
             'alternative_product': 'WRONG OLD', 'alternative_variation': '30 ML', 'locked': True},
            {'id': 'two', 'original_product': 'SHOPEE PERFUME A', 'original_variation': '30 ML',
             'alternative_product': 'DIFFERENT NEW', 'alternative_variation': '30 ML', 'locked': True}], 0)
        catalog = names.migrate_registry(self.path, self.registry)
        self.assertEqual({r['id'] for r in catalog['rows']}, {'one', 'two'})
        self.assertTrue(all(r.get('review_notes') for r in catalog['rows']))
        with self.assertRaisesRegex(PackingError, 'conflicting'):
            names.apply_names([self.batch()], catalog, self.registry)
        before = self.path.read_bytes()
        with self.assertRaisesRegex(PackingError, 'conflicting'):
            names.encounter(self.path, [self.batch()], self.registry)
        self.assertEqual(self.path.read_bytes(), before)

    def test_unknown_sku_does_not_borrow_known_sku_with_same_title(self):
        catalog = self.defined()
        batch = self.batch('DIFFERENT-30ML', product='SHOPEE PERFUME A', variation='30 ML')
        catalog = names.encounter(self.path, [batch], self.registry)
        self.assertEqual(len(catalog['rows']), 2)
        shown = names.apply_names([batch], catalog, self.registry)[0].orders[0].items[0]
        self.assertEqual(shown.product, 'SHOPEE PERFUME A')
        self.assertEqual(next(r for r in catalog['rows'] if r['sku'] == 'DIFFERENT-30ML')['alternative_product'], '')

    def test_unknown_structured_sku_is_platform_scoped_and_rename_stable(self):
        first = self.batch('NEW-CODE', product='FIRST', platform='SPX')
        catalog = names.encounter(self.path, [first])
        catalog['rows'][0]['alternative_product'] = 'INTERNAL FIRST'
        catalog = names.save_catalog(self.path, catalog['rows'], catalog['revision'])
        renamed = self.batch('NEW-CODE', product='CHANGED', platform='FLASH')
        catalog = names.encounter(self.path, [renamed])
        self.assertEqual(len(catalog['rows']), 1)
        self.assertEqual(names.apply_names([renamed], catalog)[0].orders[0].items[0].product, 'INTERNAL FIRST')
        other = self.batch('NEW-CODE', product='OTHER PRODUCT', platform='TIKTOK')
        catalog = names.encounter(self.path, [other])
        self.assertEqual(len(catalog['rows']), 2)
        self.assertEqual(names.apply_names([other], catalog)[0].orders[0].items[0].product, 'OTHER PRODUCT')

    def test_generic_tiktok_size_never_combines_different_perfumes(self):
        first = self.batch('30 ml', product='PERFUME A', platform='TIKTOK')
        second = self.batch('30 ml', product='PERFUME B', platform='TIKTOK')
        catalog = names.encounter(self.path, [first, second], self.registry)
        fallback = [r for r in catalog['rows'] if not r.get('sku')]
        self.assertEqual(len(fallback), 2)
        self.assertTrue(all(r['review_notes'] for r in fallback))
        fallback[0]['alternative_product'] = 'INTERNAL A'
        shown = names.apply_names([first, second], catalog, self.registry)
        self.assertNotEqual(shown[0].orders[0].items[0].product, shown[1].orders[0].items[0].product)

    def test_unknown_numeric_seller_code_from_lazada_keeps_source_identity(self):
        batch = self.batch('9988776655', product='UNLISTED PRODUCT', platform='LAZADA')
        catalog = names.encounter(self.path, [batch], self.registry)
        row = next(r for r in catalog['rows'] if r.get('sku') == '9988776655')
        row['alternative_product'] = 'INTERNAL NUMERIC PRODUCT'
        catalog = names.save_catalog(self.path, catalog['rows'], catalog['revision'])
        changed = self.batch('9988776655', product='CHANGED TITLE', platform='LAZADA')
        shown = names.apply_names([changed], catalog, self.registry)[0].orders[0].items[0]
        self.assertEqual(shown.product, 'INTERNAL NUMERIC PRODUCT')

    def test_excel_roundtrip_retains_metadata_and_hidden_sort_columns(self):
        catalog = self.defined()
        names.export_excel(self.book, catalog)
        book = load_workbook(self.book)
        sheet = book['Names']
        self.assertEqual(sheet['G1'].value, 'SKU')
        self.assertEqual(sheet['C1'].value, 'Product Name')
        self.assertTrue(sheet.column_dimensions['H'].hidden)
        self.assertTrue(sheet['G2'].protection.locked)
        self.assertTrue(sheet.auto_filter.ref.startswith('A1:H'))
        book.close()
        self.assertEqual(names.import_excel(self.book, catalog['rows'])['rows'], catalog['rows'])

    def test_excel_rejects_changing_existing_sku_even_after_sheet_unprotection(self):
        catalog = self.defined()
        names.export_excel(self.book, catalog)
        self.edit(lambda sheet: sheet.__setitem__('G2', 'DIFFERENT-CODE'))
        with self.assertRaisesRegex(PackingError, 'existing SKU'):
            names.import_excel(self.book, catalog['rows'])
        self.assertEqual(names.load_catalog(self.path), catalog)

    def test_excel_single_shared_edit_propagates_but_same_sku_cannot_split_names(self):
        catalog = self.defined()
        catalog['rows'][0]['locked'] = False
        catalog = names.save_catalog(self.path, catalog['rows'], catalog['revision'])
        names.export_excel(self.book, catalog)
        self.edit(lambda sheet: sheet.__setitem__('C2', 'NEW INTERNAL'))
        imported = names.import_excel(self.book, catalog['rows'])
        self.assertEqual(len(imported['rows']), 1)
        self.assertEqual(imported['rows'][0]['alternative_product'], 'NEW INTERNAL')
        names.export_excel(self.book, catalog)
        self.edit(lambda sheet: (sheet.__setitem__('C2', 'FIRST INTERNAL'), sheet.__setitem__('C3', 'SECOND INTERNAL')))
        with self.assertRaisesRegex(PackingError, 'same SKU'):
            names.import_excel(self.book, catalog['rows'])

    def test_manual_sku_can_be_added_without_invented_ecommerce_names(self):
        result = names.validate_rows([{'sku': 'NEW-100ML', 'alternative_product': 'INTERNAL', 'alternative_variation': '100 ML'}])
        self.assertEqual(result[0]['sku'], 'NEW-100ML')
        self.assertEqual(result[0]['original_product'], 'NEW-100ML')

    def test_manual_generic_or_duplicate_sku_cannot_silently_reassign_a_product(self):
        catalog = self.defined()
        for sku in ('30 ml', '123456'):
            with self.subTest(sku=sku), self.assertRaisesRegex(PackingError, 'unique product SKU'):
                names.validate_rows([{'sku': sku, 'alternative_product': 'NEW'}])
        rows = copy.deepcopy(catalog['rows'])
        rows.append({'sku': 'PERF-A-30ML', 'alternative_product': 'DIFFERENT', 'alternative_variation': '30 ML'})
        with self.assertRaisesRegex(PackingError, 'already belongs'):
            names.save_catalog(self.path, rows, catalog['revision'])
        self.assertEqual(names.load_catalog(self.path), catalog)

    def test_legacy_excel_without_sku_metadata_keeps_saved_identity(self):
        catalog = self.defined()
        names.export_excel(self.book, catalog)
        self.edit(lambda sheet: sheet.delete_cols(7, 2))
        self.assertEqual(names.import_excel(self.book, catalog['rows'])['rows'], catalog['rows'])

    def test_distinct_linked_skus_split_in_excel_without_losing_their_identity(self):
        self.registry = SkuRegistry({'schema': 1, 'records': [record(),
            record('PERF-B-100ML', 'SHOPEE PERFUME B', '100 ML', 'OLD LISTING B', '654321')]})
        catalog = names.migrate_registry(self.path, self.registry)
        for row in catalog['rows']:
            row.update(alternative_product='SAME INTERNAL', alternative_variation='30 ML')
        ids = [row['id'] for row in catalog['rows']]
        catalog = names.save_catalog(self.path, catalog['rows'], catalog['revision'], [ids])
        self.assertEqual(len(catalog['rows']), 1)
        names.export_excel(self.book, catalog)
        def edit(sheet):
            for number in range(2, sheet.max_row + 1):
                is_b = ' B' in sheet.cell(number, 1).value
                sheet.cell(number, 3, 'INTERNAL B' if is_b else 'INTERNAL A')
                sheet.cell(number, 4, '100 ML' if is_b else '30 ML')
        self.edit(edit)
        imported = names.import_excel(self.book, catalog['rows'])
        self.assertEqual(len(imported['rows']), 2)
        imported = names.save_catalog(self.path, imported['rows'], catalog['revision'])
        by_sku = {row['sku']: row for row in imported['rows']}
        self.assertEqual(set(by_sku), {'PERF-A-30ML', 'PERF-B-100ML'})
        for sku, expected in [('PERF-A-30ML', 'INTERNAL A'), ('PERF-B-100ML', 'INTERNAL B')]:
            shown = names.apply_names([self.batch(sku)], imported, self.registry)[0].orders[0].items[0]
            self.assertEqual(shown.product, expected)
        self.assertEqual(by_sku['PERF-A-30ML']['id'], ids[0])

    def test_visible_only_excel_sort_does_not_assign_stale_sku_metadata(self):
        self.registry = SkuRegistry({'schema': 1, 'records': [record(),
            record('PERF-B-100ML', 'SHOPEE PERFUME B', '100 ML', 'OLD LISTING B', '654321')]})
        catalog = names.migrate_registry(self.path, self.registry)
        names.export_excel(self.book, catalog)
        def edit(sheet):
            values = [[sheet.cell(number, column).value for column in range(1, 5)]
                      for number in range(2, sheet.max_row + 1)]
            for number, values in enumerate(reversed(values), 2):
                for column, value in enumerate(values, 1):
                    sheet.cell(number, column, value)
        self.edit(edit)
        self.assertEqual(names.import_excel(self.book, catalog['rows'])['rows'], catalog['rows'])

    def test_same_title_for_unrelated_platform_skus_does_not_force_a_merge(self):
        first = self.batch('NEW-CODE', product='COMMON TITLE', platform='SPX')
        second = self.batch('NEW-CODE', product='COMMON TITLE', platform='TIKTOK')
        catalog = names.encounter(self.path, [first, second])
        self.assertEqual(len(catalog['rows']), 2)
        self.assertEqual({r['sku_aliases'][0]['platform'] for r in catalog['rows']}, {'SHOPEE', 'TIKTOK'})

    def test_unknown_platform_code_equal_to_master_sku_stays_separate_after_rescan(self):
        catalog = self.defined()
        tiktok = self.batch('PERF-A-30ML', product='UNRELATED TIKTOK PRODUCT', platform='TIKTOK')
        catalog = names.encounter(self.path, [tiktok], self.registry)
        self.assertEqual(len(catalog['rows']), 2)
        self.assertEqual(names.apply_names([tiktok], catalog, self.registry)[0].orders[0].items[0].product,
                         'UNRELATED TIKTOK PRODUCT')
        catalog = names.encounter(self.path, [tiktok], self.registry)
        self.assertEqual(len(catalog['rows']), 2)
        self.assertEqual(names.apply_names([tiktok], catalog, self.registry)[0].orders[0].items[0].product,
                         'UNRELATED TIKTOK PRODUCT')
        scoped = next(row for row in catalog['rows'] if row['sku_scope'] == 'TIKTOK')
        self.assertTrue(all(alias['platform'] == 'TIKTOK' for alias in scoped['sku_aliases']))

    def test_existing_sku_cannot_be_cleared_or_reassigned_on_unlocked_save(self):
        catalog = self.defined()
        for sku in ('', 'OTHER-CODE'):
            rows = copy.deepcopy(catalog['rows'])
            rows[0].update(locked=False, sku=sku)
            with self.subTest(sku=sku), self.assertRaisesRegex(PackingError, 'existing SKU'):
                names.save_catalog(self.path, rows, catalog['revision'])
        self.assertEqual(names.load_catalog(self.path), catalog)

    def test_readonly_identity_cannot_be_bypassed_by_editing_hidden_scope_or_links(self):
        catalog = self.defined()
        edits = [lambda row: row.__setitem__('sku_scope', 'TIKTOK'),
                 lambda row: row['sku_aliases'].append({'platform': 'TIKTOK', 'sku': 'OTHER-PRODUCT'}),
                 lambda row: row.__setitem__('sku_aliases', []),
                 lambda row: row.__setitem__('ecommerce_sources', [])]
        for edit in edits:
            rows = copy.deepcopy(catalog['rows'])
            rows[0]['locked'] = False
            edit(rows[0])
            with self.assertRaisesRegex(PackingError, 'SKU source'):
                names.save_catalog(self.path, rows, catalog['revision'])
        self.assertEqual(names.load_catalog(self.path), catalog)


class SkuExtractionTests(unittest.TestCase):
    def test_lazada_keeps_raw_wrapped_sku_and_original_description(self):
        class Page:
            def extract_text(self):
                return 'Lazada Seller Center Order Number:123'
            def extract_tables(self):
                return [[['Item Name', 'Item SKU Seller', 'Item Variation', 'Quantity'],
                         ['Listing\nName', 'PERF-A-\n30ML', 'Size:30 ml', '2']]]
        order = lazada_page(Page(), 1)
        item = order.items[0]
        self.assertEqual((item.seller_sku, item.raw_seller_sku), ('PERF-A-30ML', 'PERF-A-\n30ML'))
        self.assertEqual((item.product, item.variation, item.pieces), ('Listing Name', 'Size:30 ml', 2))
        self.assertEqual(order.fingerprint(), Order('LAZADA', '123', [Item('Listing Name', 'Size:30 ml', 2)]).fingerprint())


if __name__ == '__main__':
    unittest.main(verbosity=2)
