"""Visible-column pastes/sorts must not apply leftover IDs to other originals."""
import copy
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

APP=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(APP/'src'))
import names
from extractors import PackingError

class ExcelIdentityTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=APP/'qa');self.root=Path(self.temp.name)
        self.path=self.root/'ProductNames.json';self.book=self.root/'Names.xlsx'
        self.catalog=names.save_catalog(self.path,[
            dict(id='base',original_product='A ORIGINAL',original_variation='SMALL',alternative_product='STANDARD A',alternative_variation='100ML',aliases=[dict(original_product='B LINK',original_variation='SMALL')]),
            dict(id='other',original_product='C ORIGINAL',original_variation='BIG',alternative_product='STANDARD C',alternative_variation='1L')],0)
        names.export_excel(self.book,self.catalog)
    def tearDown(self):self.temp.cleanup()
    def edit(self,fn):
        wb=load_workbook(self.book);fn(wb['Names']);wb.save(self.book);wb.close()
    def paste(self,ws,values):
        for number,values in enumerate(values,2):
            for column,value in enumerate(values,1):ws.cell(number,column,value)
    def test_larger_pasted_list_preserves_all_originals_and_ids(self):
        def edit(ws):self.paste(ws,[['NEW ITEM','30ML','NEW','30ML'],['C ORIGINAL','BIG','STANDARD C','1L'],['A ORIGINAL','SMALL','STANDARD A','100ML'],['B LINK','SMALL','STANDARD A','100ML']])
        self.edit(edit);before=self.path.read_bytes();result=names.import_excel(self.book,self.catalog['rows'])
        self.assertEqual(self.path.read_bytes(),before);self.assertEqual(result['repaired'],3);self.assertEqual(result['added'],1)
        current={r['id']:r for r in result['rows']}
        for row in self.catalog['rows']:self.assertEqual(current[row['id']],row)
        self.assertEqual(result['priority'][-1],'base')
        self.assertEqual(sum(len(names.source_pairs(r)) for r in result['rows']),4)
    def test_visible_only_sort_moves_alternatives_without_moving_ids(self):
        def edit(ws):
            visible=[[ws.cell(n,c).value for c in range(1,5)] for n in range(2,5)]
            self.paste(ws,list(reversed(visible)))
        self.edit(edit);result=names.import_excel(self.book,self.catalog['rows'])
        self.assertEqual(result['rows'],self.catalog['rows']);self.assertEqual(result['repaired'],2)
        self.assertEqual(result['priority'],['other','base'])
    def test_pasted_new_alternatives_reach_the_correct_originals(self):
        def edit(ws):self.paste(ws,[['C ORIGINAL','BIG','UPDATED C','2L'],['B LINK','SMALL','SECOND','250ML'],['A ORIGINAL','SMALL','FIRST','30ML']])
        self.edit(edit);result=names.import_excel(self.book,self.catalog['rows'])
        values={r['original_product']:(r['alternative_product'],r['alternative_variation']) for r in result['rows']}
        self.assertEqual(values,{'A ORIGINAL':('FIRST','30ML'),'B LINK':('SECOND','250ML'),'C ORIGINAL':('UPDATED C','2L')})
        self.assertEqual(result['separated'],1)
    def test_pasted_inventory_keeps_single_changed_size_on_its_own_original(self):
        def edit(ws):self.paste(ws,[['C ORIGINAL','BIG','STANDARD C','1L'],['B LINK','SMALL','STANDARD A','25KG'],['A ORIGINAL','SMALL','STANDARD A','100ML']])
        self.edit(edit);result=names.import_excel(self.book,self.catalog['rows'])
        by_product={r['original_product']:r for r in result['rows']}
        self.assertEqual(by_product['A ORIGINAL']['alternative_variation'],'100ML')
        self.assertEqual(by_product['B LINK']['alternative_variation'],'25KG')
        self.assertEqual(result['separated'],1)
    def test_locked_originals_remain_locked_after_visible_sort(self):
        locked=copy.deepcopy(self.catalog['rows'])
        for row in locked:row['locked']=True
        self.catalog=names.save_catalog(self.path,locked,self.catalog['revision']);names.export_excel(self.book,self.catalog)
        self.edit(lambda ws:self.paste(ws,[['C ORIGINAL','BIG','STANDARD C','1L'],['B LINK','SMALL','STANDARD A','100ML'],['A ORIGINAL','SMALL','STANDARD A','100ML']]))
        self.assertEqual(names.import_excel(self.book,self.catalog['rows'])['rows'],self.catalog['rows'])
        self.edit(lambda ws:ws.__setitem__('C2','FORBIDDEN'))
        preview=names.import_excel(self.book,self.catalog['rows'])
        self.assertEqual(preview['skipped_locked'],[2])
        self.assertEqual(preview['rows'],self.catalog['rows'])
        self.assertEqual(names.load_catalog(self.path),self.catalog)
    def test_genuine_original_rename_still_uses_its_hidden_identity(self):
        self.edit(lambda ws:ws.__setitem__('A3','RENAMED LINK'))
        result=names.import_excel(self.book,self.catalog['rows'])
        row=next(r for r in result['rows'] if r['id']=='base')
        self.assertEqual(row['aliases'],[dict(original_product='RENAMED LINK',original_variation='SMALL')]);self.assertEqual(result['repaired'],0)
    def test_bad_pair_id_recovers_when_visible_original_is_known(self):
        self.edit(lambda ws:ws.__setitem__('F3','broken pair'))
        self.assertEqual(names.import_excel(self.book,self.catalog['rows'])['rows'],self.catalog['rows'])
    def test_bad_pair_id_cannot_guess_a_new_original(self):
        def edit(ws):ws['A3']='UNMATCHED NEW NAME';ws['F3']='broken pair'
        self.edit(edit);before=self.path.read_bytes()
        with self.assertRaisesRegex(PackingError,'cannot be matched safely'):names.import_excel(self.book,self.catalog['rows'])
        self.assertEqual(self.path.read_bytes(),before)
    def test_true_duplicate_original_reports_both_excel_rows(self):
        self.edit(lambda ws:ws.append([ws.cell(2,c).value for c in range(1,7)]))
        with self.assertRaisesRegex(PackingError,'rows 2 and 5'):names.import_excel(self.book,self.catalog['rows'])
    def test_excel_filter_sort_includes_hidden_identity_columns(self):
        wb=load_workbook(self.book)
        self.assertEqual(wb['Names'].auto_filter.ref,'A1:F4');wb.close()

if __name__=='__main__':unittest.main()
