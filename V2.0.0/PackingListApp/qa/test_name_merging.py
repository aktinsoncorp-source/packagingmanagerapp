"""Confirmed merges, original-pair provenance, normalization and Excel regression."""
import copy
import json
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook

APP=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(APP/'src'))
import names
from extractors import Batch,Order,Item,PackingError
from picking import picking_rows

class MergeTests(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory(dir=APP/'qa');self.root=Path(self.temp.name)
        self.path=self.root/'ProductNames.json';self.book=self.root/'Names.xlsx'
        self.batch=Batch('source.pdf','SPX','original-hash',[Order('SPX','ORDER1',[Item('Original A','1kg',2),Item('Original B','',3)])])
        self.catalog=names.encounter(self.path,[self.batch])
        for row in self.catalog['rows']:
            row['alternative_product']='Standard Item '
            row['alternative_variation']='1kg  '
        self.catalog['rows'][1]['alternative_product']='standard item'

    def tearDown(self):self.temp.cleanup()
    def merged(self):
        preview=names.prepare_save(self.catalog['rows'],[self.catalog['rows'][1]['id'],self.catalog['rows'][0]['id']])
        return names.save_catalog(self.path,preview['rows'],self.catalog['revision'],[preview['merges'][0]['ids']])
    def edit_book(self,edit):
        wb=load_workbook(self.book);edit(wb['Names']);wb.save(self.book);wb.close()

    def test_prompt_preview_is_case_insensitive_without_saving_or_merging(self):
        before=self.path.read_bytes();preview=names.prepare_save(self.catalog['rows'])
        self.assertEqual(len(preview['merges']),1);self.assertEqual(len(preview['rows']),2)
        self.assertEqual(preview['merges'][0]['alternative_product'],'STANDARD ITEM')
        self.assertEqual(self.path.read_bytes(),before)
        declined=names.save_catalog(self.path,preview['rows'],self.catalog['revision'])
        self.assertEqual(len(declined['rows']),2)

    def test_yes_retains_first_selected_base_and_all_original_pairs(self):
        saved=self.merged();self.assertEqual(len(saved['rows']),1)
        row=saved['rows'][0];self.assertEqual(row['id'],self.catalog['rows'][1]['id'])
        self.assertEqual((row['original_product'],row['original_variation']),('ORIGINAL B','—'))
        self.assertEqual(row['aliases'],[{'original_product':'ORIGINAL A','original_variation':'1KG'}])
        self.assertEqual(names.load_catalog(self.path),saved)

    def test_rescan_does_not_recreate_merged_rows_and_all_sources_use_alias(self):
        saved=self.merged();self.assertEqual(names.encounter(self.path,[self.batch]),saved)
        changed=copy.deepcopy(self.batch)
        changed.orders[0].items[0].product='original a  '
        changed.orders[0].items[0].variation='1KG '
        self.assertEqual(names.encounter(self.path,[changed]),saved)
        printed=names.apply_names([changed],saved)
        self.assertEqual(picking_rows(printed),[('STANDARD ITEM','1KG',5)])
        self.assertEqual(changed.orders[0].items[0].product,'original a  ')
        self.assertEqual(printed[0].orders[0].pieces,5)

    def test_both_alternatives_required_no_automatic_blank_merges(self):
        for field in names.FIELDS[2:]:
            rows=copy.deepcopy(self.catalog['rows'])
            for row in rows:row[field]=''
            self.assertEqual(names.prepare_save(rows)['merges'],[])
            with self.assertRaises(PackingError):names.merge_rows(rows,[[r['id'] for r in rows]])

    def test_trim_and_uppercase_all_fields_keep_internal_spaces_and_symbols(self):
        row={'original_product':'Original / A  ', 'original_variation':'Light blue 1kg ', 'alternative_product':'Standard  name - (new) ', 'alternative_variation':'1.5 L / 50% '}
        row=names.validate_rows([row])[0]
        self.assertEqual([row[f] for f in names.FIELDS],['ORIGINAL / A','LIGHT BLUE 1KG','STANDARD  NAME - (NEW)','1.5 L / 50%'])
        self.assertEqual(names.validate_alternative('   ','Alternative'),'')
        self.assertEqual(names.validate_alternative(' @ leading','Alternative'),'LEADING')

    def test_all_fields_reject_paragraph_and_control_characters(self):
        for field in names.FIELDS:
            for sep in ['\n','\r','\r\n','\t','\u2028','\u2029','\x85']:
                rows=copy.deepcopy(self.catalog['rows']);rows[0][field]='FIRST'+sep+'SECOND'
                with self.subTest(field=field,sep=sep),self.assertRaises(PackingError):names.validate_rows(rows)

    def test_case_only_originals_in_confirmed_merge_keep_first_selected_pair(self):
        rows=copy.deepcopy(self.catalog['rows'])
        rows[0]['original_product']='Same original';rows[0]['original_variation']='same variation'
        rows[1]['original_product']='SAME ORIGINAL ';rows[1]['original_variation']='SAME VARIATION '
        preview=names.prepare_save(rows,[rows[1]['id'],rows[0]['id']])
        merged=names.merge_rows(preview['rows'],[preview['merges'][0]['ids']])
        self.assertEqual(len(merged),1);self.assertEqual(merged[0]['id'],rows[1]['id'])
        self.assertEqual(merged[0]['original_product'],'SAME ORIGINAL');self.assertEqual(merged[0]['aliases'],[])

    def test_excel_originals_becoming_same_pair_can_merge_after_confirmation(self):
        current=names.save_catalog(self.path,self.catalog['rows'],self.catalog['revision'])
        names.export_excel(self.book,current)
        def edit(ws):ws['A3']=ws['A2'].value.lower();ws['B3']=ws['B2'].value.lower()
        self.edit_book(edit)
        imported=names.import_excel(self.book,current['rows'])
        preview=names.prepare_save(imported['rows'],imported['priority'])
        self.assertEqual(len(preview['merges']),1)
        saved=names.save_catalog(self.path,preview['rows'],current['revision'],[preview['merges'][0]['ids']])
        self.assertEqual(len(saved['rows']),1);self.assertEqual(saved['rows'][0]['aliases'],[])

    def test_case_only_originals_are_one_pair_on_encounter(self):
        variant=copy.deepcopy(self.batch)
        variant.orders[0].items=[Item('original a ','1KG ',4),Item('ORIGINAL B','—',1)]
        self.assertEqual(len(names.encounter(self.path,[variant])['rows']),2)

    def test_delete_merged_definition_rediscovers_unmerged_blank_rows(self):
        merged=self.merged();names.save_catalog(self.path,[],merged['revision'])
        new=names.encounter(self.path,[self.batch]);self.assertEqual(len(new['rows']),2)
        self.assertTrue(all(not r['aliases'] and not r['alternative_product'] and not r['alternative_variation'] for r in new['rows']))
        self.assertNotIn(merged['rows'][0]['id'],[r['id'] for r in new['rows']])

    def test_excel_all_linked_originals_editable_and_alternative_edit_propagates(self):
        merged=self.merged();names.export_excel(self.book,merged)
        def edit(ws):
            self.assertEqual(ws.max_row,3);self.assertTrue(ws.column_dimensions['F'].hidden)
            ws['A3']='renamed linked original  ';ws['B3']='500ml '
            ws['C3']='revised standard ';ws['D2']='250ml '
        self.edit_book(edit)
        preview=names.import_excel(self.book,merged['rows']);saved=names.save_catalog(self.path,preview['rows'],merged['revision'])
        self.assertEqual(len(saved['rows']),1);row=saved['rows'][0]
        self.assertEqual(row['id'],merged['rows'][0]['id'])
        self.assertEqual(row['aliases'],[{'original_product':'RENAMED LINKED ORIGINAL','original_variation':'500ML'}])
        self.assertEqual(row['alternative_product'],'REVISED STANDARD');self.assertEqual(row['alternative_variation'],'250ML')

    def test_excel_missing_linked_rows_retained(self):
        merged=self.merged();names.export_excel(self.book,merged)
        self.edit_book(lambda ws:ws.delete_rows(3))
        self.assertEqual(names.import_excel(self.book,merged['rows'])['rows'],merged['rows'])

    def test_excel_different_alternatives_split_without_saving_partial_edits(self):
        merged=self.merged();names.export_excel(self.book,merged)
        def edit(ws):ws['C2']='first name';ws['C3']='second name'
        self.edit_book(edit);before=self.path.read_bytes()
        imported=names.import_excel(self.book,merged['rows'])
        self.assertEqual(self.path.read_bytes(),before)
        self.assertEqual(imported['separated'],1)
        self.assertEqual(len(imported['rows']),2)
        by_original={r['original_product']:r for r in imported['rows']}
        self.assertEqual(by_original['ORIGINAL B']['id'],merged['rows'][0]['id'])
        self.assertEqual(by_original['ORIGINAL B']['alternative_product'],'FIRST NAME')
        self.assertEqual(by_original['ORIGINAL A']['alternative_product'],'SECOND NAME')
        self.assertTrue(all(not r['aliases'] and not r['locked'] for r in imported['rows']))
        self.assertEqual(imported['priority'],[by_original['ORIGINAL B']['id'],by_original['ORIGINAL A']['id']])
        saved=names.save_catalog(self.path,imported['rows'],merged['revision'])
        self.assertEqual(names.encounter(self.path,[self.batch]),saved)
        printed=names.apply_names([self.batch],saved)
        self.assertEqual([(i.product,i.variation,i.pieces) for i in printed[0].orders[0].items],[('SECOND NAME','1KG',2),('FIRST NAME','1KG',3)])

    def test_excel_split_keeps_complete_pairs_omitted_originals_and_base_id(self):
        merged=self.merged();row=merged['rows'][0]
        row['aliases'].append({'original_product':'ORIGINAL C','original_variation':'1KG'})
        row['aliases'].append({'original_product':'ORIGINAL D','original_variation':'1KG'})
        merged=names.save_catalog(self.path,merged['rows'],merged['revision'])
        names.export_excel(self.book,merged)
        def edit(ws):
            ws['C2']='first name';ws['D2']='100ml'
            ws['A3']='renamed listing';ws['C3']='second name';ws['D3']=''
            ws['C4']='first name';ws['D4']='100ml'
            ws.delete_rows(5)
        self.edit_book(edit)
        imported=names.import_excel(self.book,merged['rows'])
        self.assertEqual(imported['separated'],2)
        by_product={r['original_product']:r for r in imported['rows']}
        self.assertEqual(by_product['ORIGINAL B']['id'],row['id'])
        self.assertEqual(by_product['ORIGINAL B']['aliases'],[{'original_product':'ORIGINAL C','original_variation':'1KG'}])
        self.assertEqual((by_product['RENAMED LISTING']['alternative_product'],by_product['RENAMED LISTING']['alternative_variation']),('SECOND NAME',''))
        self.assertEqual((by_product['ORIGINAL D']['alternative_product'],by_product['ORIGINAL D']['alternative_variation']),('STANDARD ITEM','1KG'))
        names.export_excel(self.book,{'rows':imported['rows']})
        self.assertEqual(names.import_excel(self.book,imported['rows'])['rows'],imported['rows'])

    def test_excel_variation_conflicts_split_without_combining_name_edits(self):
        merged=self.merged();names.export_excel(self.book,merged)
        def edit(ws):
            ws['C2']='changed name';ws['D2']='100ml';ws['D3']='1l'
        self.edit_book(edit)
        imported=names.import_excel(self.book,merged['rows'])
        by_product={r['original_product']:r for r in imported['rows']}
        self.assertEqual((by_product['ORIGINAL B']['alternative_product'],by_product['ORIGINAL B']['alternative_variation']),('CHANGED NAME','100ML'))
        self.assertEqual((by_product['ORIGINAL A']['alternative_product'],by_product['ORIGINAL A']['alternative_variation']),('STANDARD ITEM','1L'))

    def test_excel_split_matches_merge_only_after_app_confirmation(self):
        merged=self.merged();names.export_excel(self.book,merged)
        def edit(ws):
            ws['C2']='first name';ws['C3']='second name'
            ws.append(['Original C','100ml','second name','1kg'])
        self.edit_book(edit)
        imported=names.import_excel(self.book,merged['rows'])
        preview=names.prepare_save(imported['rows'],imported['priority'])
        self.assertEqual(len(preview['merges']),1)
        self.assertEqual(preview['merges'][0]['base_product'],'ORIGINAL A')
        separate=names.save_catalog(self.path,preview['rows'],merged['revision'])
        self.assertEqual(len(separate['rows']),3)
        joined=names.save_catalog(self.path,preview['rows'],separate['revision'],[preview['merges'][0]['ids']])
        self.assertEqual(len(joined['rows']),2)
        self.assertEqual(sum(len(names.source_pairs(r)) for r in joined['rows']),3)

    def test_excel_split_validation_error_preserves_saved_catalog(self):
        merged=self.merged();names.export_excel(self.book,merged)
        def edit(ws):
            ws['C2']='first name';ws['C3']='second name';ws['A3']=ws['A2'].value;ws['B3']=ws['B2'].value
        self.edit_book(edit);before=self.path.read_bytes()
        with self.assertRaisesRegex(PackingError,'already belong'):names.import_excel(self.book,merged['rows'])
        self.assertEqual(self.path.read_bytes(),before)

    def test_excel_new_matching_definition_prompts_with_workbook_base_order(self):
        merged=self.merged();names.export_excel(self.book,merged)
        self.edit_book(lambda ws:ws.append(['Original C','250ml','standard item','1kg']))
        imported=names.import_excel(self.book,merged['rows']);preview=names.prepare_save(imported['rows'],imported['priority'])
        self.assertEqual(len(preview['merges']),1);self.assertEqual(preview['merges'][0]['ids'][0],merged['rows'][0]['id'])
        saved=names.save_catalog(self.path,preview['rows'],merged['revision'],[preview['merges'][0]['ids']])
        self.assertEqual(len(saved['rows']),1);self.assertEqual(len(saved['rows'][0]['aliases']),2)

    def test_excel_paragraphs_in_every_column_rejected_and_unchanged(self):
        merged=self.merged();before=self.path.read_bytes()
        for column in 'ABCD':
            names.export_excel(self.book,merged)
            self.edit_book(lambda ws:ws.__setitem__(column+'2','First\nSecond'))
            with self.subTest(column=column),self.assertRaises(PackingError):names.import_excel(self.book,merged['rows'])
        self.assertEqual(self.path.read_bytes(),before)

    def test_existing_catalog_migrates_without_losing_user_deletions_or_ids(self):
        old={'schema':1,'revision':7,'rows':[{'id':'keep-this','original_product':'Old name ','original_variation':'old variation ','alternative_product':'Saved alternate ','alternative_variation':''}]}
        self.path.write_text(json.dumps(old),encoding='utf-8')
        new=names.load_catalog(self.path)
        self.assertEqual((new['schema'],new['revision'],len(new['rows'])),(2,8,1))
        self.assertEqual(new['rows'][0]['id'],'keep-this');self.assertEqual(new['rows'][0]['alternative_product'],'SAVED ALTERNATE')
        self.assertEqual(json.loads(self.path.with_suffix('.json.bak').read_text(encoding='utf-8')),old)

if __name__=='__main__':unittest.main(verbosity=2)
