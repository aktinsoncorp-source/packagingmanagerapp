import copy
import json
import re
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
ROOT = Path(__import__('os').environ.get('PACKING_TEST_DATA', APP.parent))
sys.path.insert(0, str(APP / 'src'))
import engine
from generation_test_support import generate_approved
from documents import build_document
from extractors import Batch, Item, Order, PackingError, file_hash, parse_pdf
from docx import Document
from pypdf import PdfReader, PdfWriter


def fake_layout(paths):
    return [{'path': str(p), 'pages': 1, 'orders': len(Document(p).tables)-1} for p in paths]


class PackingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='packing-test-', dir=APP / 'qa')
        self.folder = Path(self.temp.name) / '20260926'
        self.folder.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def copy_source(self, name, sub='SPX'):
        path = self.folder / sub / name
        path.parent.mkdir(exist_ok=True)
        shutil.copyfile(ROOT / '20260926' / sub / name, path)
        return path

    def entries(self):
        return engine.scan(self.folder)['entries']

    def select(self, statuses=('READY',)):
        return [{'path': e['path'], 'sha256': e['sha256']} for e in self.entries() if e['status'] in statuses]

    def test_every_approved_order_matches_products_variations_and_quantities(self):
        from lazada_history import legacy_key
        expected = {}
        for group in ['SPX','FLASH','LAZADA_TIKTOK']:
            expected.update(engine.read_completed_doc(ROOT/'20260926'/'000 Output'/f'{group}_Packing_Checklist_2026-09-26_1.docx', group))
        actual, tracking_keys = {}, set()
        for path in (ROOT/'20260926').rglob('*.pdf'):
            for order in parse_pdf(path).orders:
                self.assertNotIn(order.key, tracking_keys)
                tracking_keys.add(order.key)
                if order.platform == 'LAZADA':
                    self.assertTrue(order.tracking_id, str(path))
                # Approved pre-tracking documents identify these unsplit fixture
                # orders by number only; never collapse multiple current parcels.
                key = legacy_key(order.key)
                self.assertNotIn(key, actual)
                actual[key] = order
        self.assertEqual(set(expected), set(actual))
        self.assertEqual(len(actual),139)
        self.assertEqual(sum(o.pieces for o in actual.values()),318)
        for key, order in actual.items():
            self.assertEqual(order.fingerprint(),expected[key]['fingerprint'],key)

    def test_all_sample_dates_cover_every_source_order_id(self):
        import pdfplumber
        total_files = 0
        for date in ['20260924','20260925','20260926']:
            for path in (ROOT/date).rglob('*.pdf'):
                batch = parse_pdf(path)
                with pdfplumber.open(path) as pdf:
                    text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
                ids = set(re.findall(r'Order\s+(?:ID|Number)\s*[:：]\s*([A-Z0-9]+)',text))
                self.assertEqual(ids,{o.order_id for o in batch.orders},str(path))
                total_files += 1
        self.assertEqual(total_files,27)

    def test_legacy_log_is_verified_and_links_saved_without_reprinting(self):
        shutil.copytree(ROOT/'20260926',self.folder,dirs_exist_ok=True)
        log = self.folder/'000 Output'/engine.LOG
        before = file_hash(self.folder/'000 Output'/'Processed_Files.txt')
        out = self.folder/'000 Output'
        documents_before = {p.name: file_hash(p) for p in out.iterdir() if p.suffix.lower() in ('.docx', '.pdf')}
        entries = self.entries()
        self.assertEqual(len(entries),9)
        lazada = [e for e in entries if e['platform'] == 'LAZADA']
        self.assertTrue(lazada)
        self.assertTrue(all(e['status'] == 'REVIEW' and e['requires_revision'] for e in lazada))
        self.assertTrue(all('no verified Tracking ID' in e['message'] for e in lazada))
        self.assertTrue(all(e['status'] == 'DONE' for e in entries if e['platform'] != 'LAZADA'))
        # The old text log has no source hashes: identical item contents alone
        # cannot prove which tracking number an earlier Lazada table covered.
        with self.assertRaisesRegex(PackingError, 'explicit revision review'):
            engine.generate(self.folder, [{'path': e['path'], 'sha256': e['sha256']} for e in lazada], verifier=fake_layout)
        self.assertEqual(before,file_hash(log))
        history=json.loads((self.folder/'000 Output'/engine.HISTORY).read_text())
        self.assertEqual(len(history['outputs']),3)
        self.assertTrue(all(o['sources'] and o['order_fingerprints'] for o in history['outputs'].values()))
        self.assertTrue(all(not o.get('tracking_order_keys') for o in history['outputs'].values()))
        self.assertEqual(documents_before, {p.name: file_hash(p) for p in out.iterdir() if p.suffix.lower() in ('.docx', '.pdf')})

    def test_three_output_groups_and_later_suffix_without_overwrite(self):
        self.copy_source('20260926_SPX_PL_1.pdf')
        self.copy_source('20260926_Flash_PL_1.pdf','FLASH')
        self.copy_source('20260926_Lazada_PL_1.pdf','LAZADA & TIKTOK')
        self.copy_source('20260926_Tiktok_PL_1.pdf','LAZADA & TIKTOK')
        result = generate_approved(self.folder,self.select(),verifier=fake_layout)
        self.assertEqual(len(result['outputs']),3)
        before = {p['filename']:file_hash(p['path']) for p in result['outputs']}
        combined = Document(self.folder/'000 Output'/'LAZADA_TIKTOK_Packing_Checklist_2026-09-26_1.docx')
        self.assertEqual(len(combined.tables),8)
        self.assertTrue(all(e['status']=='DONE' for e in self.entries()))
        self.copy_source('20260926_SPX_PL_4.pdf')
        shutil.copyfile(self.folder/'SPX/20260926_SPX_PL_1.pdf',self.folder/'SPX/renamed.pdf')
        states = {e['path']:e['status'] for e in self.entries()}
        self.assertEqual(states['SPX/renamed.pdf'],'DUPLICATE')
        result = generate_approved(self.folder,self.select(),verifier=fake_layout)
        self.assertEqual(len(result['outputs']),1)
        self.assertEqual(result['outputs'][0]['filename'],'SPX_Packing_Checklist_2026-09-26_2.docx')
        self.assertEqual(len(Document(result['outputs'][0]['path']).tables),8)
        for name,hash_value in before.items():self.assertEqual(hash_value,file_hash(self.folder/'000 Output'/name))
        self.assertIn('SPX/20260926_SPX_PL_4.pdf',(self.folder/'000 Output'/engine.LOG).read_text())

    def test_changed_binary_requires_review_and_explicit_ignore(self):
        path = self.copy_source('20260926_SPX_PL_1.pdf')
        generate_approved(self.folder,self.select(),verifier=fake_layout)
        with path.open('ab') as f:f.write(b'\n% revised source metadata\n')
        entry = self.entries()[0]
        self.assertEqual(entry['status'],'REVIEW')
        selected={'path':entry['path'],'sha256':entry['sha256']}
        with self.assertRaises(PackingError):generate_approved(self.folder,[selected],verifier=fake_layout)
        engine.acknowledge(self.folder,selected)
        self.assertEqual(self.entries()[0]['status'],'DONE')
        self.assertEqual(len(list((self.folder/'000 Output').glob('*.docx'))),1)

    def test_reviewed_revision_creates_next_file_and_supersedes_original(self):
        path=self.copy_source('20260926_SPX_PL_1.pdf')
        first=generate_approved(self.folder,self.select(),verifier=fake_layout)['outputs'][0]
        with path.open('ab') as f:f.write(b'\n% updated export\n')
        entry=self.entries()[0]
        selection={'path':entry['path'],'sha256':entry['sha256'],'revision':True}
        new=generate_approved(self.folder,[selection],verifier=fake_layout)['outputs'][0]
        self.assertTrue(new['filename'].endswith('_2.docx'))
        self.assertIn('REVISION',Document(new['path']).paragraphs[1].text)
        h=json.loads((self.folder/'000 Output'/engine.HISTORY).read_text())
        self.assertEqual(h['outputs'][first['filename']]['status'],'superseded')

    def test_file_changed_since_preview_and_failed_word_check_commit_nothing(self):
        path=self.copy_source('20260926_SPX_PL_1.pdf')
        selection=self.select()
        with path.open('ab') as f:f.write(b'\n% changed while reviewing\n')
        with self.assertRaises(PackingError):generate_approved(self.folder,selection,verifier=fake_layout)
        def fail(_):raise PackingError('Word unavailable')
        with self.assertRaises(PackingError):generate_approved(self.folder,self.select(),verifier=fail)
        out=self.folder/'000 Output'
        self.assertFalse(list(out.glob('*.docx')))
        self.assertFalse((out/engine.HISTORY).exists())
        self.assertFalse((out/engine.LOG).exists())

    def test_interrupted_commit_recovers_without_extra_batch(self):
        path=self.copy_source('20260926_SPX_PL_1.pdf')
        out=self.folder/'000 Output';out.mkdir()
        stage=out/'.pending-recovery';stage.mkdir()
        dest=stage/'SPX_Packing_Checklist_2026-09-26_1.docx'
        batch=parse_pdf(path);build_document([batch],'SPX','20260926',dest)
        history={'schema':1,'date':'20260926','sources':{},'orders':{},'outputs':{}}
        plan={'id':'test-recovery','files':[{'target':dest.name,'stage':dest.relative_to(out).as_posix(),'sha256':file_hash(dest)}],
              'history':history,'stage_dir':stage.name,'log':'Recovery test'}
        engine.atomic_json(out/'.packing-transaction.json',plan)
        with engine.folder_lock(self.folder):
            engine.recover(out);engine.recover(out)
        self.assertTrue((out/dest.name).exists())
        self.assertFalse((out/engine.JOURNAL).exists())
        self.assertEqual((out/engine.LOG).read_text().count('Transaction: test-recovery'),1)

    def test_import_keeps_original_and_detects_duplicate(self):
        source=ROOT/'20260926/SPX/20260926_SPX_PL_1.pdf';before=file_hash(source)
        r=engine.import_pdfs(self.folder,[str(source)])
        self.assertEqual(len(r['imported']),1)
        self.assertEqual(file_hash(source),before)
        r=engine.import_pdfs(self.folder,[str(source)])
        self.assertEqual(len(r['skipped']),1)
        self.assertFalse(r['imported'])

    def test_unsupported_pdf_is_hidden_without_partial_output(self):
        folder=self.folder/'SPX';folder.mkdir()
        writer=PdfWriter();writer.add_blank_page(width=595,height=842)
        with (folder/'blank.pdf').open('wb') as f:writer.write(f)
        entries=self.entries()
        self.assertEqual(entries,[])

    def test_missing_app_history_is_not_silently_recreated(self):
        self.copy_source('20260926_SPX_PL_1.pdf')
        generate_approved(self.folder,self.select(),verifier=fake_layout)
        (self.folder/'000 Output'/engine.HISTORY).unlink()
        with self.assertRaisesRegex(PackingError,'Restore'):self.entries()

    def test_explicit_courier_selection_overrides_filename(self):
        r=engine.import_pdfs(self.folder,[str(ROOT/'20260926/FLASH/20260926_Flash_PL_1.pdf')],carrier='SPX')
        self.assertEqual(len(r['imported']),1)
        self.assertFalse(r['errors'])
        self.assertEqual(self.entries()[0]['platform'],'SPX')
        self.assertEqual(self.entries()[0]['courier'],'SPX')

    def test_revision_tracks_orders_removed_from_source(self):
        path=self.copy_source('20260926_SPX_PL_1.pdf')
        original=generate_approved(self.folder,self.select(),verifier=fake_layout)['outputs'][0]['filename']
        shutil.copyfile(ROOT/'20260926/SPX/20260926_SPX_PL_4.pdf',path)
        entry=self.entries()[0]
        result=generate_approved(self.folder,[{'path':entry['path'],'sha256':entry['sha256'],'revision':True}],verifier=fake_layout)
        history=json.loads((self.folder/'000 Output'/engine.HISTORY).read_text())
        self.assertEqual(history['outputs'][original]['status'],'superseded')
        self.assertEqual(len(history['outputs'][result['outputs'][0]['filename']]['removed_orders']),8)


if __name__=='__main__':unittest.main(verbosity=2)
